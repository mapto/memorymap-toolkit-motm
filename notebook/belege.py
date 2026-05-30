"""Belege (Extraction) import logic.

Reads belege_*.xlsx files and imports them into the Extraction model
via the Django REST API.
"""

from __future__ import annotations

from pathlib import Path

from fast_langdetect import detect as detect_language
from openpyxl import load_workbook
from tqdm.auto import tqdm

from api_client import (
    api_get,
    api_patch,
    api_post,
    clean,
    find_person_by_identifier,
    get_interview_id,
    get_or_create_concept,
    is_empty,
)


def import_belege(xlsx_path: str | Path, sheet_name: str = "Belege") -> list[int]:
    """Import all rows from a belege xlsx file via the API."""
    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path)
    sheet = wb[sheet_name]

    # Count data rows (skip first 3: title, blank, header)
    all_rows = list(sheet.iter_rows(values_only=True))
    data_rows = [r for r in all_rows[3:] if r and not all(cell is None for cell in r)]

    results: list[int] = []
    errors: list[str] = []

    for row in tqdm(data_rows, desc=xlsx_path.stem, unit="row"):
        cells = (list(row) + [None] * 13)[:13]
        beleg_id = clean(cells[0])  # identifier
        quelle = clean(cells[1])  # source reference (Markdown link)
        interview_id = clean(cells[2])  # interview archive_id
        # cells[3]: interview_datum (informational only)
        quelle_sprecher = clean(cells[4])  # speaker → Interview.interviewee
        betrifft = clean(cells[5])  # people mentioned (person identifiers)
        timecode = clean(cells[6])  # timecode
        themen = clean(cells[7])  # topics/concepts
        zitat = clean(cells[8])  # quote
        markierung = clean(cells[9])  # classification
        # cells[10]: event_ids (not yet linked)
        event_confidence = clean(cells[11])
        notizen = clean(cells[12])  # notes

        if is_empty(beleg_id):
            continue

        # Check if already imported
        existing = api_get("extractions", params={"search": beleg_id})
        if any(e["identifier"] == beleg_id for e in existing):
            continue

        # Resolve person mentioned (first identifier)
        person_id = None
        if betrifft:
            first_person_id = betrifft.split(",")[0].strip()
            person_id = find_person_by_identifier(first_person_id)

        # Resolve interview
        interview_db_id = get_interview_id(interview_id, quelle=quelle)

        # Link speaker → Person → Interview.interviewee
        if not is_empty(quelle_sprecher) and interview_db_id:
            speaker_person_id = find_person_by_identifier(quelle_sprecher)
            if speaker_person_id:
                try:
                    api_patch(
                        "interviews",
                        interview_db_id,
                        {"interviewee": speaker_person_id},
                    )
                except Exception as e:
                    errors.append(f"interviewee {interview_id}: {e}")

        # Resolve all concepts
        concept_ids = []
        if themen:
            for topic in themen.split(","):
                topic = topic.strip()
                if not is_empty(topic):
                    cid = get_or_create_concept(topic)
                    if cid:
                        concept_ids.append(cid)

        payload: dict = {
            "identifier": beleg_id,
            "timecode": timecode,
            "quote": zitat,
            "classification": markierung,
            "event_confidence": event_confidence,
            "notes": notizen,
        }
        # Detect language (only if long enough and confidence > 70%)
        if len(zitat) >= 20:
            try:
                lang_result = detect_language(zitat)
                if lang_result and float(lang_result[0]["score"]) > 0.7:
                    payload["language"] = str(lang_result[0]["lang"])
            except Exception:
                pass
        if person_id:
            payload["people_mentioned"] = person_id
        if interview_db_id:
            payload["interview"] = interview_db_id
        if concept_ids:
            payload["concepts"] = concept_ids

        try:
            created = api_post("extractions", payload)
            results.append(created["id"])
        except Exception as e:
            errors.append(f"{beleg_id}: {e}")

    if errors:
        print(f"\n\u26a0 {len(errors)} errors:")
        for err in errors:
            print(f"  \u274c {err}")
    print(f"\u2705 Imported {len(results)} extractions from {xlsx_path.name}")
    return results


def import_all_belege(directory: str | Path | None = None) -> list[int]:
    """Find and import all belege_*.xlsx files from the data directory."""
    if directory is None:
        directory = Path.cwd().parent / "data" / "schede mappatura"
    directory = Path(directory)
    all_results: list[int] = []
    for xlsx_path in sorted(directory.rglob("belege_*.xlsx")):
        print(f"\nProcessing: {xlsx_path.name}")
        try:
            ids = import_belege(xlsx_path)
            all_results.extend(ids)
        except Exception as e:
            print(f"\u274c ERROR: {e}")
    print(f"\n\u2705 Total imported: {len(all_results)} extractions")
    return all_results


def backfill_languages(
    min_chars: int = 20, min_confidence: float = 0.7
) -> tuple[int, int]:
    """Detect and set language on extractions that don't have one yet."""
    extractions = api_get("extractions")
    updated = 0
    skipped = 0

    for ext in tqdm(extractions, desc="Language detect", unit="ext"):
        if ext.get("language", ""):
            continue
        quote = ext.get("quote", "").strip()
        if len(quote) < min_chars:
            skipped += 1
            continue
        result = detect_language(quote)
        if result and float(result[0]["score"]) > min_confidence:
            lang = str(result[0]["lang"])
            api_patch("extractions", ext["id"], {"language": lang})
            updated += 1
        else:
            skipped += 1

    print(f"\u2705 Updated {updated} extractions, skipped {skipped}")
    return updated, skipped
