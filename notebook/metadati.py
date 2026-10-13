"""Metadati (person metadata) import logic.

Reads metadati_*.xlsx files and imports person records, relationships,
interviews, and birth places via the Django REST API.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from api_client import (
    BASE_URL,
    SESSION,
    api_get,
    api_patch,
    api_post,
    build_interview_payload,
    get_or_create_url,
    upsert_interviews,
)
from utils import is_empty
from locations import (
    get_or_create_location as get_or_create_location_point,
    load_locations_db,
)
from timespan import parse_date, parse_date_range, parse_timespan
from person_taxonomy import (
    canonical_person_name,
    is_placeholder_name,
    metadata_relative_name,
    person_taxonomy_alias,
)


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

RELATIONSHIP_MAP = {
    "vater": "father",
    "mutter": "mother",
    "schwester": "sister",
    "bruder": "brother",
    "großvater": "grandfather",
    "grossvater": "grandfather",
    "großvater väterlicherseits": "paternal grandfather",
    "grossvater väterlicherseits": "paternal grandfather",
    "großvater väterl": "paternal grandfather",
    "großmutter väterlicherseits": "paternal grandmother",
    "großmutter väterl": "paternal grandmother",
    "grossmutter väterlicherseits": "paternal grandmother",
    "großmutter mütterlicherseits": "maternal grandmother",
    "grossmutter mütterlicherseits": "maternal grandmother",
    "großmutter mütt": "maternal grandmother",
    "großvater mütterlicherseits": "maternal grandfather",
    "großvater mütt": "maternal grandfather",
    "tante": "aunt",
    "onkel": "uncle",
    "onkel väterl": "paternal uncle",
    "bruder aus erster ehe": "brother from the first marriage",
    "halbschwester": "half sister",
    "1 frau erste ehe": "first wife",
    "stiefmutter": "stepmother",
    "cousin": "cousin",
    "vetter": "cousin",
    "cousine": "cousin",
    "base": "cousin",
}


# ---------------------------------------------------------------------------
# Caches
# ---------------------------------------------------------------------------

_timespan_cache: dict = {}
_person_cache: dict = {}
_relationship_type_cache: dict = {}
_relationship_cache: dict = {}
_location_workbook_cache = None

PERSON_METADATA_FIELDS = (
    "education",
    "occupation_at_emigration",
    "occupation_at_source",
    "residence_at_source",
)


# ---------------------------------------------------------------------------
# Utility helpers (metadati-specific)
# ---------------------------------------------------------------------------


def clean_label(label):
    label = label.strip().lower()
    label = label.replace("(", "").replace(")", "")
    label = label.replace(".", "")
    label = " ".join(label.split())
    return label


def clean_value(value):
    if not value:
        return None
    value = str(value).strip()
    if is_empty(value):
        return None
    return value


def clean_metadata_value(value):
    value = clean_value(value)
    return " ".join(value.split()) if value else None


def _date_interval(raw):
    """Preserve the precision interval represented by one metadata date."""
    parsed = parse_timespan(raw)
    if parsed.start is None:
        return None
    end = parsed.end or parsed.start
    return parsed.start.isoformat(), end.isoformat()


def _split_biographical_date_range(raw):
    """Return separate birth/death intervals from a metadata family date cell."""
    value = clean_value(raw)
    if not value:
        return None, None

    separators = [r"\s*[–—]\s*", r"\s+-\s+", r"-(?=[A-Za-zÄÖÜäöü])"]
    if re.fullmatch(r"\d{4}-\d{2,4}\??", value):
        separators.append(r"-")
    if not re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}\*?", value):
        separators.append(r"\s*/\s*")

    for separator in separators:
        match = re.search(separator, value)
        if not match:
            continue
        birth_interval = _date_interval(value[: match.start()])
        death_interval = _date_interval(value[match.end() :])
        if birth_interval and death_interval:
            return birth_interval, death_interval

    return _date_interval(value), None


def parse_coordinates(text):
    match = re.search(r"([\d.]+)\s*/\s*([\d.]+)", text)
    if match:
        return {"lat": float(match.group(1)), "lon": float(match.group(2))}
    return None


def parse_name(full_name):
    if not full_name:
        return None, None
    full_name = full_name.strip()
    if "," in full_name:
        parts = full_name.split(",")
        family_name = parts[0].strip()
        given_name = parts[1].strip() if len(parts) > 1 else "UNKNOWN"
        return given_name, family_name
    parts = full_name.split()
    if len(parts) > 1:
        return parts[0], " ".join(parts[1:])
    return full_name, "UNKNOWN"


def parse_source_label(value):
    """Extract interviewer text and corpus ID from a metadata source label."""
    source_label = clean_value(value)
    if not source_label:
        return None, None
    match = re.search(r"\b([A-Z]+_[A-Z]_\d+)\b", source_label)
    if not match:
        return source_label, None
    interviewer = source_label[: match.start()].strip(" ,\u2013-+") or None
    return interviewer, match.group(1)


# ---------------------------------------------------------------------------
# Timespan (metadati-specific: searches API before creating)
# ---------------------------------------------------------------------------


def get_or_create_timespan(start=None, end=None):
    """Get or create a Timespan singleton. Returns id or None."""
    if start is None and end is None:
        return None
    key = (start, end)
    if key in _timespan_cache:
        return _timespan_cache[key]
    existing = api_get("timespans")
    for ts in existing:
        if ts.get("start") == start and ts.get("end") == end:
            _timespan_cache[key] = ts["id"]
            return ts["id"]
    created = api_post("timespans", {"start": start, "end": end})
    _timespan_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# XLSX reader
# ---------------------------------------------------------------------------


def _read_sheet_rows(sheet):
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        field = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        value = row[1] if len(row) > 1 else None
        extra = row[2] if len(row) > 2 else None
        rows.append((field, value, extra))
    return rows


def _parse_identifier(field, value_str):
    if field == "ID Sprecher":
        return value_str

    combined = " ".join([x for x in [field, value_str] if x]).replace("\u2013", "-")
    match = re.search(r"Metadaten\s*-\s*([A-Za-z0-9_]+)", combined)
    return match.group(1) if match else None


def _update_birth_place(birth_place, value_str, extra_str):
    metadata = " | ".join(value for value in (value_str, extra_str) if value)
    name_match = re.search(
        r"(?:^|\||\n)\s*Ortsname\s*:\s*([^|\r\n]+)",
        metadata,
        flags=re.IGNORECASE,
    )
    birth_place["name"] = (
        name_match.group(1).strip() if name_match else value_str.splitlines()[0].strip()
    )

    for part in (p.strip() for p in re.split(r"\||\r?\n", metadata) if p.strip()):
        key, separator, detail = part.partition(":")
        if not separator:
            continue
        normalized_key = key.strip().casefold()
        detail = detail.strip()
        if normalized_key in {"land", "region", "kreis"}:
            birth_place["regions"].append(detail)
        elif "koordinat" in normalized_key:
            birth_place["coordinates"] = parse_coordinates(detail)
        elif normalized_key in {"wikidata", "wikidata id"}:
            birth_place["wikidata_id"] = detail


def _parse_person_field(person, birth_place, field, value_str, extra_str):
    if field == "Vorname":
        person["given_name"] = value_str
    elif field in {"ehem. Vorname", "ehemaliger Vorname"}:
        person["previous_given_name"] = value_str
    elif field == "Nachname":
        person["family_name"] = value_str
    elif field in {"ehem. Nachname", "ehemaliger Nachname"}:
        person["previous_family_name"] = value_str
    elif field == "Geschlecht":
        person["gender"] = value_str
    elif field == "Geburtsdatum":
        person["birth_date"] = parse_date(value_str) if value_str else None
        person["birth_date_interval"] = _date_interval(value_str)
    elif field in {"Verfolgtengruppe NS", "Verfolgtengruppe(n) NS"}:
        person["attributes"]["ns_persecution_group"] = value_str
    elif field == "Schulbildung / Abschluss":
        person["education"] = clean_metadata_value(value_str)
    elif field == "Beruf (zum Zeitpunkt der Emigration)":
        person["occupation_at_emigration"] = clean_metadata_value(value_str)
    elif field == "Beruf (zum Zeitpunkt der Quelle)":
        person["occupation_at_source"] = clean_metadata_value(value_str)
    elif field == "Wohnort zum Zeitpunkt der Quelle":
        person["residence_at_source"] = clean_metadata_value(value_str)
    elif field == "Geburtsort":
        _update_birth_place(birth_place, value_str, extra_str)


def _new_interview():
    return {
        "type": None,
        "place": None,
        "date": None,
        "interviewer": None,
        "archive_id": None,
        "url": None,
    }


def _parse_source_field(field, value, value_str, interviews):
    match = re.match(r"Quelle\s+(\d+)\s*[-\u2013\u2014]\s*(.+)", field)
    if not match:
        return

    index = int(match.group(1))
    subfield = match.group(2)
    while len(interviews) < index:
        interviews.append(_new_interview())
    current = interviews[index - 1]

    if subfield == "Typ":
        current["type"] = clean_value(value)
    elif subfield == "Ort":
        current["place"] = clean_value(value)
    elif subfield == "Datum":
        current["date"] = parse_date(value_str) if value_str else None
    elif subfield == "Gesprächspartner/in":
        current["interviewer"] = clean_value(value)
    elif subfield == "Bezeichnung":
        interviewer, archive_id = parse_source_label(value)
        current["interviewer"] = interviewer
        current["archive_id"] = archive_id
    elif subfield == "Handle/URL":
        current["url"] = clean_value(value)
    elif subfield == "ID":
        current["archive_id"] = clean_value(value)


def _parse_person_sheet(rows):
    person = {
        "given_name": None,
        "previous_given_name": None,
        "family_name": None,
        "previous_family_name": None,
        "gender": None,
        "birth_date": None,
        "birth_date_interval": None,
        "education": None,
        "occupation_at_emigration": None,
        "occupation_at_source": None,
        "residence_at_source": None,
        "attributes": {},
    }
    birth_place = {
        "name": None,
        "coordinates": None,
        "wikidata_id": None,
        "regions": [],
    }
    interviews = []
    identifier = None
    in_sources_section = False

    for field, value, extra in rows:
        value_str = clean_value(value)
        extra_str = clean_value(extra)

        if identifier is None:
            identifier = _parse_identifier(field, value_str)

        if field == "QUELLEN":
            in_sources_section = True
            continue
        if field == "MIGRATION":
            in_sources_section = False

        _parse_person_field(person, birth_place, field, value_str, extra_str)

        if in_sources_section or field.startswith("Quelle "):
            _parse_source_field(field, value, value_str, interviews)

    interviews = [
        interview
        for interview in interviews
        if any(
            [
                interview["type"],
                interview["place"],
                interview["date"],
                interview["interviewer"],
                interview["archive_id"],
                interview["url"],
            ]
        )
    ]
    return identifier, person, birth_place, interviews


def _parse_family_sheet(sheet):
    family = []
    for row in sheet.iter_rows(min_row=4, max_col=7, values_only=True):
        if all(cell is None for cell in row):
            continue
        (
            _,
            relation,
            given_name,
            family_name,
            birth_date_raw,
            birth_place,
            notes,
        ) = row
        birth_date_interval, death_date_interval = _split_biographical_date_range(
            birth_date_raw
        )
        birth_date, death_date = (
            parse_date_range(str(birth_date_raw)) if birth_date_raw else (None, None)
        )
        family.append(
            {
                "relation": clean_value(relation),
                "given_name": clean_value(given_name),
                "family_name": clean_value(family_name),
                "birth_date": birth_date,
                "birth_date_interval": birth_date_interval,
                "death_date": death_date,
                "death_date_interval": death_date_interval,
                "birth_place": clean_value(birth_place),
                "notes": clean_value(notes),
            }
        )
    return family


def read_person_record(xlsx_path, sheet_name="Scheda individuale"):
    """Parse a metadati xlsx file into a structured record dict."""
    wb = load_workbook(xlsx_path)
    sheet = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    assert sheet is not None, f"No sheet found in {xlsx_path}"

    identifier, person, birth_place, interviews = _parse_person_sheet(
        _read_sheet_rows(sheet)
    )
    family = _parse_family_sheet(wb["Familie"]) if "Familie" in wb.sheetnames else []

    return {
        "identifier": identifier,
        "person": person,
        "birth_place": birth_place,
        "interviews": interviews,
        "family": family,
    }


# ---------------------------------------------------------------------------
# Location (birth place)
# ---------------------------------------------------------------------------


def get_or_create_region_hierarchy(region_list):
    """Create or retrieve a chain of LocationRegion objects via the API.
    Returns the innermost region's id, or None."""
    if not region_list:
        return None

    parent_id = None
    for name in region_list:
        if is_empty(name):
            continue

        existing = api_get("regions", params={"search": name})
        match = next((r for r in existing if r["name"] == name), None)

        if match:
            region_id = match["id"]
            if parent_id and match.get("part_of") != parent_id:
                api_patch("regions", region_id, {"part_of": parent_id})
        else:
            payload = {"name": name}
            if parent_id:
                payload["part_of"] = parent_id
            created = api_post("regions", payload)
            region_id = created["id"]

        parent_id = region_id

    return parent_id


def get_or_create_location(bp):
    """Create or retrieve a LocationPoint from a birth_place dict."""
    if is_empty(bp.get("name")):
        return None
    global _location_workbook_cache
    if _location_workbook_cache is None:
        _location_workbook_cache = load_locations_db(enrich_coordinates=False)
    loc_id = get_or_create_location_point(
        bp["name"],
        wikidata_qid=bp.get("wikidata_id"),
        locations_db=_location_workbook_cache,
    )
    if loc_id is None:
        return None
    region_id = get_or_create_region_hierarchy(bp.get("regions"))
    if region_id:
        api_patch("locations", loc_id, {"region": region_id})
    return loc_id


# ---------------------------------------------------------------------------
# Person matching
# ---------------------------------------------------------------------------


def _data_date_interval(data, field_name, value):
    interval = data.get(f"{field_name}_date_interval")
    if isinstance(interval, (tuple, list)) and len(interval) == 2:
        start = parse_timespan(interval[0])
        end = parse_timespan(interval[1])
        if start.start and end.start:
            return start.start, end.end or end.start
    parsed = parse_timespan(value)
    if parsed.start is None:
        return None
    return parsed.start, parsed.end or parsed.start


def _compatible_date_or_existing(current, incoming, interval, field_name, person_id):
    if not incoming:
        return current
    if not current:
        return incoming
    if current == incoming:
        return current

    existing = parse_timespan(current)
    existing_start = existing.start
    existing_end = existing.end or existing.start
    incoming_start, incoming_end = interval or (None, None)
    if (
        existing_start
        and existing_end
        and incoming_start
        and incoming_end
        and existing_start <= incoming_end
        and incoming_start <= existing_end
    ):
        # Keep an existing, more specific value when an imported partial date
        # merely confirms that it falls within the same precision interval.
        return current

    raise ValueError(
        f"Refusing to overwrite Person {person_id} {field_name} date "
        f"{current!r} with conflicting metadata value {incoming!r}"
    )


def _maybe_update_lifespan(person_id, data):
    """Fill a person's lifespan without overwriting conflicting existing dates."""
    birth = data.get("birth_date")
    death = data.get("death_date")
    if not birth and not death:
        return
    person = SESSION.get(f"{BASE_URL}/persons/{person_id}/").json()
    existing_ts_id = person.get("lifespan")
    cur_start = None
    cur_end = None
    if existing_ts_id:
        ts = SESSION.get(f"{BASE_URL}/timespans/{existing_ts_id}/").json()
        cur_start = ts.get("start")
        cur_end = ts.get("end")
        new_start = _compatible_date_or_existing(
            cur_start,
            birth,
            _data_date_interval(data, "birth", birth),
            "birth",
            person_id,
        )
        new_end = _compatible_date_or_existing(
            cur_end,
            death,
            _data_date_interval(data, "death", death),
            "death",
            person_id,
        )
        if new_start == cur_start and new_end == cur_end:
            return  # no change
    else:
        new_start = birth
        new_end = death
    new_ts_id = get_or_create_timespan(start=new_start, end=new_end)
    if new_ts_id and new_ts_id != existing_ts_id:
        SESSION.patch(f"{BASE_URL}/persons/{person_id}/", json={"lifespan": new_ts_id})
        # Remove old timespan from cache and delete if orphaned
        if existing_ts_id:
            old_key = (cur_start, cur_end)
            _timespan_cache.pop(old_key, None)
            # Delete orphaned timespan (no other person references it after update)
            referencing = api_get("persons", params={"search": ""})
            still_used = any(p.get("lifespan") == existing_ts_id for p in referencing)
            if not still_used:
                SESSION.delete(f"{BASE_URL}/timespans/{existing_ts_id}/")


def get_or_create_person(data, identifier=None):
    """Create or match a Person via the API. Returns person id or None."""
    source_given_name = data.get("given_name")
    source_family_name = data.get("family_name")
    canonical = canonical_person_name(source_given_name, source_family_name)
    if canonical is not None:
        given_name, family_name = canonical.given_name, canonical.family_name
    else:
        given_name, family_name = source_given_name, source_family_name
    if (given_name, family_name) != (source_given_name, source_family_name):
        data = dict(data)
        if (
            not data.get("previous_given_name")
            and source_given_name
            and source_given_name != given_name
        ):
            data["previous_given_name"] = source_given_name
        if (
            not data.get("previous_family_name")
            and source_family_name
            and source_family_name != family_name
        ):
            data["previous_family_name"] = source_family_name
        data["given_name"] = given_name
        data["family_name"] = family_name
    source_name = " ".join(
        part for part in (source_given_name, source_family_name) if part
    )
    alias_target = person_taxonomy_alias(source_name)

    if is_empty(given_name) and is_empty(family_name):
        return None

    def metadata_payload():
        payload = {
            field: data[field]
            for field in PERSON_METADATA_FIELDS
            if not is_empty(data.get(field))
        }
        prosecuted_group = data.get("attributes", {}).get("ns_persecution_group")
        if not is_empty(prosecuted_group):
            payload["description"] = prosecuted_group
        return payload

    def update_existing(existing):
        patch = metadata_payload()
        if existing:
            if data.get("gender") and not existing.get("gender"):
                patch["gender"] = data["gender"]
            if identifier and not existing.get("identifier"):
                patch["identifier"] = identifier
            if existing.get("given_name") != given_name:
                patch["given_name"] = given_name or ""
            if existing.get("family_name") != family_name:
                patch["family_name"] = family_name or ""
            if data.get("previous_given_name") and not existing.get(
                "previous_given_name"
            ):
                patch["previous_given_name"] = data["previous_given_name"]
            if data.get("previous_family_name") and not existing.get(
                "previous_family_name"
            ):
                patch["previous_family_name"] = data["previous_family_name"]
        if patch:
            api_patch("persons", existing["id"], patch)

    if alias_target:
        from api_client import get_person_id

        person_id = get_person_id(source_name, source_name, identifier=identifier)
        if not person_id:
            return None
        existing = api_get(f"persons/{person_id}")
        _maybe_update_lifespan(person_id, data)
        update_existing(existing)
        _person_cache[(given_name or "").lower(), (family_name or "").lower()] = (
            person_id
        )
        return person_id

    # Match on identifier
    if identifier:
        existing = api_get("persons", params={"search": identifier})
        match = next((p for p in existing if p.get("identifier") == identifier), None)
        if not match and not is_empty(given_name) and not is_empty(family_name):
            existing = api_get("persons", params={"search": family_name})
            match = next(
                (
                    p
                    for p in existing
                    if not p.get("identifier")
                    and p.get("given_name", "").casefold() == given_name.casefold()
                    and p.get("family_name", "").casefold() == family_name.casefold()
                ),
                None,
            )
            if match:
                api_patch("persons", match["id"], {"identifier": identifier})
        if match:
            key = (match["given_name"].lower(), match["family_name"].lower())
            _person_cache[key] = match["id"]
            _maybe_update_lifespan(match["id"], data)
            update_existing(match)
            return match["id"]
        payload = {
            "identifier": identifier,
            "given_name": given_name or "",
            "family_name": family_name or "",
        }
        payload.update(metadata_payload())
        lifespan_id = get_or_create_timespan(
            start=data.get("birth_date"),
            end=data.get("death_date"),
        )
        if lifespan_id:
            payload["lifespan"] = lifespan_id
        if data.get("gender"):
            payload["gender"] = data["gender"]
        if data.get("previous_given_name"):
            payload["previous_given_name"] = data["previous_given_name"]
        if data.get("previous_family_name"):
            payload["previous_family_name"] = data["previous_family_name"]
        if data.get("attributes", {}).get("ns_persecution_group"):
            payload["description"] = data["attributes"]["ns_persecution_group"]
        created = api_post("persons", payload)
        key = ((given_name or "").lower(), (family_name or "").lower())
        _person_cache[key] = created["id"]
        return created["id"]

    # Check local cache first
    key = ((given_name or "").lower(), (family_name or "").lower())
    if not is_empty(given_name) and not is_empty(family_name):
        if key in _person_cache:
            _maybe_update_lifespan(_person_cache[key], data)
            api_patch("persons", _person_cache[key], metadata_payload())
            return _person_cache[key]

    # Match on name (case-insensitive) via API
    if not is_empty(given_name) and not is_empty(family_name):
        existing = api_get("persons", params={"search": family_name})
        for p in existing:
            if (
                p["given_name"].lower() == given_name.lower()
                and p["family_name"].lower() == family_name.lower()
            ):
                _person_cache[key] = p["id"]
                _maybe_update_lifespan(p["id"], data)
                update_existing(p)
                return p["id"]

    payload = {
        "given_name": given_name or "",
        "family_name": family_name or "",
    }
    payload.update(metadata_payload())
    if data.get("gender"):
        payload["gender"] = data["gender"]
    lifespan_id = get_or_create_timespan(
        start=data.get("birth_date"),
        end=data.get("death_date"),
    )
    if lifespan_id:
        payload["lifespan"] = lifespan_id
    created = api_post("persons", payload)
    key = ((given_name or "").lower(), (family_name or "").lower())
    _person_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# Interview
# ---------------------------------------------------------------------------


def create_interview(person_id, data):
    """Upsert an Interview from one metadata source record.

    Person-specific parsing stays here, while archive-ID normalization and the
    actual Interview write are shared with the Belege importer.
    """
    interviewer_id = None
    if not is_empty(data.get("interviewer")):
        given, family = parse_name(data["interviewer"])
        interviewer_id = get_or_create_person(
            {
                "given_name": given,
                "family_name": family,
            }
        )

    url_id = get_or_create_url(data.get("url")) if data.get("url") else None

    payload = build_interview_payload(
        data.get("archive_id"),
        interviewee=person_id,
        interviewer=interviewer_id,
        interview_type=data.get("type"),
        interview_date=data.get("date"),
        place=data.get("place"),
        urls=[url_id] if url_id else None,
    )
    if not payload:
        return None

    interview_ids, errors = upsert_interviews([payload])
    for error in errors:
        raise RuntimeError(error)
    return interview_ids.get(payload["archive_id"])


# ---------------------------------------------------------------------------
# Relationship
# ---------------------------------------------------------------------------


def get_or_create_relationship_type(label):
    """Get or create a RelationshipType via the API."""
    if is_empty(label):
        return None

    label_clean = clean_label(label)
    mapped = RELATIONSHIP_MAP.get(label_clean, label_clean)

    if mapped in _relationship_type_cache:
        return _relationship_type_cache[mapped]

    existing = api_get("relationship-types", params={"search": mapped})
    match = next((rt for rt in existing if rt["name"] == mapped), None)
    if match:
        _relationship_type_cache[mapped] = match["id"]
        return match["id"]

    created = api_post(
        "relationship-types",
        {
            "name": mapped,
            "original_label": label.strip(),
        },
    )
    _relationship_type_cache[mapped] = created["id"]
    return created["id"]


def _is_unnamed_relative(person, family_name):
    """True for a placeholder record that stands for the same relative."""
    if not is_placeholder_name(person.get("given_name")):
        return False
    stored_family = person.get("family_name")
    return is_placeholder_name(stored_family) or (
        str(stored_family).strip().casefold() == family_name.casefold()
    )


def _merge_unnamed_relatives(
    main_person_id, rel_type_id, canonical_id, given_name, family_name
):
    """Merge placeholder duplicates of a linked relative into the canonical person.

    Each duplicate is renamed before the merge, so the backend merge does not
    copy the placeholder into the canonical person's previous-name fields.
    Records with an identifier are never merged automatically.
    """
    duplicate_ids = []
    for relationship in api_get("relationships"):
        if (
            relationship["person_from"] != main_person_id
            or relationship["relationship_type"] != rel_type_id
            or relationship["person_to"] in (canonical_id, *duplicate_ids)
        ):
            continue
        candidate = api_get(f"persons/{relationship['person_to']}")
        if candidate.get("identifier") or not _is_unnamed_relative(
            candidate, family_name
        ):
            continue
        duplicate_ids.append(candidate["id"])
    if not duplicate_ids:
        return
    for duplicate_id in duplicate_ids:
        api_patch(
            "persons",
            duplicate_id,
            {"given_name": given_name, "family_name": family_name},
        )
    api_post(f"persons/{canonical_id}/merge", {"duplicate_ids": duplicate_ids})


def create_relationship(
    main_person_id, data, fallback_family_name, source_workbook=None
):
    """Create a Relationship via the API.

    An unnamed relative with a reviewed link for the source workbook is attached
    to that Personae person, and placeholder duplicates of the same relation
    are merged into it.
    """
    if is_empty(data.get("relation")):
        return None

    given_name = data.get("given_name")
    family_name = data.get("family_name")
    if is_empty(family_name):
        family_name = fallback_family_name

    linked_name = None
    if is_placeholder_name(given_name):
        linked_name = metadata_relative_name(source_workbook, data.get("relation"))
    if linked_name is not None:
        given_name = linked_name.given_name
        family_name = linked_name.family_name

    related_id = get_or_create_person(
        {
            "given_name": given_name,
            "family_name": family_name,
            "birth_date": data.get("birth_date"),
            "death_date": data.get("death_date"),
        }
    )
    if not related_id:
        return None

    rel_type_id = get_or_create_relationship_type(data.get("relation"))
    if not rel_type_id:
        return None

    if linked_name:
        _merge_unnamed_relatives(
            main_person_id, rel_type_id, related_id, given_name, family_name
        )

    # Check local cache first
    cache_key = (main_person_id, related_id, rel_type_id)
    if cache_key in _relationship_cache:
        return _relationship_cache[cache_key]

    # Check for existing relationship via API
    existing = api_get("relationships")
    for r in existing:
        if (
            r["person_from"] == main_person_id
            and r["person_to"] == related_id
            and r["relationship_type"] == rel_type_id
        ):
            _relationship_cache[cache_key] = r["id"]
            return r["id"]

    created = api_post(
        "relationships",
        {
            "relationship_type": rel_type_id,
            "person_from": main_person_id,
            "person_to": related_id,
            "description": data.get("notes") or "",
        },
    )
    _relationship_cache[cache_key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# Import orchestration
# ---------------------------------------------------------------------------


def _maybe_update_person_place(person_id, field_name, location_id, source_name):
    person = api_get(f"persons/{person_id}")
    existing_id = person.get(field_name) if isinstance(person, dict) else None
    if existing_id:
        if str(existing_id) == str(location_id):
            return
        existing_location = api_get(f"locations/{existing_id}")
        existing_name = (
            existing_location.get("current_name") or existing_location.get("name")
            if isinstance(existing_location, dict)
            else None
        )
        from person_biography import place_names_equivalent

        if existing_name and place_names_equivalent(source_name, existing_name):
            return
        raise ValueError(
            f"Refusing to overwrite Person {person_id} {field_name} "
            f"{existing_name!r} with conflicting metadata value {source_name!r}"
        )
    api_patch("persons", person_id, {field_name: location_id})


def import_record(record, source_workbook=None):
    """Import a full person record (parsed JSON) via API calls.

    source_workbook is the metadata file name; it scopes the reviewed relative links.
    """
    person_id = get_or_create_person(
        record["person"],
        identifier=record.get("identifier"),
    )
    if not person_id:
        print("Skipped: invalid person")
        return None

    # Birth place
    birth_place = record.get("birth_place", {})
    location_id = get_or_create_location(birth_place)
    if location_id:
        _maybe_update_person_place(
            person_id,
            "birth_place",
            location_id,
            birth_place.get("name"),
        )

    # Interviews
    for interview_data in record.get("interviews", []):
        create_interview(person_id, interview_data)

    # Family / relationships
    person_data = SESSION.get(f"{BASE_URL}/persons/{person_id}/").json()
    main_family_name = person_data.get("family_name", "")
    for family_data in record.get("family", []):
        create_relationship(person_id, family_data, main_family_name, source_workbook)

    return person_id


def import_from_xlsx(xlsx_path):
    """Parse an xlsx file and import the person record via the API."""
    record = read_person_record(xlsx_path)
    person_id = import_record(record, source_workbook=Path(xlsx_path).name)
    if person_id:
        print(f"\u2705 Imported: {xlsx_path}")
    else:
        print(f"\u26a0\ufe0f Skipped: {xlsx_path}")
    return person_id


def import_all_metadati(directory=None):
    """Find and import all metadati_*.xlsx files from the data directory."""
    if directory is None:
        directory = Path(__file__).resolve().parent.parent / "data" / "nuove schede"
    directory = Path(directory)
    results = []
    for xlsx_path in sorted(directory.rglob("metadati_*.xlsx")):
        if "template" in xlsx_path.name:
            continue
        print(f"\nProcessing: {xlsx_path}")
        try:
            person_id = import_from_xlsx(xlsx_path)
            if person_id:
                results.append(person_id)
        except Exception as e:
            print(f"\u274c ERROR in {xlsx_path.name}: {e}")
    print(f"\n\u2705 Imported {len(results)} valid records")
    return results
