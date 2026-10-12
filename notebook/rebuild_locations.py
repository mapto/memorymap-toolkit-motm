"""Rebuild the location workbook from chronotopi and geospatial sources.

The rebuild is deliberately separate from the API import path. It produces a
stable workbook that can be consumed by ``locations.load_locations_db`` and
keeps the source provenance in ``coordinate_source`` while retaining every
existing workbook column. Geography is derived from the imported GO taxonomy;
addresses and descriptive source rows use their city context as the nearest
parent.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
import time
from typing import Callable

import openpyxl
import pandas as pd
import requests
from unidecode import unidecode

from locations import (
    extract_geonames_coordinates,
    extract_geonames_id,
    extract_wikidata_coordinates,
    normalize_location,
    query_geocoder,
    taxonomy_alias,
    taxonomy_match_excluded,
    tokenize_location,
)
from taxonomy import Taxonomy, _extract_key, _label_key, load_taxonomy
from utils import extract_id_from_url, extract_url_map


DEFAULT_INPUT = Path(__file__).resolve().parent / "locations.xlsx"
DEFAULT_OUTPUT = DEFAULT_INPUT
DEFAULT_DATA_ROOT = Path(__file__).resolve().parents[1] / "data"
DEFAULT_TAXONOMY = DEFAULT_DATA_ROOT / "maxqda" / "MAXQDA_Code_System.mmd"

BASE_COLUMNS = (
    "location",
    "bag_of_words",
    "super_region",
    "lat",
    "long",
    "label",
    "www.geonames.org",
    "www.wikidata.org",
)

DERIVED_COLUMNS = ("wider_region", "coordinate_source")

NON_LOCATION_NAMES = {
    "?",
    "aus Gießen nach Palästina",
    "Association of Israelis of Central European Origin (AICEO)",
    "Digitales Archiv vom Else-Lasker-Schüler-Gesellschaft",
}

QUERY_OVERRIDES = {
    "Kibbuz": "Ben Shemen, Israel",
    "Esther-Stern-Platz Gießen": "Gießen, Germany",
    "Steinstraße 19,": "Steinstraße 19, Gießen, Germany",
    "Marktplatz 11": "Marktplatz 11, Gießen, Germany",
    "Löberstraße 20": "Löberstraße 20, Gießen, Germany",
    "Walltorstraße 48": "Walltorstraße 48, Gießen, Germany",
    "Gießen Schillerschule (Volksschule)": "Schillerschule, Gießen, Germany",
    "Haifa, Universität, Bibliothek": "Haifa, Israel",
    "Kibbuz Rodges, besteht nicht mehr (in der Nähe von Petah Tikva)": "Petah Tikva, Israel",
}

@dataclass
class SourceLocation:
    """Merged source evidence for one normalized location name."""

    name: str
    direct_coordinates: list[tuple[float, float]] = field(default_factory=list)
    geonames_ids: set[str] = field(default_factory=set)
    wikidata_ids: set[str] = field(default_factory=set)
    details: set[str] = field(default_factory=set)
    categories: set[str] = field(default_factory=set)
    urls: dict[str, str] = field(default_factory=dict)


def clean_value(value) -> str:
    value = str(value or "").strip()
    return "" if value.lower() in {"nan", "none"} else value


def latinize_text(value: str) -> str:
    """Transliterate a region label so the exported value uses Latin text."""
    return " ".join(unidecode(clean_value(value)).split())


def location_key(value: str) -> str:
    """Normalize punctuation and uncertain suffixes for source merging."""
    cleaned = clean_value(value).rstrip("?").strip()
    cleaned = taxonomy_alias(cleaned)
    return normalize_location(cleaned)


def coordinate_pair(latitude, longitude) -> tuple[float, float] | None:
    try:
        latitude = float(latitude)
        longitude = float(longitude)
    except (TypeError, ValueError):
        return None
    if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
        return None
    return latitude, longitude


def read_existing_workbook(path: Path) -> tuple[list[str], dict[str, dict]]:
    """Read meaningful existing rows while retaining all workbook columns."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [clean_value(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]
    records = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = {header: value for header, value in zip(headers, values)}
        name = clean_value(record.get("location"))
        if not name or name.startswith("http") or name in NON_LOCATION_NAMES:
            continue
        records[location_key(name)] = record
    return headers, records


def _header_row(path: Path, sheet_name: str) -> int | None:
    preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=12)
    for index, row in preview.iterrows():
        values = {clean_value(value) for value in row if pd.notna(value)}
        if "event_label" in values:
            return int(index)
    return None


def _add_url(urls: dict[str, str], value: str) -> None:
    for host, url in extract_url_map(value)[0].items():
        if host and url:
            urls[host] = url


def read_chronotopi_sources(data_root: Path) -> dict[str, SourceLocation]:
    """Collect coordinates, identifiers, context, and URLs from all chronotopi workbooks."""
    sources: dict[str, SourceLocation] = {}
    for path in sorted(data_root.rglob("*chronotop*.xlsx")):
        try:
            workbook = pd.ExcelFile(path)
        except Exception:
            continue
        for sheet_name in workbook.sheet_names:
            if sheet_name.startswith("_"):
                continue
            header = _header_row(path, sheet_name)
            if header is None:
                continue
            frame = pd.read_excel(path, sheet_name=sheet_name, header=header)
            if "place_name" not in frame.columns:
                continue
            for _, row in frame.iterrows():
                name = clean_value(row.get("place_name"))
                if not name or name in NON_LOCATION_NAMES:
                    continue
                key = location_key(name)
                source = sources.setdefault(key, SourceLocation(name=name.rstrip("?").strip()))
                coordinates = coordinate_pair(
                    clean_value(row.get("place_lat")),
                    clean_value(row.get("place_lon")),
                )
                if coordinates:
                    source.direct_coordinates.append(coordinates)
                geonames_id = extract_geonames_id(clean_value(row.get("geonames_id")))
                if geonames_id:
                    source.geonames_ids.add(geonames_id)
                wikidata_id = clean_value(row.get("wikidata_qid")) or None
                if wikidata_id:
                    source.wikidata_ids.add(wikidata_id)
                detail = clean_value(row.get("place_detail"))
                if detail:
                    source.details.add(detail.rstrip("?").strip())
                category = clean_value(row.get("place_category"))
                if category:
                    source.categories.update(
                        part.strip() for part in category.splitlines() if part.strip()
                    )
                _add_url(source.urls, clean_value(row.get("external_links")))
    return sources


def read_metadata_sources(data_root: Path) -> dict[str, SourceLocation]:
    """Collect birth-place evidence from all metadata workbooks."""
    from metadati import read_person_record

    sources: dict[str, SourceLocation] = {}
    for path in sorted(data_root.rglob("metadati*.xlsx")):
        try:
            record = read_person_record(path)
        except Exception:
            continue
        birth_place = record.get("birth_place") or {}
        name = clean_value(birth_place.get("name"))
        if not name or name in NON_LOCATION_NAMES:
            continue
        key = location_key(name)
        source = sources.setdefault(key, SourceLocation(name=name))
        coordinates = birth_place.get("coordinates") or {}
        pair = coordinate_pair(coordinates.get("lat"), coordinates.get("lon"))
        if pair:
            source.direct_coordinates.append(pair)
        wikidata_id = clean_value(birth_place.get("wikidata_id"))
        if wikidata_id:
            source.wikidata_ids.add(wikidata_id)
        source.details.update(
            clean_value(region)
            for region in birth_place.get("regions", [])
            if clean_value(region)
        )
    return sources


def merge_source_maps(
    target: dict[str, SourceLocation], extra: dict[str, SourceLocation]
) -> None:
    """Merge source evidence from another import source into the location map."""
    for key, source in extra.items():
        if key not in target:
            target[key] = source
            continue
        current = target[key]
        current.direct_coordinates.extend(source.direct_coordinates)
        current.geonames_ids.update(source.geonames_ids)
        current.wikidata_ids.update(source.wikidata_ids)
        current.details.update(source.details)
        current.categories.update(source.categories)
        current.urls.update(source.urls)


def merge_sources(
    headers: list[str],
    existing: dict[str, dict],
    sources: dict[str, SourceLocation],
) -> dict[str, dict]:
    """Union existing workbook metadata with chronotopi source evidence."""
    merged = {key: dict(record) for key, record in existing.items()}
    for key, source in sources.items():
        record = merged.setdefault(key, {header: "" for header in headers})
        record.setdefault("location", source.name)
        if not clean_value(record.get("location")):
            record["location"] = source.name
        if source.geonames_ids and not clean_value(record.get("www.geonames.org")):
            record["www.geonames.org"] = (
                f"https://www.geonames.org/{sorted(source.geonames_ids)[0]}"
            )
        if source.wikidata_ids and not clean_value(record.get("www.wikidata.org")):
            record["www.wikidata.org"] = (
                f"https://www.wikidata.org/wiki/{sorted(source.wikidata_ids)[0]}"
            )
        for host, url in source.urls.items():
            if host not in record or not clean_value(record.get(host)):
                record[host] = url
        if source.categories:
            current_label = clean_value(record.get("label"))
            if not current_label or current_label.startswith("http"):
                record["label"] = ",".join(sorted(source.categories))
    return merged


def extend_headers(headers: list[str], sources: dict[str, SourceLocation]) -> list[str]:
    """Retain existing columns and append new URL hosts found in chronotopi."""
    extended = list(headers)
    for source in sources.values():
        for host in source.urls:
            if host not in extended:
                extended.append(host)
    return extended


def order_output_headers(headers: list[str]) -> list[str]:
    """Place derived geography columns directly after ``super_region``."""
    ordered = [header for header in headers if header not in DERIVED_COLUMNS]
    insert_at = ordered.index("super_region") + 1
    return ordered[:insert_at] + list(DERIVED_COLUMNS) + ordered[insert_at:]


def build_geo_paths(taxonomy: Taxonomy) -> dict[str, tuple[str, ...]]:
    """Index GO taxonomy labels by their broad-to-specific source path."""
    paths = {}
    for occurrence in taxonomy.occurrences:
        if _label_key(_extract_key(occurrence.top_level)) != "go":
            continue
        paths[_label_key(occurrence.label)] = occurrence.path
    return paths


def choose_direct_coordinates(source: SourceLocation) -> tuple[float, float] | None:
    if not source.direct_coordinates:
        return None
    rounded = Counter((round(lat, 6), round(lon, 6)) for lat, lon in source.direct_coordinates)
    return rounded.most_common(1)[0][0]


def source_ids(record: dict, source: SourceLocation | None) -> tuple[str | None, str | None]:
    geonames_url = clean_value(record.get("www.geonames.org"))
    wikidata_url = clean_value(record.get("www.wikidata.org"))
    geonames_id = extract_geonames_id(geonames_url)
    wikidata_id = extract_id_from_url(wikidata_url)
    if source:
        geonames_id = geonames_id or next(iter(sorted(source.geonames_ids)), None)
        wikidata_id = wikidata_id or next(iter(sorted(source.wikidata_ids)), None)
    return geonames_id, wikidata_id


def build_query(
    name: str,
    source: SourceLocation | None,
    record: dict,
    geo_paths: dict[str, tuple[str, ...]] | None = None,
) -> str:
    if name in QUERY_OVERRIDES:
        return QUERY_OVERRIDES[name]
    if geo_paths:
        normalized_name = _label_key(name)
        if normalized_name in geo_paths:
            return name
    context = next(iter(sorted(source.details)), "") if source and source.details else ""
    if context and len(context) < 80 and not context.startswith("http"):
        return f"{name}, {context}, Germany"
    existing_parent = clean_value(record.get("super_region"))
    if existing_parent and not existing_parent.startswith("GO –"):
        return f"{name}, {existing_parent}, Germany"
    return name


def geocode_nominatim(query: str, *, request_get: Callable = requests.get, delay: float = 1.1):
    """Geocode a final fallback query and return coordinates plus address metadata."""
    response = request_get(
        "https://nominatim.openstreetmap.org/search",
        params={"format": "jsonv2", "addressdetails": 1, "limit": 1, "q": query},
        headers={"User-Agent": "memorymap-toolkit/1.0"},
        timeout=20,
    )
    response.raise_for_status()
    results = response.json()
    if not results:
        return None
    time.sleep(delay)
    result = results[0]
    coordinates = coordinate_pair(result.get("lat"), result.get("lon"))
    if not coordinates:
        return None
    return {"coordinates": coordinates, "address": result.get("address", {})}


def reverse_geocode_nominatim(
    coordinates: tuple[float, float],
    *,
    request_get: Callable = requests.get,
    delay: float = 1.1,
):
    """Reverse-geocode coordinates to recover city/country containment."""
    latitude, longitude = coordinates
    response = request_get(
        "https://nominatim.openstreetmap.org/reverse",
        params={
            "format": "jsonv2",
            "addressdetails": 1,
            "lat": latitude,
            "lon": longitude,
        },
        headers={"User-Agent": "memorymap-toolkit/1.0"},
        timeout=20,
    )
    response.raise_for_status()
    time.sleep(delay)
    return response.json().get("address", {})


def resolve_coordinates(
    name: str,
    record: dict,
    source: SourceLocation | None,
    *,
    own_geocoder: Callable = query_geocoder,
    nominatim: Callable = geocode_nominatim,
    geo_paths: dict[str, tuple[str, ...]] | None = None,
) -> tuple[tuple[float, float], str, dict]:
    """Resolve coordinates in source-priority order and return provenance."""
    direct = choose_direct_coordinates(source) if source else None
    if direct:
        return direct, "chronotopi", {}

    query = build_query(name, source, record, geo_paths)
    if name in QUERY_OVERRIDES:
        own = own_geocoder(query)
        if own:
            pair = coordinate_pair(own.get("lat"), own.get("long"))
            if pair:
                return pair, "own-geocoder", {}
        result = nominatim(query)
        if result:
            return result["coordinates"], "nominatim", result.get("address", {})

    geonames_id, wikidata_id = source_ids(record, source)
    if geonames_id:
        url = f"https://www.geonames.org/{geonames_id}"
        try:
            coordinates = extract_geonames_coordinates(url)
        except Exception:
            coordinates = None
        if coordinates:
            pair = coordinate_pair(coordinates.get("lat"), coordinates.get("long"))
            if pair:
                return pair, "geonames", {}
    if wikidata_id:
        url = f"https://www.wikidata.org/wiki/{wikidata_id}"
        try:
            coordinates = extract_wikidata_coordinates(url)
        except Exception:
            coordinates = None
        if coordinates:
            pair = coordinate_pair(coordinates.get("lat"), coordinates.get("long"))
            if pair:
                return pair, "wikidata", {}

    own = own_geocoder(query)
    if own:
        pair = coordinate_pair(own.get("lat"), own.get("long"))
        if pair:
            return pair, "own-geocoder", {}

    result = nominatim(query)
    if result:
        return result["coordinates"], "nominatim", result.get("address", {})
    raise RuntimeError(f"No coordinates found for {name!r} using {query!r}")


def _taxonomy_parent_chain(name: str, geo_paths: dict[str, tuple[str, ...]]):
    if taxonomy_match_excluded(name):
        return []
    canonical = taxonomy_alias(name)
    path = _geo_path_for(canonical, geo_paths)
    if not path:
        return []
    if _label_key(canonical) == _label_key(name):
        return list(path[:-1])
    return list(path)


def _geo_path_for(label: str, geo_paths: dict[str, tuple[str, ...]]):
    """Find a GO path using exact or Latin-transliterated label keys."""
    path = geo_paths.get(_label_key(label))
    if path:
        return path
    latin_key = latinize_text(label).casefold()
    return next(
        (
            path
            for candidate, path in geo_paths.items()
            if latinize_text(candidate).casefold() == latin_key
        ),
        None,
    )


def wider_region_for(
    name: str,
    record: dict,
    source: SourceLocation | None,
    geo_paths: dict[str, tuple[str, ...]],
    address: dict,
) -> str:
    """Return broad-to-nearest containment for a location row."""
    exact = _taxonomy_parent_chain(name, geo_paths)
    if exact:
        return " > ".join(exact)

    normalized_name = _label_key(name)
    embedded_candidates = [
        (label, path)
        for label, path in geo_paths.items()
        if label and label in normalized_name
    ]
    embedded = max(embedded_candidates, key=lambda item: len(item[0]), default=None)
    if embedded:
        return " > ".join(embedded[1])

    context_candidates = []
    if source:
        context_candidates.extend(source.details)
    existing_parent = clean_value(record.get("super_region"))
    if existing_parent:
        context_candidates.append(existing_parent)
    for candidate in context_candidates:
        alias = taxonomy_alias(candidate)
        path = _geo_path_for(alias, geo_paths)
        if path:
            return " > ".join(path)

    address_names = [
        address.get("city"),
        address.get("town"),
        address.get("village"),
        address.get("municipality"),
        address.get("state"),
        address.get("country"),
    ]
    for candidate in address_names:
        if not candidate:
            continue
        alias = taxonomy_alias(candidate)
        path = _geo_path_for(alias, geo_paths)
        if path:
            return " > ".join(path)
    values = [value for value in address_names if value]
    if values:
        return " > ".join(["GO – geographischer Ort", *reversed(values)])
    return ""


def rebuild_workbook(
    input_path: Path = DEFAULT_INPUT,
    output_path: Path = DEFAULT_OUTPUT,
    data_root: Path = DEFAULT_DATA_ROOT,
    taxonomy_path: Path = DEFAULT_TAXONOMY,
    *,
    own_geocoder: Callable = query_geocoder,
    nominatim: Callable = geocode_nominatim,
    reverse_nominatim: Callable = reverse_geocode_nominatim,
) -> dict[str, int]:
    """Rebuild the workbook and assert every retained row has coordinates."""
    headers, existing = read_existing_workbook(input_path)
    sources = read_chronotopi_sources(data_root)
    merge_source_maps(sources, read_metadata_sources(data_root))
    headers = extend_headers(headers, sources)
    merged = merge_sources(headers, existing, sources)
    taxonomy = load_taxonomy(taxonomy_path)
    geo_paths = build_geo_paths(taxonomy)
    headers = order_output_headers(headers)

    output_rows = []
    dropped = 0
    source_counts = Counter()
    for key in sorted(merged, key=lambda item: clean_value(merged[item].get("location")).casefold()):
        record = merged[key]
        name = clean_value(record.get("location"))
        if not name or name in NON_LOCATION_NAMES or name == "?":
            dropped += 1
            continue
        source = sources.get(key)
        coordinates, coordinate_source, address = resolve_coordinates(
            name,
            record,
            source,
            own_geocoder=own_geocoder,
            nominatim=nominatim,
            geo_paths=geo_paths,
        )
        record["lat"], record["long"] = coordinates
        record["bag_of_words"] = " ".join(sorted(tokenize_location(name)))
        if not wider_region_for(name, record, source, geo_paths, address):
            try:
                address = reverse_nominatim(coordinates)
            except Exception:
                address = {}
        wider = wider_region_for(name, record, source, geo_paths, address)
        record["wider_region"] = wider
        record["super_region"] = latinize_text(wider.split(" > ")[-1]) if wider else ""
        record["coordinate_source"] = coordinate_source
        for header in headers:
            record.setdefault(header, "")
        output_rows.append(record)
        source_counts[coordinate_source] += 1

    if not output_rows or any(not coordinate_pair(row.get("lat"), row.get("long")) for row in output_rows):
        raise RuntimeError("Rebuilt workbook still contains rows without valid coordinates")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "locations"
    sheet.append(headers)
    for record in output_rows:
        sheet.append([record.get(header, "") for header in headers])
    workbook.save(output_path)
    return {
        "rows": len(output_rows),
        "dropped": dropped,
        **{f"coordinates_{key}": value for key, value in sorted(source_counts.items())},
    }


if __name__ == "__main__":
    print(rebuild_workbook())
