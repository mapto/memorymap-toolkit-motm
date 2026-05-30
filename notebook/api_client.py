"""API client for Memory Map Toolkit Django REST API.

Provides session management, CRUD helpers, and domain-specific
get_or_create_* functions for concepts, locations, persons, timespans, and URLs.
"""

from __future__ import annotations

import os
import re
import time
from typing import Any

import requests
from requests.exceptions import ConnectionError as ReqConnectionError


# ---------------------------------------------------------------------------
# Session & config
# ---------------------------------------------------------------------------

BASE_URL = os.environ.get("MMT_API_URL", "http://localhost:8000/motm/api")
LOGIN_URL = os.environ.get("MMT_LOGIN_URL", "http://localhost:8000/admin/login/")

SESSION = requests.Session()

_MAX_RETRIES = 5
_RETRY_DELAY = 3  # seconds


def _retry(fn, *args, **kwargs):  # type: ignore[no-untyped-def]
    """Retry *fn* on connection errors (e.g. Django auto-reload)."""
    for attempt in range(1, _MAX_RETRIES + 1):
        try:
            return fn(*args, **kwargs)
        except (
            ReqConnectionError,
            requests.exceptions.ChunkedEncodingError,
            ConnectionResetError,
        ) as exc:
            if attempt == _MAX_RETRIES:
                raise
            print(
                f"  \u21bb connection lost (attempt {attempt}/{_MAX_RETRIES}), "
                f"retrying in {_RETRY_DELAY}s \u2026 ({exc.__class__.__name__})"
            )
            time.sleep(_RETRY_DELAY)
    raise RuntimeError("unreachable")  # keeps pyright happy


def login(username="admin", password="admin"):
    """Authenticate with the Django admin and configure CSRF headers."""
    SESSION.get(LOGIN_URL)
    SESSION.post(
        LOGIN_URL,
        data={
            "username": username,
            "password": password,
            "csrfmiddlewaretoken": SESSION.cookies["csrftoken"],
            "next": "/admin/",
        },
    )
    SESSION.headers["X-CSRFToken"] = SESSION.cookies.get("csrftoken") or ""
    SESSION.headers["Referer"] = LOGIN_URL
    ok = bool(SESSION.cookies.get("sessionid"))
    print("Authenticated" if ok else "Login failed")
    return ok


# ---------------------------------------------------------------------------
# Generic REST helpers
# ---------------------------------------------------------------------------


def api_get(endpoint, params=None):
    if params is None:
        params = {}
    params.setdefault("limit", 10000)
    resp = _retry(SESSION.get, f"{BASE_URL}/{endpoint}/", params=params)
    resp.raise_for_status()
    data = resp.json()
    return data["results"] if isinstance(data, dict) and "results" in data else data


def api_post(endpoint, payload):
    resp = _retry(SESSION.post, f"{BASE_URL}/{endpoint}/", json=payload)
    resp.raise_for_status()
    return resp.json()


def api_patch(endpoint, obj_id, payload):
    resp = _retry(SESSION.patch, f"{BASE_URL}/{endpoint}/{obj_id}/", json=payload)
    resp.raise_for_status()
    return resp.json()


# ---------------------------------------------------------------------------
# String utilities
# ---------------------------------------------------------------------------


def clean_str(val):
    """Return cleaned string or empty string for nan/blank values."""
    s = str(val).strip()
    if s in ("", "nan"):
        return ""
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


# Alias for backward compatibility
clean = clean_str


def is_empty(val) -> bool:
    """Check if a value is None, empty, nan, or a dash placeholder."""
    return val is None or str(val).strip() in ("", "nan", "-", "\u2013", "\u2014")


def extract_urls_from_text(text):
    """Extract URLs from a text field."""
    if not text or str(text).strip() in ("", "nan"):
        return []
    return re.findall(r"https?://[^\s<>\"{}|\\^`\[\]]+", str(text))


# ---------------------------------------------------------------------------
# Concepts
# ---------------------------------------------------------------------------

_concept_cache = {}


def get_or_create_concept(label):
    """Get or create a concept.  If *label* contains ' > ' it is treated as a
    hierarchy, e.g. "Gegenstand > Grammophon" creates parent "Gegenstand" and
    child "Grammophon" with the parent relationship set."""
    label = label.strip()
    if not label or label == "nan":
        return None

    # Handle hierarchical labels like "Parent > Child"
    if " > " in label:
        parts = [p.strip() for p in label.split(" > ", 1)]
        parent_id = get_or_create_concept(parts[0])
        child_label = parts[1]
        child_key = child_label.lower()
        if child_key in _concept_cache:
            return _concept_cache[child_key]
        existing = api_get("concepts", params={"search": child_label})
        match = next((c for c in existing if c["label"].lower() == child_key), None)
        if match:
            _concept_cache[child_key] = match["id"]
            if parent_id and match.get("parent") != parent_id:
                api_patch("concepts", match["id"], {"parent": parent_id})
            return match["id"]
        payload = {"label": child_label}
        if parent_id:
            payload["parent"] = parent_id
        created = api_post("concepts", payload)
        _concept_cache[child_key] = created["id"]
        return created["id"]

    key = label.lower()
    if key in _concept_cache:
        return _concept_cache[key]
    existing = api_get("concepts", params={"search": label})
    match = next((c for c in existing if c["label"].lower() == key), None)
    if match:
        _concept_cache[key] = match["id"]
        return match["id"]
    created = api_post("concepts", {"label": label})
    _concept_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# Locations (API layer — uses locations_db for coordinates)
# ---------------------------------------------------------------------------

_location_cache = {}


def get_or_create_location(
    name, wikidata_qid=None, geonames_id=None, *, locations_db=None, normalize_fn=None
):
    """Create or retrieve a LocationPoint, using locations_db for coordinates."""
    if not name or str(name).strip() in ("", "nan"):
        return None
    name = str(name).strip()
    norm = normalize_fn or (lambda n: n.lower())
    key = norm(name)
    if key in _location_cache:
        return _location_cache[key]
    existing = api_get("locations", params={"search": name})
    match = next((loc for loc in existing if norm(loc["current_name"]) == key), None)
    if match:
        _location_cache[key] = match["id"]
        return match["id"]
    db_entry = (locations_db or {}).get(key, {})
    payload: dict[str, Any] = {"current_name": name}
    lat = db_entry.get("lat")
    lng = db_entry.get("long")
    if lat and lng:
        try:
            payload["latitude"] = float(lat)
            payload["longitude"] = float(lng)
        except (ValueError, TypeError):
            pass
    wid = db_entry.get("wikidata_id") or (
        wikidata_qid if wikidata_qid and wikidata_qid not in ("", "nan") else None
    )
    gid = db_entry.get("geonames_id") or (
        geonames_id if geonames_id and geonames_id not in ("", "nan") else None
    )
    if wid:
        payload["wikidata_id"] = wid
    if gid:
        payload["geonames_id"] = gid
    created = api_post("locations", payload)
    _location_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# Persons
# ---------------------------------------------------------------------------

_person_cache = {}


def find_person_by_identifier(identifier):
    """Find a Person by identifier (lookup only, does not create).

    Returns the person's API id or None.
    """
    if is_empty(identifier):
        return None
    existing = api_get("persons", params={"search": identifier})
    match = next((p for p in existing if p.get("identifier") == identifier), None)
    return match["id"] if match else None


def get_person_id(protagonist_col, name_col):
    """Get or create a Person from the protagonist + name columns."""
    raw_name = clean_str(name_col)
    if not raw_name:
        raw_name = clean_str(protagonist_col)
    if not raw_name:
        return None
    key = raw_name.lower()
    if key in _person_cache:
        return _person_cache[key]
    parts = raw_name.split()
    given = parts[0] if parts else raw_name
    family = " ".join(parts[1:]) if len(parts) > 1 else raw_name
    existing = api_get("persons", params={"search": raw_name})
    for p in existing:
        if (
            p["given_name"].lower() == given.lower()
            and p["family_name"].lower() == family.lower()
        ):
            _person_cache[key] = p["id"]
            return p["id"]
    created = api_post("persons", {"given_name": given, "family_name": family})
    _person_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# Interviews
# ---------------------------------------------------------------------------


def parse_interview_id_from_quelle(quelle):
    """Extract the interview ID from a quelle Markdown link.

    The quelle column format is:
        Interviewer Name – [IS_E_00124](https://dgd.ids-mannheim.de/...)
    Returns the link caption (e.g. 'IS_E_00124') or None.
    """
    if is_empty(quelle):
        return None
    m = re.search(r"\[([A-Z]+_[A-Z]_\d+)\]", quelle)
    return m.group(1) if m else None


def get_interview_id(archive_id, quelle=None):
    """Find an Interview by archive_id.

    Falls back to parsing quelle if archive_id is empty or not found.
    Returns the interview's API id or None.
    """
    if is_empty(archive_id) and quelle:
        archive_id = parse_interview_id_from_quelle(quelle)
    if is_empty(archive_id):
        return None
    existing = api_get("interviews", params={"search": archive_id})
    match = next((i for i in existing if i["archive_id"] == archive_id), None)
    if match:
        return match["id"]
    # archive_id from the explicit column didn't match; try quelle as fallback
    if quelle:
        fallback_id = parse_interview_id_from_quelle(quelle)
        if fallback_id and fallback_id != archive_id:
            existing = api_get("interviews", params={"search": fallback_id})
            match = next((i for i in existing if i["archive_id"] == fallback_id), None)
            if match:
                return match["id"]
    return None


# ---------------------------------------------------------------------------
# Timespans
# ---------------------------------------------------------------------------

_timespan_cache = {}


def get_or_create_timespan(start, end, certainty=None):
    """Get or create a Timespan object."""
    key = (str(start), str(end), certainty)
    if key in _timespan_cache:
        return _timespan_cache[key]
    payload = {}
    if start:
        payload["start"] = str(start)
    if end:
        payload["end"] = str(end)
    if certainty:
        payload["certainty"] = certainty
    if not payload:
        return None
    created = api_post("timespans", payload)
    _timespan_cache[key] = created["id"]
    return created["id"]


# ---------------------------------------------------------------------------
# URLs
# ---------------------------------------------------------------------------

_url_cache = {}


def get_or_create_url(url_str):
    url_str = url_str.strip()
    if not url_str or url_str == "nan":
        return None
    if url_str in _url_cache:
        return _url_cache[url_str]
    existing = api_get("urls")
    for u in existing:
        if u["url"] == url_str:
            _url_cache[url_str] = u["id"]
            return u["id"]
    try:
        created = api_post("urls", {"url": url_str})
        _url_cache[url_str] = created["id"]
        return created["id"]
    except Exception as e:
        print(f"  URL creation failed for {url_str}: {e}")
        return None


# ---------------------------------------------------------------------------
# Regions — link locations to super-regions via the API
# ---------------------------------------------------------------------------

_region_cache = {}  # region_name_lower -> region_id


def _get_or_create_region(name):
    """Get or create a LocationRegion."""
    key = name.strip().lower()
    if key in _region_cache:
        return _region_cache[key]
    existing = api_get("regions", params={"search": name})
    for r in existing:
        if r["name"].strip().lower() == key:
            _region_cache[key] = r["id"]
            return r["id"]
    created = api_post("regions", {"name": name.strip()})
    _region_cache[key] = created["id"]
    return created["id"]


def link_locations_to_regions(locations_db):
    """Create LocationRegions from super_region data and link LocationPoints."""
    import openpyxl as _xl
    import os

    xlsx_path = os.path.join(
        os.path.dirname(__file__) if "__file__" in dir() else ".",
        "locations.xlsx",
    )
    if not os.path.exists(xlsx_path):
        print("  locations.xlsx not found, skipping region linking")
        return

    wb = _xl.load_workbook(xlsx_path)
    ws = wb.active
    assert ws is not None, "locations.xlsx has no active sheet"
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    sr_col = (headers.index("super_region") + 1) if "super_region" in headers else None
    if not sr_col:
        print("  No super_region column in locations.xlsx, skipping")
        return

    # Collect location→super_region pairs
    pairs = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        sr = ws.cell(r, sr_col).value
        if name and sr and str(sr).strip() not in ("", "nan"):
            pairs.append((str(name).strip(), str(sr).strip()))

    # Create regions and link locations
    linked = 0
    for loc_name, region_name in pairs:
        # Find the location in the API
        existing_locs = api_get("locations", params={"search": loc_name})
        loc_match = next(
            (
                loc
                for loc in existing_locs
                if loc["current_name"].strip().lower() == loc_name.strip().lower()
            ),
            None,
        )
        if not loc_match:
            continue
        region_id = _get_or_create_region(region_name)
        if loc_match.get("region") != region_id:
            api_patch("locations", loc_match["id"], {"region": region_id})
            linked += 1

    print(
        f"  Linked {linked} locations to regions ({len(_region_cache)} regions created/found)"
    )
