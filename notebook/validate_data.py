"""Report data inconsistencies in the source folders without changing anything.

The checks reuse the import pipeline's own code, so the findings match what the
import would do: header verification, taxonomy augmentation, a dry-run taxonomy
import with fake IDs, the location and person resolvers, and the biography
planner. The chronotopi table is built the way cell 9 of ``chronotopi.ipynb``
builds it, except that the header row is found by searching for ``event_label``
instead of using the notebook's ``skip`` table.

The script is read-only and offline. It does not log in, call the API, use a
geocoder, or write any file. It only reports; fixes need a separate decision.

Usage, from the ``notebook`` folder::

    ./.venv/bin/python validate_data.py          # up to 15 findings per list
    ./.venv/bin/python validate_data.py --all    # every finding

The exit status is 1 when at least one ERROR is reported. WARNING and INFO
findings do not change it.
"""

from __future__ import annotations

import argparse
import difflib
import itertools
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

import openpyxl
import pandas as pd

from locations import (
    LOCATION_TAXONOMY_SYNONYMS_PATH,
    LOCATIONS_WORKBOOK_PATH,
    taxonomy_label_key,
    taxonomy_match_excluded,
)
from metadati import read_person_record
from person_biography import (
    PersonBiographyConflictError,
    PersonBiographyError,
    build_person_biography_plan,
)
from person_taxonomy import (
    PERSON_TAXONOMY_SYNONYMS_PATH,
    PersonaeAugmentation,
    PersonTaxonomyError,
    _is_collective_label,
    _is_personae_occurrence,
    _metadata_name_variants,
    _occurrences_for,
    augment_personae_taxonomy,
    is_placeholder_name,
    person_taxonomy_alias,
    person_taxonomy_key,
)
from rebuild_locations import (
    coordinate_pair,
    read_chronotopi_sources,
    read_existing_workbook,
    read_metadata_sources,
)
from taxonomy import (
    Taxonomy,
    TaxonomyError,
    _extract_key,
    _label_key,
    load_taxonomy,
    parse_mermaid_occurrences,
)
from taxonomy_import import TaxonomyImport, import_taxonomy
from timespan import parse_date
from utils import clean_str
from verify_excel_headers import verify_files

NOTEBOOK_DIR = Path(__file__).resolve().parent
REPO_ROOT = NOTEBOOK_DIR.parent
DATA_DIR = REPO_ROOT / "data"
SOURCE_DIR = DATA_DIR / "nuove schede"
TEMPLATE_DIR = DATA_DIR / "template"
MAXQDA_DIR = DATA_DIR / "maxqda"
MAXQDA_TAXONOMY = MAXQDA_DIR / "MAXQDA_Code_System.mmd"
DOCS_TAXONOMY = REPO_ROOT / "docs" / "mqda_coding_taxonomy.mmd"
LOCATION_WORKBOOK = LOCATIONS_WORKBOOK_PATH

ERROR = "ERROR"
WARNING = "WARNING"
INFO = "INFO"
SEVERITIES = (ERROR, WARNING, INFO)

FILES = "Files"
HEADERS = "Headers"
TAXONOMY = "Taxonomy"
SYNONYMS = "Synonym files"
PLACES = "Places"
PERSONS = "Persons"
SHEETS = "Chronotopi sheets"
VALUES = "Chronotopi values"
BIOGRAPHIES = "Biographies"
SECTIONS = (
    FILES,
    HEADERS,
    TAXONOMY,
    SYNONYMS,
    PLACES,
    PERSONS,
    SHEETS,
    VALUES,
    BIOGRAPHIES,
)

DEFAULT_LIMIT = 15
ROW_LIST_LIMIT = 8
CERTAINTY_VALUES = frozenset({"certain", "probable", "uncertain", "disputed"})
PRECISION_VALUES = frozenset({"day", "month", "year", "range", "after", "unknown"})
ENUM_COLUMNS: dict[str, frozenset[str]] = {
    "date_certainty": CERTAINTY_VALUES,
    "date_precision": PRECISION_VALUES,
}
CONCEPT_COLUMNS = ("event_type", "transport_mode", "place_category")
TIME_COLUMN = "date_label"
DATE_COLUMNS = ("date_start", "date_end")
LOCATION_SECTIONS = frozenset(
    {
        "aliases",
        "canonical_parents",
        "excluded_contains",
        "query_overrides",
        "reviewed_unmapped",
    }
)
PERSON_SECTIONS = frozenset({"aliases", "metadata_by_chronotopi", "metadata_relatives"})
WIKIDATA_ID = re.compile(r"(?:https?://www\.wikidata\.org/wiki/)?Q\d+")


@dataclass(frozen=True)
class Finding:
    """One problem or note, with its severity and the section it belongs to."""

    severity: str
    section: str
    message: str


class Report:
    """Collects findings and renders them grouped by section."""

    def __init__(self) -> None:
        self.findings: list[Finding] = []

    def add(self, severity: str, section: str, message: str) -> None:
        self.findings.append(Finding(severity, section, message))

    def count(self, severity: str) -> int:
        return sum(finding.severity == severity for finding in self.findings)

    def render(self, limit: int | None) -> str:
        lines: list[str] = []
        for section in SECTIONS:
            found = [finding for finding in self.findings if finding.section == section]
            lines.append(f"== {section}")
            if not found:
                lines.append("  no findings")
            for severity in SEVERITIES:
                messages = [
                    finding.message for finding in found if finding.severity == severity
                ]
                shown = messages if limit is None else messages[:limit]
                lines.extend(f"  {severity}: {message}" for message in shown)
                if len(shown) < len(messages):
                    lines.append(
                        f"  {severity}: ... {len(messages) - len(shown)} more (use --all)"
                    )
        lines.append(
            f"Summary: {self.count(ERROR)} errors, {self.count(WARNING)} warnings, "
            f"{self.count(INFO)} notes"
        )
        return "\n".join(lines)


@dataclass(frozen=True)
class SynonymPair:
    """One source-to-target entry of a synonym section."""

    source: str
    target: str


@dataclass(frozen=True)
class LocationSynonyms:
    """Parsed entries of data/location_taxonomy_synonyms.json."""

    aliases: list[SynonymPair]
    parents: list[SynonymPair]
    excluded: list[str]
    reviewed: list[str]

    def alias_map(self) -> dict[str, str]:
        """Alias targets keyed by normalized source label."""
        return {taxonomy_label_key(pair.source): pair.target for pair in self.aliases}

    def parent_map(self) -> dict[str, str]:
        """Canonical parents keyed by normalized child label."""
        return {taxonomy_label_key(pair.source): pair.target for pair in self.parents}


@dataclass(frozen=True)
class PersonSynonyms:
    """Parsed entries of data/person_taxonomy_synonyms.json."""

    aliases: list[SynonymPair]
    associations: list[SynonymPair]

    def alias_target_keys(self) -> set[str]:
        """Normalized keys of all alias targets: the names a relative may link to."""
        return {person_taxonomy_key(pair.target) for pair in self.aliases}

    def association_map(self) -> dict[str, str]:
        """Metadata workbook file name keyed by chronotopi workbook file name."""
        return {
            Path(pair.source).name: Path(pair.target).name for pair in self.associations
        }


@dataclass(frozen=True)
class LoadedSheet:
    """One chronotopi sheet, identified by its workbook file name."""

    workbook: str
    sheet: str

    @property
    def label(self) -> str:
        return f"{self.workbook} [{self.sheet}]"


@dataclass
class ChronotopiData:
    """Chronotopi tables built the way cells 5 (overview) and 9 (frame) build them."""

    overview: dict[str, list[str]]
    frame: pd.DataFrame
    loaded: list[LoadedSheet]


def cell(row: Mapping[Any, Any], column: str) -> str:
    """Return a cleaned cell; the notebook's astype(str) turns blanks into "nan"."""
    return clean_str(row.get(column, ""))


def has_outer_spaces(raw: object) -> bool:
    text = "" if raw is None else str(raw)
    return text not in ("", "nan", "None") and text != text.strip()


def location_of(row: Mapping[Any, Any]) -> str:
    return f"{row['source_workbook']} [{row['name']}] row {row['excel_row']}"


def format_rows(numbers: Sequence[int]) -> str:
    shown = ", ".join(str(number) for number in numbers[:ROW_LIST_LIMIT])
    noun = "row" if len(numbers) == 1 else "rows"
    more = " ..." if len(numbers) > ROW_LIST_LIMIT else ""
    return f"{noun} {shown}{more}"


def summarize(counts: Mapping[str, int], limit: int = 10) -> str:
    items = [f"{value!r} ({count})" for value, count in sorted(counts.items())[:limit]]
    hidden = len(counts) - limit
    return ", ".join(items) + (f", and {hidden} more" if hidden > 0 else "")


# --- Files and headers -------------------------------------------------------


def chronotopi_workbooks() -> list[Path]:
    """Mirror cell 2: top-level chronotopi files, without hidden or marker files."""
    return sorted(
        path
        for path in SOURCE_DIR.glob("*chronotop*")
        if not path.name.startswith(".") and not path.name.endswith(":Zone.Identifier")
    )


def check_files(report: Report) -> list[Path]:
    """Check the source folder for unreadable workbooks and non-.xlsx chronotopi files."""
    for path in chronotopi_workbooks():
        if path.suffix != ".xlsx":
            report.add(
                ERROR,
                FILES,
                f"{path.name}: chronotopi file is not an .xlsx workbook; the notebook cannot read it",
            )
    for path in sorted(SOURCE_DIR.glob("*.xlsx")):
        try:
            openpyxl.load_workbook(path, read_only=True).close()
        except Exception as exc:
            report.add(
                ERROR, FILES, f"{path.name}: cannot be opened as a workbook ({exc})"
            )
    return [path for path in chronotopi_workbooks() if path.suffix == ".xlsx"]


def check_headers(report: Report) -> None:
    """Compare every source workbook header with its template (verify_excel_headers)."""
    result = verify_files(SOURCE_DIR, TEMPLATE_DIR)
    if result.error:
        report.add(ERROR, HEADERS, f"header verification failed: {result.error}")
        return
    for file_result in result.files:
        for error in file_result.errors:
            details = []
            if error.missing:
                details.append(f"missing: {', '.join(error.missing)}")
            if error.extra:
                details.append(f"extra: {', '.join(error.extra)}")
            severity = ERROR if error.missing else WARNING
            suffix = f" ({'; '.join(details)})" if details else ""
            report.add(
                severity,
                HEADERS,
                f"{file_result.name} [{error.sheet}]: {error.message}{suffix}",
            )


def check_rebuild_scope(report: Report) -> None:
    """The location rebuild reads every match under data/, so flag matches outside the source folder."""
    for pattern in ("*chronotop*.xlsx", "metadati*.xlsx"):
        for path in sorted(DATA_DIR.rglob(pattern)):
            if path.parent != SOURCE_DIR:
                report.add(
                    WARNING,
                    PLACES,
                    f"{path.relative_to(REPO_ROOT)}: the location rebuild reads this workbook as a source, "
                    f"but it is outside {SOURCE_DIR.relative_to(REPO_ROOT)}",
                )


# --- Taxonomy ----------------------------------------------------------------


def taxonomy_label_set(path: Path) -> set[str]:
    return {
        _label_key(occurrence.label) for occurrence in parse_mermaid_occurrences(path)
    }


def check_taxonomy_copies(report: Report) -> None:
    """Compare the MAXQDA code-system copies; only MAXQDA_Code_System.mmd feeds the import."""
    reference = taxonomy_label_set(MAXQDA_TAXONOMY)
    for copy in [DOCS_TAXONOMY, *sorted(MAXQDA_DIR.glob("*.mmd"))]:
        if copy == MAXQDA_TAXONOMY or not copy.is_file():
            continue
        try:
            labels = taxonomy_label_set(copy)
        except Exception as exc:
            report.add(
                WARNING,
                TAXONOMY,
                f"{copy.relative_to(REPO_ROOT)} cannot be parsed ({exc})",
            )
            continue
        if labels != reference:
            report.add(
                WARNING,
                TAXONOMY,
                f"{copy.relative_to(REPO_ROOT)} differs from {MAXQDA_TAXONOMY.name}: "
                f"only in the copy {sorted(labels - reference)[:5]}, "
                f"only in the used file {sorted(reference - labels)[:5]}",
            )


def dry_run_import(
    report: Report, taxonomy: Taxonomy, person_aliases: Mapping[str, str]
) -> TaxonomyImport | None:
    """Run the real import with fake IDs, so resolvers can be queried without writes."""
    ids = itertools.count(1)

    def create_concept(label: str, **_options: object) -> int:
        return next(ids)

    try:
        importer = import_taxonomy(
            taxonomy, create_concept, person_aliases=dict(person_aliases)
        )
    except (TaxonomyError, PersonTaxonomyError) as exc:
        report.add(ERROR, TAXONOMY, f"dry-run import failed: {exc}")
        return None
    report.add(
        INFO,
        TAXONOMY,
        f"dry-run import resolved {len(importer.ids_by_key)} taxonomy keys with fake IDs",
    )
    return importer


# --- Synonym files -----------------------------------------------------------


def load_strict_json(path: Path, report: Report) -> dict[str, Any] | None:
    """Parse a synonym file and report duplicate keys, which json.loads would hide."""
    duplicates: list[str] = []

    def keep_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                duplicates.append(key)
            result[key] = value
        return result

    try:
        data = json.loads(
            path.read_text(encoding="utf-8"), object_pairs_hook=keep_pairs
        )
    except (OSError, json.JSONDecodeError) as exc:
        report.add(ERROR, SYNONYMS, f"{path.name} cannot be read as JSON: {exc}")
        return None
    for key in dict.fromkeys(duplicates):
        report.add(
            ERROR,
            SYNONYMS,
            f"{path.name}: key {key!r} is defined more than once; only the last value is used",
        )
    if not isinstance(data, dict):
        report.add(ERROR, SYNONYMS, f"{path.name}: the top level must be a JSON object")
        return None
    return data


def pairs_of(
    section: object, file_name: str, key: str, report: Report
) -> list[SynonymPair]:
    """Return an object section as SynonymPair entries with string values."""
    if not isinstance(section, dict):
        report.add(
            ERROR, SYNONYMS, f"{file_name}: section {key!r} must be a JSON object"
        )
        return []
    pairs: list[SynonymPair] = []
    for source, target in section.items():
        if isinstance(target, str):
            pairs.append(SynonymPair(str(source), target))
        else:
            report.add(
                ERROR,
                SYNONYMS,
                f"{file_name}: {key!r} value for {source!r} must be a string",
            )
    return pairs


def items_of(section: object, file_name: str, key: str, report: Report) -> list[str]:
    """Return an array section as strings."""
    if not isinstance(section, list) or not all(
        isinstance(item, str) for item in section
    ):
        report.add(
            ERROR,
            SYNONYMS,
            f"{file_name}: section {key!r} must be a JSON array of strings",
        )
        return []
    return list(section)


def text_problems(text: str, *, punctuation: bool) -> list[str]:
    problems: list[str] = []
    if text != text.strip():
        problems.append("leading or trailing spaces")
    if "  " in text:
        problems.append("double spaces")
    if punctuation and re.search(r"[,;:]$", text.strip()):
        problems.append("trailing punctuation")
    return problems


def check_text(
    report: Report, file_name: str, role: str, text: str, *, punctuation: bool
) -> None:
    for problem in text_problems(text, punctuation=punctuation):
        report.add(WARNING, SYNONYMS, f"{file_name}: {role} {text!r} has {problem}")


def report_collisions(
    report: Report,
    file_name: str,
    pairs: Sequence[SynonymPair],
    normalize: Callable[[str], str],
) -> None:
    """Flag sources that normalize to one key; the loaders keep only the last value."""
    groups: defaultdict[str, list[SynonymPair]] = defaultdict(list)
    for pair in pairs:
        groups[normalize(pair.source)].append(pair)
    for key, members in groups.items():
        if len(members) < 2:
            continue
        conflicting = len({normalize(pair.target) for pair in members}) > 1
        severity = ERROR if conflicting else WARNING
        kind = "conflicting" if conflicting else "identical"
        listing = "; ".join(f"{pair.source!r} -> {pair.target!r}" for pair in members)
        report.add(
            severity,
            SYNONYMS,
            f"{file_name}: sources normalize to {key!r} with {kind} targets: {listing}",
        )


def read_location_synonyms(report: Report) -> LocationSynonyms | None:
    name = LOCATION_TAXONOMY_SYNONYMS_PATH.name
    data = load_strict_json(LOCATION_TAXONOMY_SYNONYMS_PATH, report)
    if data is None:
        return None
    for section in sorted(set(data) - LOCATION_SECTIONS):
        report.add(WARNING, SYNONYMS, f"{name}: unknown section {section!r} is ignored")
    synonyms = LocationSynonyms(
        aliases=pairs_of(data.get("aliases", {}), name, "aliases", report),
        parents=pairs_of(
            data.get("canonical_parents", {}), name, "canonical_parents", report
        ),
        excluded=items_of(
            data.get("excluded_contains", []), name, "excluded_contains", report
        ),
        reviewed=items_of(
            data.get("reviewed_unmapped", []), name, "reviewed_unmapped", report
        ),
    )
    for pair in [*synonyms.aliases, *synonyms.parents]:
        check_text(report, name, "key", pair.source, punctuation=True)
        check_text(report, name, "target", pair.target, punctuation=True)
    for term in [*synonyms.excluded, *synonyms.reviewed]:
        check_text(report, name, "entry", term, punctuation=False)
    report_collisions(report, name, synonyms.aliases, taxonomy_label_key)
    report_collisions(report, name, synonyms.parents, taxonomy_label_key)
    report_collisions(
        report,
        name,
        [SynonymPair(term, term) for term in synonyms.excluded],
        taxonomy_label_key,
    )
    report_collisions(
        report,
        name,
        [SynonymPair(term, term) for term in synonyms.reviewed],
        taxonomy_label_key,
    )
    return synonyms


def check_metadata_relatives(
    report: Report, file_name: str, section: object, synonyms: PersonSynonyms
) -> None:
    """Relative links must point at a reviewed 'Surname, Given' alias target."""
    if not isinstance(section, dict):
        report.add(
            ERROR,
            SYNONYMS,
            f"{file_name}: section 'metadata_relatives' must be a JSON object",
        )
        return
    alias_targets = synonyms.alias_target_keys()
    for workbook, links in section.items():
        pairs = pairs_of(links, file_name, f"metadata_relatives {workbook}", report)
        for link in pairs:
            check_text(
                report, file_name, "relative link", link.source, punctuation=False
            )
            target_key = person_taxonomy_key(link.target)
            if "," not in link.target or target_key not in alias_targets:
                report.add(
                    ERROR,
                    SYNONYMS,
                    f"{file_name}: relative {link.source!r} of {workbook} "
                    f"must link to a reviewed alias target, not {link.target!r}",
                )


def read_person_synonyms(report: Report) -> PersonSynonyms | None:
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    data = load_strict_json(PERSON_TAXONOMY_SYNONYMS_PATH, report)
    if data is None:
        return None
    for section in sorted(set(data) - PERSON_SECTIONS):
        report.add(WARNING, SYNONYMS, f"{name}: unknown section {section!r} is ignored")
    synonyms = PersonSynonyms(
        aliases=pairs_of(data.get("aliases", {}), name, "aliases", report),
        associations=pairs_of(
            data.get("metadata_by_chronotopi", {}),
            name,
            "metadata_by_chronotopi",
            report,
        ),
    )
    for pair in synonyms.aliases:
        target = pair.target
        check_text(report, name, "alias key", pair.source, punctuation=True)
        check_text(report, name, "alias target", target, punctuation=True)
        if "," not in target:
            report.add(
                WARNING,
                SYNONYMS,
                f"{name}: alias target {target!r} is not in 'Surname, Given' form",
            )
            continue
        family, given = (part.strip() for part in target.split(",", 1))
        if is_placeholder_name(family) or is_placeholder_name(given):
            report.add(
                ERROR,
                SYNONYMS,
                f"{name}: alias target {target!r} has a placeholder name part",
            )
    for pair in synonyms.associations:
        check_text(report, name, "association key", pair.source, punctuation=False)
        check_text(report, name, "association value", pair.target, punctuation=False)
    check_metadata_relatives(report, name, data.get("metadata_relatives", {}), synonyms)
    report_collisions(report, name, synonyms.aliases, person_taxonomy_key)
    return synonyms


# --- Places ------------------------------------------------------------------


def parent_cycles(parents: Mapping[str, str]) -> list[list[str]]:
    """Return each distinct cycle in the normalized parent map once."""
    cycles: list[list[str]] = []
    seen: set[frozenset[str]] = set()
    for start in parents:
        path: list[str] = []
        node: str | None = start
        while node is not None and node in parents and node not in path:
            path.append(node)
            node = taxonomy_label_key(parents[node])
        if node is not None and node in path:
            cycle = path[path.index(node) :]
            signature = frozenset(cycle)
            if signature not in seen:
                seen.add(signature)
                cycles.append([*cycle, node])
    return cycles


def check_location_graph(
    report: Report, synonyms: LocationSynonyms, go_keys: set[str]
) -> None:
    """Check the normalized alias and parent maps for chains, cycles, and overlaps."""
    name = LOCATION_TAXONOMY_SYNONYMS_PATH.name
    aliases = synonyms.alias_map()
    parents = synonyms.parent_map()
    keys = set(aliases) | set(parents)
    for source, target in aliases.items():
        target_key = taxonomy_label_key(target)
        if target_key != source and target_key in aliases and target_key not in go_keys:
            report.add(
                WARNING,
                PLACES,
                f"{name}: alias {source!r} -> {target!r}: the target is itself an alias and not a "
                "GO label, and lookups stop after one step",
            )
    for cycle in parent_cycles(parents):
        report.add(
            ERROR, PLACES, f"{name}: canonical_parents cycle {' -> '.join(cycle)}"
        )
    for term in sorted({taxonomy_label_key(term) for term in synonyms.reviewed}):
        if term in keys:
            report.add(
                WARNING,
                PLACES,
                f"{name}: reviewed_unmapped {term!r} also has an alias or parent entry; "
                "the exclusion wins and that entry is never used",
            )
    for term in sorted({taxonomy_label_key(term) for term in synonyms.excluded}):
        hits = sorted(key for key in keys if term and term in key)
        if hits:
            report.add(
                WARNING,
                PLACES,
                f"{name}: excluded_contains {term!r} matches {hits[:5]}; "
                "those alias or parent entries are never used",
            )


def check_places(
    report: Report, importer: TaxonomyImport | None, synonyms: LocationSynonyms | None
) -> None:
    """Dry-run the location linking for every name the rebuild and sync would see."""
    if importer is None:
        report.add(
            INFO,
            PLACES,
            "place checks skipped: the dry-run taxonomy import did not complete",
        )
        return
    names = {source.name for source in read_chronotopi_sources(DATA_DIR).values()}
    names.update(source.name for source in read_metadata_sources(DATA_DIR).values())
    if LOCATION_WORKBOOK.is_file():
        _, records = read_existing_workbook(LOCATION_WORKBOOK)
        names.update(clean_str(record.get("location")) for record in records.values())
    names.discard("")
    excluded: list[str] = []
    supplemental: list[str] = []
    for name in sorted(names, key=str.casefold):
        if taxonomy_match_excluded(name):
            excluded.append(name)
            continue
        direct = importer.resolve_location_exact(name) is not None
        try:
            importer.ensure_location_concept(name)
        except TaxonomyError as exc:
            report.add(ERROR, PLACES, f"place {name!r} cannot be linked: {exc}")
            continue
        if not direct:
            supplemental.append(name)
    if synonyms is not None:
        for pair in synonyms.parents:
            try:
                importer.ensure_location_concept(pair.target)
            except TaxonomyError as exc:
                report.add(
                    ERROR,
                    PLACES,
                    f"parent hint {pair.target!r} cannot be linked: {exc}",
                )
    if excluded:
        report.add(
            INFO,
            PLACES,
            f"{len(excluded)} place names are reviewed as unmapped: "
            f"{', '.join(sorted(excluded, key=str.casefold))}",
        )
    if supplemental:
        report.add(
            INFO,
            PLACES,
            f"{len(supplemental)} place names have no direct GO label or alias and become "
            f"supplemental concepts: {', '.join(sorted(supplemental, key=str.casefold))}",
        )


# --- Persons -----------------------------------------------------------------


def check_person_associations(
    report: Report, workbooks: Sequence[Path], synonyms: PersonSynonyms
) -> None:
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    mapping = synonyms.association_map()
    chrono_names = {path.name for path in workbooks}
    metadata_files = {path.name for path in SOURCE_DIR.glob("metadati*.xlsx")}
    for chrono in sorted(chrono_names - mapping.keys()):
        report.add(
            ERROR,
            PERSONS,
            f"{chrono}: no metadata association in {name}; augmentation stops at this workbook",
        )
    for chrono in sorted(mapping.keys() - chrono_names):
        report.add(
            WARNING,
            PERSONS,
            f"{name}: association for {chrono} matches no chronotopi workbook in {SOURCE_DIR.name}",
        )
    for chrono, metadata in sorted(mapping.items()):
        if metadata not in metadata_files:
            report.add(
                ERROR,
                PERSONS,
                f"{name}: {chrono} points to {metadata}, which does not exist in {SOURCE_DIR.name}",
            )
    for metadata in sorted(metadata_files - set(mapping.values())):
        report.add(
            WARNING,
            PERSONS,
            f"{metadata}: not associated with any chronotopi workbook; its person is not placed under Personae",
        )


def raw_birth_date(path: Path) -> str:
    """Return the Geburtsdatum text as written in the 'Scheda individuale' sheet."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for row in workbook["Scheda individuale"].iter_rows(values_only=True):
            if row and clean_str(row[0]) == "Geburtsdatum":
                return clean_str(row[1]) if len(row) > 1 else ""
        return ""
    finally:
        workbook.close()


def read_metadata_records(report: Report) -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for path in sorted(SOURCE_DIR.glob("metadati*.xlsx")):
        try:
            records[path.name] = read_person_record(path)
        except Exception as exc:
            report.add(ERROR, PERSONS, f"{path.name}: metadata cannot be read ({exc})")
    return records


def check_metadata_records(
    report: Report, records: Mapping[str, Mapping[str, Any]]
) -> None:
    owners: defaultdict[str, list[str]] = defaultdict(list)
    for file_name, record in records.items():
        person = record.get("person") or {}
        identifier = clean_str(record.get("identifier"))
        if identifier:
            owners[identifier].append(file_name)
        else:
            report.add(
                WARNING,
                PERSONS,
                f"{file_name}: no metadata identifier (Metadaten ID) found",
            )
        given = person.get("given_name")
        family = person.get("family_name")
        if is_placeholder_name(given) or is_placeholder_name(family):
            report.add(
                WARNING,
                PERSONS,
                f"{file_name}: main name is missing or a placeholder "
                f"(given {clean_str(given)!r}, family {clean_str(family)!r})",
            )
        if person.get("birth_date_interval") is None:
            raw = raw_birth_date(SOURCE_DIR / file_name)
            if raw:
                report.add(
                    WARNING,
                    PERSONS,
                    f"{file_name}: birth date {raw!r} cannot be parsed; it is not compared with chronotopi",
                )
    for identifier, files in owners.items():
        if len(files) > 1:
            report.add(
                ERROR,
                PERSONS,
                f"identifier {identifier} is used by {', '.join(sorted(files))}",
            )


def augment_dry_run(
    report: Report, taxonomy: Taxonomy, overview: Mapping[str, Sequence[str]]
) -> PersonaeAugmentation | None:
    try:
        return augment_personae_taxonomy(taxonomy, overview, data_root=SOURCE_DIR)
    except (PersonTaxonomyError, OSError, KeyError, ValueError) as exc:
        report.add(ERROR, PERSONS, f"augmentation failed: {exc}")
        return None


def check_person_placement(
    report: Report,
    overview: Mapping[str, Sequence[str]],
    taxonomy: Taxonomy,
    synonyms: PersonSynonyms,
    records: Mapping[str, Mapping[str, Any]],
    person_aliases: Mapping[str, str],
) -> None:
    """Find aliases the augmentation silently drops and sheets that become raw leaves."""
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    for pair in synonyms.aliases:
        if person_taxonomy_key(pair.source) not in person_aliases:
            report.add(
                WARNING,
                PERSONS,
                f"{name}: alias {pair.source!r} -> {pair.target!r} is ignored; "
                "its target is not a Personae leaf and no sheet uses it",
            )
    occurrences = _occurrences_for(taxonomy)
    parent_indexes = {
        occurrence.parent_index
        for occurrence in occurrences
        if occurrence.parent_index is not None
    }
    known = {
        person_taxonomy_key(occurrence.label)
        for index, occurrence in enumerate(occurrences)
        if _is_personae_occurrence(occurrence)
        and index not in parent_indexes
        and not _is_collective_label(occurrence.label)
    }
    for record in records.values():
        for variant in _metadata_name_variants(record.get("person") or {}):
            known.add(person_taxonomy_key(variant))
    for pair in synonyms.aliases:
        known.add(person_taxonomy_key(pair.source))
        known.add(person_taxonomy_key(pair.target))
    for path_text, sheets in overview.items():
        workbook = Path(path_text).name
        for sheet in sheets:
            if sheet.startswith("_"):
                continue
            if (
                person_taxonomy_key(sheet) in known
                or person_taxonomy_alias(sheet) is not None
            ):
                continue
            report.add(
                WARNING,
                PERSONS,
                f"{workbook} [{sheet}]: no metadata record, MAXQDA person, or alias matches; "
                "the sheet becomes a raw Personae leaf",
            )


# --- Chronotopi sheets -------------------------------------------------------


def read_sheet(report: Report, path: Path, sheet: str) -> pd.DataFrame | None:
    """Load one sheet as cell 9 does, locating the header by its event_label cell."""
    where = f"{path.name} [{sheet}]"
    preview = pd.read_excel(path, sheet_name=sheet, header=None, nrows=10)
    preview_rows = preview.itertuples(index=False, name=None)
    header_row = next(
        (
            position
            for position, values in enumerate(preview_rows)
            if "event_label"
            in {str(value).strip() for value in values if pd.notna(value)}
        ),
        None,
    )
    if header_row is None:
        report.add(
            ERROR,
            SHEETS,
            f"{where}: no 'event_label' header in the first 10 rows; the notebook reads row 1 as the header",
        )
        header_row = 0
    frame = pd.read_excel(path, sheet_name=sheet, header=header_row)
    if "Unnamed:" in str(frame.columns[0]):
        report.add(
            WARNING,
            SHEETS,
            f"{where}: legacy layout without named columns; the notebook renames them and this check skips the sheet",
        )
        return None
    duplicated = [str(column) for column in frame.columns[frame.columns.duplicated()]]
    if duplicated:
        report.add(
            WARNING,
            SHEETS,
            f"{where}: duplicate column headers {duplicated}; only the first column of each name is read",
        )
    frame = frame.astype(str)
    frame["name"] = sheet
    frame["source_workbook"] = path.name
    frame["excel_row"] = [header_row + index + 2 for index in range(len(frame))]
    return frame


def load_chronotopi(report: Report, workbooks: Sequence[Path]) -> ChronotopiData:
    overview: dict[str, list[str]] = {}
    loaded: list[LoadedSheet] = []
    frames: list[pd.DataFrame] = []
    for path in workbooks:
        try:
            workbook = openpyxl.load_workbook(path, read_only=True)
        except Exception:
            continue  # already reported under Files
        sheet_names = list(workbook.sheetnames)
        workbook.close()
        overview[str(path)] = sheet_names
        for sheet in sheet_names:
            if sheet.startswith("_"):
                continue
            frame = read_sheet(report, path, sheet)
            if frame is not None:
                frames.append(frame)
                loaded.append(LoadedSheet(path.name, sheet))
    if frames:
        combined = pd.concat(frames, ignore_index=True).astype(str)
        add_locations(combined)
    else:
        combined = pd.DataFrame(
            columns=["name", "source_workbook", "excel_row", "place_name"]
        )
    return ChronotopiData(overview, combined, loaded)


def add_locations(frame: pd.DataFrame) -> None:
    """Mirror the end of cell 9: start_location is the previous place of the same sheet."""
    if "place_name" not in frame.columns:
        frame["place_name"] = ""
    previous: dict[str, str] = {}
    starts: list[str] = []
    for name, place in zip(frame["name"], frame["place_name"]):
        end = str(place).strip()
        starts.append(previous.get(str(name), end))
        if end:
            previous[str(name)] = end
    frame["start_location"] = starts
    frame["end_location"] = frame["place_name"]


def check_sheets(
    report: Report,
    chrono: ChronotopiData,
    workbook_count: int,
    importer: TaxonomyImport | None,
) -> None:
    report.add(
        INFO,
        SHEETS,
        f"{len(chrono.loaded)} sheets loaded from {workbook_count} chronotopi workbooks",
    )
    skipped = [
        f"{Path(path).name} [{sheet}]"
        for path, sheets in chrono.overview.items()
        for sheet in sheets
        if sheet.startswith("_")
    ]
    if skipped:
        report.add(
            INFO,
            SHEETS,
            f"{len(skipped)} sheets starting with '_' are skipped by design: {', '.join(skipped)}",
        )
    workbooks_by_sheet: defaultdict[str, set[str]] = defaultdict(set)
    for item in chrono.loaded:
        workbooks_by_sheet[item.sheet].add(item.workbook)
        if item.sheet != item.sheet.strip() or "  " in item.sheet:
            report.add(
                WARNING,
                SHEETS,
                f"{item.workbook} [{item.sheet!r}]: sheet name has stray spaces",
            )
    for sheet, workbooks in sorted(workbooks_by_sheet.items()):
        if len(workbooks) > 1:
            report.add(
                WARNING,
                SHEETS,
                f"sheet {sheet!r} appears in {', '.join(sorted(workbooks))}; "
                "events from both are attached to one person",
            )
    if importer is None:
        report.add(
            INFO,
            SHEETS,
            "person checks skipped: the dry-run taxonomy import did not complete",
        )
        return
    for item in chrono.loaded:
        try:
            person = importer.resolve_person(item.sheet)
        except TaxonomyError as exc:
            report.add(ERROR, SHEETS, f"{item.label}: person name is ambiguous ({exc})")
            continue
        if person is None:
            report.add(
                WARNING,
                SHEETS,
                f"{item.label}: the sheet name does not resolve to a Personae concept; "
                "its events get no person",
            )


def check_rows(report: Report, rows: Sequence[Mapping[Any, Any]]) -> None:
    """Find rows the import skips and events that repeat an earlier row."""
    missing_place: defaultdict[str, list[int]] = defaultdict(list)
    missing_event: defaultdict[str, list[int]] = defaultdict(list)
    first_seen: dict[tuple[str, ...], str] = {}
    for row in rows:
        event = cell(row, "event_label")
        place = cell(row, "place_name")
        sheet_key = f"{row['source_workbook']} [{row['name']}]"
        if event and not place:
            missing_place[sheet_key].append(int(row["excel_row"]))
        elif place and not event:
            missing_event[sheet_key].append(int(row["excel_row"]))
        elif event and place:
            key = (
                str(row["name"]),
                event,
                place,
                cell(row, TIME_COLUMN),
                cell(row, "date_start"),
                cell(row, "date_end"),
            )
            if key in first_seen:
                report.add(
                    WARNING,
                    SHEETS,
                    f"{location_of(row)} repeats {first_seen[key]} (same person, event, place, and dates)",
                )
            else:
                first_seen[key] = location_of(row)
    for sheet_key, numbers in missing_place.items():
        report.add(
            WARNING,
            SHEETS,
            f"{sheet_key}: {format_rows(numbers)} have an event but no place_name; the import skips them",
        )
    for sheet_key, numbers in missing_event.items():
        report.add(
            WARNING,
            SHEETS,
            f"{sheet_key}: {format_rows(numbers)} have a place but no event_label; the import skips them",
        )


# --- Chronotopi values -------------------------------------------------------


def closest_label(value: str, labels: Mapping[str, str]) -> str | None:
    """Suggest the taxonomy label closest to a value that matches nothing."""
    matches = difflib.get_close_matches(
        _label_key(value), list(labels), n=1, cutoff=0.85
    )
    return labels[matches[0]] if matches else None


def check_column(
    report: Report,
    column: str,
    values: Counter[str],
    *,
    allowed: frozenset[str] | None,
    resolve: Callable[[str], int | None] | None,
    labels: Mapping[str, str],
) -> None:
    """Check one column's distinct values: variants, multi-line cells, vocabulary, taxonomy."""
    grouped: defaultdict[str, list[str]] = defaultdict(list)
    for value in values:
        grouped[_label_key(value)].append(value)
    for forms in grouped.values():
        if len(forms) > 1:
            listing = ", ".join(f"{form!r} ({values[form]})" for form in sorted(forms))
            report.add(
                WARNING, VALUES, f"{column}: spelling variants of one value: {listing}"
            )
    unknown: list[str] = []
    for value in sorted(values):
        count = values[value]
        if "\n" in value:
            report.add(
                WARNING,
                VALUES,
                f"{column}: multi-line value {value!r} ({count} rows) is not split and matches nothing",
            )
            continue
        if value.startswith(("http://", "https://")):
            report.add(
                WARNING,
                VALUES,
                f"{column}: URL {value!r} ({count} rows) in a value column",
            )
            continue
        if allowed is not None:
            if value not in allowed:
                report.add(
                    WARNING,
                    VALUES,
                    f"{column}: {value!r} ({count} rows) is outside the vocabulary {sorted(allowed)}",
                )
            continue
        if resolve is None:
            continue
        try:
            found = resolve(value)
        except TaxonomyError as exc:
            report.add(
                ERROR,
                VALUES,
                f"{column}: {value!r} is ambiguous in the taxonomy ({exc})",
            )
            continue
        if found is not None:
            continue
        suggestion = closest_label(value, labels)
        if suggestion:
            report.add(
                WARNING,
                VALUES,
                f"{column}: {value!r} ({count} rows) is not a taxonomy label; closest is {suggestion!r}",
            )
        else:
            unknown.append(f"{value!r} ({count})")
    if unknown:
        report.add(
            INFO,
            VALUES,
            f"{column}: {len(unknown)} values have no taxonomy concept and are dropped: {', '.join(unknown)}",
        )


def check_values(
    report: Report,
    rows: Sequence[Mapping[Any, Any]],
    columns: Sequence[str],
    importer: TaxonomyImport | None,
    labels: Mapping[str, str],
) -> None:
    for column in (*CONCEPT_COLUMNS, TIME_COLUMN, *ENUM_COLUMNS):
        if column not in columns:
            continue
        values = Counter(
            value for value in (cell(row, column) for row in rows) if value
        )
        spaced = sum(1 for row in rows if has_outer_spaces(row.get(column)))
        if spaced:
            report.add(
                INFO,
                VALUES,
                f"{column}: {spaced} cells have leading or trailing spaces",
            )
        allowed = ENUM_COLUMNS.get(column)
        resolve: Callable[[str], int | None] | None = None
        if allowed is None and importer is not None:
            resolve = (
                importer.resolve_time_period
                if column == TIME_COLUMN
                else importer.resolve_concept
            )
        check_column(
            report, column, values, allowed=allowed, resolve=resolve, labels=labels
        )


def check_fields(report: Report, rows: Sequence[Mapping[Any, Any]]) -> None:
    """Check dates, coordinates, and identifiers, reporting unparseable values in aggregate."""
    unparsed: dict[str, Counter[str]] = {column: Counter() for column in DATE_COLUMNS}
    qids: Counter[str] = Counter()
    geonames: Counter[str] = Counter()
    for row in rows:
        for column in DATE_COLUMNS:
            value = cell(row, column)
            if value and parse_date(value) is None:
                unparsed[column][value] += 1
        start_text = cell(row, "date_start")
        end_text = cell(row, "date_end")
        start = parse_date(start_text) if start_text else None
        end = parse_date(end_text) if end_text else None
        if start and end and start > end:
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: date_start {start} is after date_end {end}",
            )
        lat = cell(row, "place_lat")
        lon = cell(row, "place_lon")
        if bool(lat) != bool(lon):
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: only one of place_lat and place_lon is set",
            )
        elif lat and coordinate_pair(lat, lon) is None:
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: coordinates {lat!r}, {lon!r} are invalid and are ignored",
            )
        qid = cell(row, "wikidata_qid")
        if qid and not WIKIDATA_ID.fullmatch(qid):
            qids[qid] += 1
        geonames_id = cell(row, "geonames_id")
        if geonames_id and not geonames_id.isdigit():
            geonames[geonames_id] += 1
    for column, counter in unparsed.items():
        if counter:
            report.add(
                WARNING,
                VALUES,
                f"{column}: {sum(counter.values())} values cannot be parsed as dates, "
                f"so those events lose the timespan: {summarize(counter)}",
            )
    if qids:
        report.add(
            WARNING,
            VALUES,
            f"wikidata_qid: values that are not Q-identifiers: {summarize(qids)}",
        )
    if geonames:
        report.add(
            WARNING,
            VALUES,
            f"geonames_id: values that are not plain numeric IDs: {summarize(geonames)}",
        )


# --- Biographies -------------------------------------------------------------


def check_biographies(report: Report, chrono: ChronotopiData) -> None:
    try:
        plan = build_person_biography_plan(chrono.overview, chrono.frame, SOURCE_DIR)
    except PersonBiographyConflictError as exc:
        for conflict in exc.conflicts:
            report.add(ERROR, BIOGRAPHIES, f"{conflict} (reported only; not fixed)")
        return
    except PersonBiographyError as exc:
        report.add(ERROR, BIOGRAPHIES, str(exc))
        return
    report.add(INFO, BIOGRAPHIES, f"{len(plan)} people have validated biography facts")


# --- Entry point -------------------------------------------------------------


def finish(report: Report, show_all: bool) -> int:
    print(report.render(None if show_all else DEFAULT_LIMIT))
    return 1 if report.count(ERROR) else 0


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Report data inconsistencies without changing anything."
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="print every finding instead of capping each list",
    )
    args = parser.parse_args(argv)

    report = Report()
    workbooks = check_files(report)
    check_headers(report)
    check_rebuild_scope(report)
    check_taxonomy_copies(report)

    try:
        taxonomy = load_taxonomy(MAXQDA_TAXONOMY)
        concept_nodes = taxonomy.concept_nodes()
    except (OSError, TaxonomyError, ValueError) as exc:
        report.add(ERROR, TAXONOMY, f"{MAXQDA_TAXONOMY.name} cannot be loaded: {exc}")
        return finish(report, args.all)
    go_keys = {
        taxonomy_label_key(node.label)
        for node in concept_nodes
        if _label_key(_extract_key(node.top_level)) == _label_key("GO")
    }
    labels = {_label_key(node.label): node.label for node in concept_nodes}

    location = read_location_synonyms(report)
    person = read_person_synonyms(report)
    if location is not None:
        check_location_graph(report, location, go_keys)
    if person is not None:
        check_person_associations(report, workbooks, person)

    chrono = load_chronotopi(report, workbooks)
    rows = chrono.frame.to_dict("records")
    records = read_metadata_records(report)
    check_metadata_records(report, records)

    importer: TaxonomyImport | None = None
    augmentation = (
        augment_dry_run(report, taxonomy, chrono.overview)
        if person is not None
        else None
    )
    if augmentation is not None and person is not None:
        importer = dry_run_import(report, augmentation.taxonomy, augmentation.aliases)
        check_person_placement(
            report, chrono.overview, taxonomy, person, records, augmentation.aliases
        )
    check_places(report, importer, location)
    check_sheets(report, chrono, len(workbooks), importer)
    check_rows(report, rows)
    check_fields(report, rows)
    check_values(report, rows, list(chrono.frame.columns), importer, labels)
    check_biographies(report, chrono)
    return finish(report, args.all)


if __name__ == "__main__":
    sys.exit(main())
