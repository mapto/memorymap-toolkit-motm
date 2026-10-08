"""Unit tests for Belege workbook imports."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from api_client import normalize_interview_date
import belege


DATA_DIR = Path(__file__).parent.parent / "data" / "nuove schede"
CURRENT_WORKBOOKS = tuple(
    path
    for path in sorted(DATA_DIR.glob("belege_*.xlsx"))
    if "template" not in path.name.lower()
)


def write_workbook(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Belege"
    sheet.append(["Belege"])
    sheet.append([])
    sheet.append(
        [
            "beleg_id",
            "quelle",
            "interview_id",
            "interview_datum",
            "quelle_sprecher",
            "betrifft_personen",
            "timecode",
            "themen",
            "zitat",
            "markierung",
            "event_ids",
            "event_ids_confidence",
            "notizen",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


@pytest.fixture
def api_stub(monkeypatch):
    calls = {"upserts": [], "interview_upserts": [], "concepts": []}

    def fake_api_get(endpoint, params=None):
        if endpoint == "persons":
            return [
                {"identifier": "P-MENTIONED", "id": 11},
                {"identifier": "P-SPEAKER", "id": 12},
            ]
        if endpoint == "interviews":
            return [{"archive_id": "I-001", "id": 21}]
        return []

    def fake_concept(label):
        calls["concepts"].append(label)
        return f"concept:{label}"

    def fake_upsert_interviews(items):
        calls["interview_upserts"].append(("interviews", items))
        return {"I-001": 21}, []

    def fake_bulk_upsert(endpoint, items):
        calls["upserts"].append((endpoint, items))
        first_id = 1000 + sum(
            len(batch) for _, batch in calls["upserts"][:-1]
        )
        return {
            "created_items": [
                {"id": first_id + index, **item}
                for index, item in enumerate(items)
            ],
            "updated_items": [],
            "errors": [],
        }

    monkeypatch.setattr(belege, "api_get", fake_api_get)
    monkeypatch.setattr(belege, "get_or_create_concept", fake_concept)
    monkeypatch.setattr(belege, "upsert_interviews", fake_upsert_interviews)
    monkeypatch.setattr(belege, "api_bulk_upsert", fake_bulk_upsert)
    monkeypatch.setattr(belege, "detect_language", lambda text: [])
    monkeypatch.setattr(belege, "tqdm", lambda rows, **kwargs: rows)
    return calls


def test_import_belege_builds_interview_and_extraction_payloads(tmp_path, api_stub):
    workbook_path = tmp_path / "belege_test.xlsx"
    write_workbook(
        workbook_path,
        [
            [
                "bel_001",
                "source",
                "I-001",
                "01.01.2000",
                "P-SPEAKER",
                "P-MENTIONED, P-OTHER",
                "00:12",
                "Topic A, Topic B",
                "A sufficiently long quote for testing language handling.",
                "highlight",
                "event-1",
                "certain",
                "A note",
            ]
        ],
    )

    result = belege.import_belege(workbook_path)

    assert result == [1000]
    assert api_stub["interview_upserts"] == [
        (
            "interviews",
            [{"archive_id": "I-001", "interviewee": 12, "date": "2000-01-01"}],
        )
    ]
    endpoint, payloads = api_stub["upserts"][0]
    assert endpoint == "extractions"
    assert payloads == [
        {
            "identifier": "bel_001",
            "timecode": "00:12",
            "quote": "A sufficiently long quote for testing language handling.",
            "classification": "highlight",
            "event_confidence": "certain",
            "notes": "A note",
            "people_mentioned": 11,
            "interview": 21,
            "concepts": ["concept:Topic A", "concept:Topic B"],
        }
    ]


def test_import_belege_creates_and_links_interview_from_quelle(tmp_path, monkeypatch):
    workbook_path = tmp_path / "belege_test.xlsx"
    write_workbook(
        workbook_path,
        [
            [
                "bel_001",
                "Anne Betten – [IS_E_170](https://example.test/interview)",
                None,
                None,
                "P-SPEAKER",
                "P-MENTIONED",
            ]
        ],
    )
    calls = {"upserts": [], "interview_upserts": []}

    def fake_api_get(endpoint, params=None):
        if endpoint == "persons":
            return [
                {"identifier": "P-MENTIONED", "id": 11},
                {"identifier": "P-SPEAKER", "id": 12},
            ]
        if endpoint == "interviews":
            return []
        return []

    def fake_bulk_upsert(endpoint, items):
        calls["upserts"].append((endpoint, items))
        return {
            "created_items": [{"id": 1000, **items[0]}],
            "updated_items": [],
            "errors": [],
        }

    monkeypatch.setattr(belege, "api_get", fake_api_get)
    monkeypatch.setattr(
        belege,
        "upsert_interviews",
        lambda items: ({"IS_E_00170": 21},
                       calls["interview_upserts"].append(("interviews", items)) or []),
    )
    monkeypatch.setattr(belege, "api_bulk_upsert", fake_bulk_upsert)
    monkeypatch.setattr(belege, "get_or_create_concept", lambda label: None)
    monkeypatch.setattr(belege, "detect_language", lambda text: [])
    monkeypatch.setattr(belege, "tqdm", lambda rows, **kwargs: rows)

    assert belege.import_belege(workbook_path) == [1000]
    assert calls["interview_upserts"] == [
        ("interviews", [{"archive_id": "IS_E_00170", "interviewee": 12}])
    ]
    assert calls["upserts"][0][1][0]["interview"] == 21


def test_parse_interview_id_from_quelle_normalizes_archive_number():
    assert belege.parse_interview_id_from_quelle(
        "Anne Betten – [IS_E_170](https://example.test/interview)"
    ) == "IS_E_00170"
    assert belege.parse_interview_id_from_quelle(
        "Anne Betten – [IS_E_00124](https://example.test/interview)"
    ) == "IS_E_00124"


def test_parse_interview_date_preserves_complete_dates():
    assert normalize_interview_date("08.05.1991") == "1991-05-08"
    assert normalize_interview_date("2010") is None


def test_import_belege_returns_created_and_updated_ids(tmp_path, monkeypatch):
    workbook_path = tmp_path / "belege_test.xlsx"
    write_workbook(workbook_path, [["bel_001"]])
    monkeypatch.setattr(belege, "api_get", lambda endpoint, params=None: [])
    monkeypatch.setattr(belege, "get_or_create_concept", lambda label: None)
    monkeypatch.setattr(belege, "tqdm", lambda rows, **kwargs: rows)
    monkeypatch.setattr(
        belege,
        "api_bulk_upsert",
        lambda endpoint, items: {
            "created_items": [{"id": 101}],
            "updated_items": [{"id": 202}],
            "errors": [],
        },
    )

    assert belege.import_belege(workbook_path) == [101, 202]


def data_rows(path):
    workbook = load_workbook(path, read_only=True)
    try:
        all_rows = list(workbook["Belege"].iter_rows(values_only=True))
        return [
            row for row in all_rows[3:] if row and any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()


def test_current_workbooks_are_discovered():
    assert CURRENT_WORKBOOKS


@pytest.mark.parametrize("workbook_path", CURRENT_WORKBOOKS)
def test_all_rows_from_current_workbook_are_imported(workbook_path, api_stub):
    rows = data_rows(workbook_path)
    result = belege.import_belege(workbook_path)

    assert len(rows) > 0
    assert len(result) == len(rows)
    assert len(api_stub["upserts"]) == 1
    _, payloads = api_stub["upserts"][0]
    assert len(payloads) == len(rows)
    assert [payload["identifier"] for payload in payloads] == [
        belege.clean_str(row[0]) for row in rows
    ]
