"""Workbook sheet parsers for metadati person records.

A metadati workbook has a main sheet, named "Scheda individuale" or, when that
name is missing, the active sheet. It lists fields as key-value rows: the German
label in column A, the value in column B, and an optional extra value in column
C. The "Familie" sheet is a table of relatives whose rows start at row 4 and span
seven columns. The parsers here turn both sheets into plain dicts, so the
workbook layout is handled in one place and the API code never sees spreadsheet
cells.

Field labels are compared as written, and source labels such as "Quelle 2 – Typ"
are read by pattern. Unknown labels are ignored. ``read_person_record`` is the
public entry point. Other packages call it to read workbooks without the API,
for example ``validate_data``. This module only reads files.
"""

from __future__ import annotations

import re
from openpyxl import load_workbook
from timespan import parse_date, parse_date_range
from metadati.values import (
    clean_value,
    clean_metadata_value,
    _date_interval,
    _split_biographical_date_range,
    parse_coordinates,
    parse_source_label,
)


# ---------------------------------------------------------------------------
# XLSX reader
# ---------------------------------------------------------------------------


def _read_sheet_rows(sheet):
    """Flatten a sheet into ``(label, value, extra)`` tuples, one per row.

    The label is column A as stripped text. ``value`` and ``extra`` are the raw
    contents of columns B and C, or None when the row is shorter. Values are not
    cleaned here, because the field parsers call ``clean_value``. Blank rows
    arrive with an empty label and match no field, so they are ignored.
    """
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        field = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        value = row[1] if len(row) > 1 else None
        extra = row[2] if len(row) > 2 else None
        rows.append((field, value, extra))
    return rows


def _parse_identifier(field, value_str):
    """Return the metadati identifier carried by a row, or None.

    The "ID Sprecher" row gives the identifier as its value. Other rows are
    searched for "Metadaten - CODE" in the label or the value, with the en dash
    read as a hyphen. ``_parse_person_sheet`` calls this for rows until one yields
    an identifier, and that identifier is later used to match the Person.
    """
    if field == "ID Sprecher":
        return value_str

    combined = " ".join([x for x in [field, value_str] if x]).replace("\u2013", "-")
    match = re.search(r"Metadaten\s*-\s*([A-Za-z0-9_]+)", combined)
    return match.group(1) if match else None


def _update_birth_place(birth_place, value_str, extra_str):
    """Fill the birth-place dict in place from the "Geburtsort" row.

    The name comes from an "Ortsname:" entry when present, and otherwise from the
    first line of the value. The value and extra cells are then split on bars and
    line breaks into "key: detail" entries. "Land", "Region", and "Kreis" append to
    ``regions``, a key containing "koordinat" is parsed by ``parse_coordinates``,
    and "Wikidata" sets ``wikidata_id``. Returns None.

    Raises AttributeError when the value is None and the extra cell has no
    "Ortsname:" entry, because the first line of the value is then read.
    """
    metadata = " | ".join(value for value in (value_str, extra_str) if value)
    name_match = re.search(
        r"(?:^|\||\n)\s*Ortsname\s*:\s*([^|\r\n]+)",
        metadata,
        flags=re.IGNORECASE,
    )
    birth_place["name"] = (
        name_match.group(1).strip() if name_match else value_str.splitlines()[0].strip()
    )

    for part in (p.strip() for p in re.split(r"\||\r?\n", metadata) if p.strip()):
        key, separator, detail = part.partition(":")
        if not separator:
            continue
        normalized_key = key.strip().casefold()
        detail = detail.strip()
        if normalized_key in {"land", "region", "kreis"}:
            birth_place["regions"].append(detail)
        elif "koordinat" in normalized_key:
            birth_place["coordinates"] = parse_coordinates(detail)
        elif normalized_key in {"wikidata", "wikidata id"}:
            birth_place["wikidata_id"] = detail


def _parse_person_field(person, birth_place, field, value_str, extra_str):
    """Store one key-value row in the person dict or the birth-place dict.

    Known labels map to person fields. For example, "Vorname" sets
    ``given_name``, and "Geburtsdatum" sets ``birth_date`` and its
    ``birth_date_interval``. "Verfolgtengruppe NS" goes to ``attributes``, and
    "Geburtsort" is passed to ``_update_birth_place``. Rows with other labels are
    ignored. ``_parse_person_sheet`` calls this once per row. Mutates both dicts
    and returns None.
    """
    if field == "Vorname":
        person["given_name"] = value_str
    elif field in {"ehem. Vorname", "ehemaliger Vorname"}:
        person["previous_given_name"] = value_str
    elif field == "Nachname":
        person["family_name"] = value_str
    elif field in {"ehem. Nachname", "ehemaliger Nachname"}:
        person["previous_family_name"] = value_str
    elif field == "Geschlecht":
        person["gender"] = value_str
    elif field == "Geburtsdatum":
        person["birth_date"] = parse_date(value_str) if value_str else None
        person["birth_date_interval"] = _date_interval(value_str)
    elif field in {"Verfolgtengruppe NS", "Verfolgtengruppe(n) NS"}:
        person["attributes"]["ns_persecution_group"] = value_str
    elif field == "Schulbildung / Abschluss":
        person["education"] = clean_metadata_value(value_str)
    elif field == "Beruf (zum Zeitpunkt der Emigration)":
        person["occupation_at_emigration"] = clean_metadata_value(value_str)
    elif field == "Beruf (zum Zeitpunkt der Quelle)":
        person["occupation_at_source"] = clean_metadata_value(value_str)
    elif field == "Wohnort zum Zeitpunkt der Quelle":
        person["residence_at_source"] = clean_metadata_value(value_str)
    elif field == "Geburtsort":
        _update_birth_place(birth_place, value_str, extra_str)


def _new_interview():
    """Return an empty interview dict with the keys the source parsers fill.

    The keys are ``type``, ``place``, ``date``, ``interviewer``, ``archive_id``,
    and ``url``, all None. ``_parse_source_field`` pads the list with these entries,
    so that a "Quelle 3" row can be stored before the rows for sources 1 and 2.
    This keeps each source number equal to its list position. Entries that stay
    empty are dropped by ``_parse_person_sheet``.
    """
    return {
        "type": None,
        "place": None,
        "date": None,
        "interviewer": None,
        "archive_id": None,
        "url": None,
    }


def _parse_source_field(field, value, value_str, interviews):
    """Store one "Quelle N - subfield" row in the interview list.

    Number N selects the interview at position N - 1 and pads the list when
    needed. The subfield sets the key: "Typ" sets ``type``, "Ort" sets ``place``,
    "Datum" sets ``date``, "Gesprächspartner/in" sets ``interviewer``,
    "Bezeichnung" sets ``interviewer`` and ``archive_id`` through
    ``parse_source_label``, "Handle/URL" sets ``url``, and "ID" sets
    ``archive_id``. Labels that do not match the pattern, and unknown subfields,
    are ignored. A number of 0 is not rejected: it selects the last interview, or
    raises IndexError when the list is empty. Mutates ``interviews`` and returns
    None.
    """
    match = re.match(r"Quelle\s+(\d+)\s*[-\u2013\u2014]\s*(.+)", field)
    if not match:
        return

    index = int(match.group(1))
    subfield = match.group(2)
    while len(interviews) < index:
        interviews.append(_new_interview())
    current = interviews[index - 1]

    if subfield == "Typ":
        current["type"] = clean_value(value)
    elif subfield == "Ort":
        current["place"] = clean_value(value)
    elif subfield == "Datum":
        current["date"] = parse_date(value_str) if value_str else None
    elif subfield == "Gesprächspartner/in":
        current["interviewer"] = clean_value(value)
    elif subfield == "Bezeichnung":
        interviewer, archive_id = parse_source_label(value)
        current["interviewer"] = interviewer
        current["archive_id"] = archive_id
    elif subfield == "Handle/URL":
        current["url"] = clean_value(value)
    elif subfield == "ID":
        current["archive_id"] = clean_value(value)


def _parse_person_sheet(rows):
    """Build the person, birth place, interviews, and identifier from sheet rows.

    Walks the rows once, in order, and uses the first row that yields an
    identifier. A "QUELLEN" row starts the sources section, and "MIGRATION" ends
    it. "Quelle" rows are read inside the section and anywhere else they appear.
    Interviews with no content are dropped at the end. Returns the tuple
    ``(identifier, person, birth_place, interviews)``. The person dict has a key
    for each person field the sheet can set, with None where the sheet is silent.
    """
    person = {
        "given_name": None,
        "previous_given_name": None,
        "family_name": None,
        "previous_family_name": None,
        "gender": None,
        "birth_date": None,
        "birth_date_interval": None,
        "education": None,
        "occupation_at_emigration": None,
        "occupation_at_source": None,
        "residence_at_source": None,
        "attributes": {},
    }
    birth_place = {
        "name": None,
        "coordinates": None,
        "wikidata_id": None,
        "regions": [],
    }
    interviews = []
    identifier = None
    in_sources_section = False

    for field, value, extra in rows:
        value_str = clean_value(value)
        extra_str = clean_value(extra)

        if identifier is None:
            identifier = _parse_identifier(field, value_str)

        if field == "QUELLEN":
            in_sources_section = True
            continue
        if field == "MIGRATION":
            in_sources_section = False

        _parse_person_field(person, birth_place, field, value_str, extra_str)

        if in_sources_section or field.startswith("Quelle "):
            _parse_source_field(field, value, value_str, interviews)

    interviews = [
        interview
        for interview in interviews
        if any(
            [
                interview["type"],
                interview["place"],
                interview["date"],
                interview["interviewer"],
                interview["archive_id"],
                interview["url"],
            ]
        )
    ]
    return identifier, person, birth_place, interviews


def _parse_family_sheet(sheet):
    """Read the relatives table of the "Familie" sheet into a list of dicts.

    The first three rows are headers and are ignored. Rows with every cell empty
    are skipped. Each dict has ``relation``, ``given_name``, ``family_name``,
    ``birth_place``, ``notes``, ``birth_date`` and ``death_date`` as ISO strings,
    and the ``*_date_interval`` pairs. A cell with two dates is split into birth
    and death, and a single date is taken as the birth date. Returns the list,
    which may be empty.
    """
    family = []
    for row in sheet.iter_rows(min_row=4, max_col=7, values_only=True):
        if all(cell is None for cell in row):
            continue
        (
            _,
            relation,
            given_name,
            family_name,
            birth_date_raw,
            birth_place,
            notes,
        ) = row
        birth_date_interval, death_date_interval = _split_biographical_date_range(
            birth_date_raw
        )
        birth_date, death_date = (
            parse_date_range(str(birth_date_raw)) if birth_date_raw else (None, None)
        )
        family.append(
            {
                "relation": clean_value(relation),
                "given_name": clean_value(given_name),
                "family_name": clean_value(family_name),
                "birth_date": birth_date,
                "birth_date_interval": birth_date_interval,
                "death_date": death_date,
                "death_date_interval": death_date_interval,
                "birth_place": clean_value(birth_place),
                "notes": clean_value(notes),
            }
        )
    return family


def read_person_record(xlsx_path, sheet_name="Scheda individuale"):
    """Parse one metadati workbook into the record dict used by the importer.

    Reads the sheet named ``sheet_name``, or the active sheet when that name is
    missing, and the "Familie" sheet when it exists. The workbook is opened with
    openpyxl and never saved. Returns a dict with the keys ``identifier``,
    ``person``, ``birth_place``, ``interviews``, and ``family``. An assertion
    checks that a sheet was found. Called by ``importer.import_from_xlsx`` and by
    other packages that need the parsed data without the API.
    """
    wb = load_workbook(xlsx_path)
    sheet = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    assert sheet is not None, f"No sheet found in {xlsx_path}"

    identifier, person, birth_place, interviews = _parse_person_sheet(
        _read_sheet_rows(sheet)
    )
    family = _parse_family_sheet(wb["Familie"]) if "Familie" in wb.sheetnames else []

    return {
        "identifier": identifier,
        "person": person,
        "birth_place": birth_place,
        "interviews": interviews,
        "family": family,
    }
