"""Report data inconsistencies in the source folders without changing anything.

The checks reuse the import pipeline's own code, so the findings match what the
import would do: header verification, taxonomy augmentation, a dry-run taxonomy
import with fake IDs, the location and person resolvers, and the biography
planner. The chronotopi table is built the way cell 9 of ``chronotopi.ipynb``
builds it, except that the header row is found by searching for ``event_label``
instead of using the notebook's ``skip`` table, and sheets without named columns
are skipped with a warning instead of being renamed.

The script is read-only and offline. It does not log in, call the API, use a
geocoder, or write any file. It only reports; fixes need a separate decision.

Usage, from the ``import`` folder::

    ./.venv/bin/python validate_data.py          # up to 15 findings per list
    ./.venv/bin/python validate_data.py --all    # every finding

The exit status is 1 when at least one ERROR is reported. WARNING and INFO
findings do not change it.
"""

from __future__ import annotations

import argparse
import difflib
import itertools
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

import openpyxl
import pandas as pd

from locations import (
    LOCATION_TAXONOMY_SYNONYMS_PATH,
    LOCATIONS_WORKBOOK_PATH,
    taxonomy_label_key,
    taxonomy_match_excluded,
)
from metadati import read_person_record
from person_biography import (
    PersonBiographyConflictError,
    PersonBiographyError,
    build_person_biography_plan,
)
from person_taxonomy import (
    PERSON_TAXONOMY_SYNONYMS_PATH,
    PersonaeAugmentation,
    PersonTaxonomyError,
    _is_collective_label,
    _is_personae_occurrence,
    _metadata_name_variants,
    _occurrences_for,
    augment_personae_taxonomy,
    is_placeholder_name,
    person_taxonomy_alias,
    person_taxonomy_key,
)
from rebuild_locations import (
    coordinate_pair,
    read_chronotopi_sources,
    read_existing_workbook,
    read_metadata_sources,
)
from taxonomy import (
    Taxonomy,
    TaxonomyError,
    _extract_key,
    _label_key,
    load_taxonomy,
    parse_mermaid_occurrences,
)
from taxonomy_import import TaxonomyImport, import_taxonomy
from timespan import parse_date
from utils import clean_str
from verify_excel_headers import verify_files

IMPORT_DIR = Path(__file__).resolve().parent
REPO_ROOT = IMPORT_DIR.parent
DATA_DIR = REPO_ROOT / "data"
SOURCE_DIR = DATA_DIR / "nuove schede"
TEMPLATE_DIR = DATA_DIR / "template"
MAXQDA_DIR = DATA_DIR / "maxqda"
MAXQDA_TAXONOMY = MAXQDA_DIR / "MAXQDA_Code_System.mmd"
DOCS_TAXONOMY = REPO_ROOT / "docs" / "mqda_coding_taxonomy.mmd"
LOCATION_WORKBOOK = LOCATIONS_WORKBOOK_PATH

ERROR = "ERROR"
WARNING = "WARNING"
INFO = "INFO"
SEVERITIES = (ERROR, WARNING, INFO)

FILES = "Files"
HEADERS = "Headers"
TAXONOMY = "Taxonomy"
SYNONYMS = "Synonym files"
PLACES = "Places"
PERSONS = "Persons"
SHEETS = "Chronotopi sheets"
VALUES = "Chronotopi values"
BIOGRAPHIES = "Biographies"
SECTIONS = (
    FILES,
    HEADERS,
    TAXONOMY,
    SYNONYMS,
    PLACES,
    PERSONS,
    SHEETS,
    VALUES,
    BIOGRAPHIES,
)

DEFAULT_LIMIT = 15
ROW_LIST_LIMIT = 8
CERTAINTY_VALUES = frozenset({"certain", "probable", "uncertain", "disputed"})
PRECISION_VALUES = frozenset({"day", "month", "year", "range", "after", "unknown"})
ENUM_COLUMNS: dict[str, frozenset[str]] = {
    "date_certainty": CERTAINTY_VALUES,
    "date_precision": PRECISION_VALUES,
}
CONCEPT_COLUMNS = ("event_type", "transport_mode", "place_category")
TIME_COLUMN = "date_label"
DATE_COLUMNS = ("date_start", "date_end")
LOCATION_SECTIONS = frozenset(
    {
        "aliases",
        "canonical_parents",
        "excluded_contains",
        "query_overrides",
        "reviewed_unmapped",
    }
)
PERSON_SECTIONS = frozenset({"aliases", "metadata_by_chronotopi", "metadata_relatives"})
WIKIDATA_ID = re.compile(r"(?:https?://www\.wikidata\.org/wiki/)?Q\d+")


@dataclass(frozen=True)
class Finding:
    """One problem or note, with its severity and the section it belongs to.

    ``severity`` is ``ERROR``, ``WARNING``, or ``INFO``, and ``section`` is one
    of ``SECTIONS``. Checks add findings through ``Report.add``. Only ERROR
    findings change the exit status of ``main``; the others are printed for
    review.
    """

    severity: str
    section: str
    message: str


class Report:
    """Collects findings and renders them grouped by section.

    ``main`` creates one ``Report``, passes it to each check, and prints
    ``render`` once at the end. Keeping the findings in a list lets ``render``
    cap each severity list and count the errors for the exit status.
    """

    def __init__(self) -> None:
        self.findings: list[Finding] = []

    def add(self, severity: str, section: str, message: str) -> None:
        self.findings.append(Finding(severity, section, message))

    def count(self, severity: str) -> int:
        return sum(finding.severity == severity for finding in self.findings)

    def render(self, limit: int | None) -> str:
        """Return the report text, with findings grouped by section and severity.

        Every section is printed, with ``no findings`` when it is empty. When
        ``limit`` is set, each section shows at most that many messages per
        severity, then a line that points to ``--all``. The last line gives the
        error, warning, and note counts.
        """
        lines: list[str] = []
        for section in SECTIONS:
            found = [finding for finding in self.findings if finding.section == section]
            lines.append(f"== {section}")
            if not found:
                lines.append("  no findings")
            for severity in SEVERITIES:
                messages = [
                    finding.message for finding in found if finding.severity == severity
                ]
                shown = messages if limit is None else messages[:limit]
                lines.extend(f"  {severity}: {message}" for message in shown)
                if len(shown) < len(messages):
                    lines.append(
                        f"  {severity}: ... {len(messages) - len(shown)} more (use --all)"
                    )
        lines.append(
            f"Summary: {self.count(ERROR)} errors, {self.count(WARNING)} warnings, "
            f"{self.count(INFO)} notes"
        )
        return "\n".join(lines)


@dataclass(frozen=True)
class SynonymPair:
    """One source-to-target entry of a synonym section."""

    source: str
    target: str


@dataclass(frozen=True)
class LocationSynonyms:
    """Parsed entries of data/location_taxonomy_synonyms.json.

    ``read_location_synonyms`` fills this from the reviewed JSON. The checks use
    it instead of the production loader in ``locations.synonyms``, so that
    problems the loader hides, such as repeated keys, can be reported. ``aliases``
    maps a source place to a GO label, ``parents`` maps a supplemental place to
    its parent hint, and ``excluded`` and ``reviewed`` list terms left unmapped.
    """

    aliases: list[SynonymPair]
    parents: list[SynonymPair]
    excluded: list[str]
    reviewed: list[str]

    def alias_map(self) -> dict[str, str]:
        """Alias targets keyed by normalized source label."""
        return {taxonomy_label_key(pair.source): pair.target for pair in self.aliases}

    def parent_map(self) -> dict[str, str]:
        """Canonical parents keyed by normalized child label."""
        return {taxonomy_label_key(pair.source): pair.target for pair in self.parents}


@dataclass(frozen=True)
class PersonSynonyms:
    """Parsed entries of data/person_taxonomy_synonyms.json.

    ``read_person_synonyms`` fills this from the reviewed JSON, for the same
    reason as ``LocationSynonyms``. ``aliases`` maps a spelling to a
    ``Surname, Given`` target, and ``associations`` maps each chronotopi workbook
    to its metadata workbook.
    """

    aliases: list[SynonymPair]
    associations: list[SynonymPair]

    def alias_target_keys(self) -> set[str]:
        """Normalized keys of all alias targets: the names a relative may link to."""
        return {person_taxonomy_key(pair.target) for pair in self.aliases}

    def association_map(self) -> dict[str, str]:
        """Metadata workbook file name keyed by chronotopi workbook file name."""
        return {
            Path(pair.source).name: Path(pair.target).name for pair in self.associations
        }


@dataclass(frozen=True)
class LoadedSheet:
    """One chronotopi sheet, identified by its workbook file name.

    ``load_chronotopi`` adds one entry for each sheet it reads, skipping sheets
    whose names start with an underscore. ``check_sheets`` resolves the sheet
    name to a person, and ``label`` is the text used in findings.
    """

    workbook: str
    sheet: str

    @property
    def label(self) -> str:
        return f"{self.workbook} [{self.sheet}]"


@dataclass
class ChronotopiData:
    """Chronotopi tables built the way cells 5 (overview) and 9 (frame) build them.

    ``overview`` maps each opened workbook path to all of its sheet names,
    including the underscore sheets that are skipped. ``frame`` has one row per
    data row, with ``name``, ``source_workbook``, and ``excel_row`` added, and
    every cell is text. A blank cell reads as ``nan``. ``loaded`` lists the sheets
    that were read. Unreadable workbooks and legacy sheets are left out, so the
    tables can differ from the notebook's. The sheet and biography checks read
    ``overview``; the row and value checks read ``frame``.
    """

    overview: dict[str, list[str]]
    frame: pd.DataFrame
    loaded: list[LoadedSheet]


def cell(row: Mapping[Any, Any], column: str) -> str:
    """Return a cleaned cell; the notebook's astype(str) turns blanks into "nan"."""
    return clean_str(row.get(column, ""))


def has_outer_spaces(raw: object) -> bool:
    text = "" if raw is None else str(raw)
    return text not in ("", "nan", "None") and text != text.strip()


def location_of(row: Mapping[Any, Any]) -> str:
    return f"{row['source_workbook']} [{row['name']}] row {row['excel_row']}"


def format_rows(numbers: Sequence[int]) -> str:
    shown = ", ".join(str(number) for number in numbers[:ROW_LIST_LIMIT])
    noun = "row" if len(numbers) == 1 else "rows"
    more = " ..." if len(numbers) > ROW_LIST_LIMIT else ""
    return f"{noun} {shown}{more}"


def summarize(counts: Mapping[str, int], limit: int = 10) -> str:
    items = [f"{value!r} ({count})" for value, count in sorted(counts.items())[:limit]]
    hidden = len(counts) - limit
    return ", ".join(items) + (f", and {hidden} more" if hidden > 0 else "")


# --- Files and headers -------------------------------------------------------


def chronotopi_workbooks() -> list[Path]:
    """Mirror cell 2: top-level chronotopi files, without hidden or marker files.

    The notebook globs ``*chronotop*`` in the source folder and drops the
    ``:Zone.Identifier`` marker files. This does the same and sorts the result.
    ``check_files`` calls it, and ``main`` passes the ``.xlsx`` paths that
    ``check_files`` returns to the later checks. Non-xlsx matches are included so
    that ``check_files`` can report them.
    """
    return sorted(
        path
        for path in SOURCE_DIR.glob("*chronotop*")
        if not path.name.startswith(".") and not path.name.endswith(":Zone.Identifier")
    )


def check_files(report: Report) -> list[Path]:
    """Check the source folder for unreadable workbooks and non-.xlsx chronotopi files.

    The notebook opens every chronotopi file with openpyxl and stops at the first
    one that fails, so this check reports each bad file before the run. Every
    ``*.xlsx`` in the source folder is also opened read-only, because the metadata
    and belege imports read workbooks from the same folder. Returns the chronotopi
    ``.xlsx`` paths in sorted order for the checks that read them.
    """
    for path in chronotopi_workbooks():
        if path.suffix != ".xlsx":
            report.add(
                ERROR,
                FILES,
                f"{path.name}: chronotopi file is not an .xlsx workbook; the notebook cannot read it",
            )
    for path in sorted(SOURCE_DIR.glob("*.xlsx")):
        try:
            openpyxl.load_workbook(path, read_only=True).close()
        except Exception as exc:
            report.add(
                ERROR, FILES, f"{path.name}: cannot be opened as a workbook ({exc})"
            )
    return [path for path in chronotopi_workbooks() if path.suffix == ".xlsx"]


def check_headers(report: Report) -> None:
    """Compare every source workbook header with its template (verify_excel_headers).

    ``verify_files`` checks each ``*.xlsx`` in the source folder against its
    template. A missing column is an ERROR, because the import reads columns by
    name. An extra column is a WARNING. If the verifier itself fails, for example
    because a folder is missing, that failure is reported as one ERROR.
    """
    result = verify_files(SOURCE_DIR, TEMPLATE_DIR)
    if result.error:
        report.add(ERROR, HEADERS, f"header verification failed: {result.error}")
        return
    for file_result in result.files:
        for error in file_result.errors:
            details = []
            if error.missing:
                details.append(f"missing: {', '.join(error.missing)}")
            if error.extra:
                details.append(f"extra: {', '.join(error.extra)}")
            severity = ERROR if error.missing else WARNING
            suffix = f" ({'; '.join(details)})" if details else ""
            report.add(
                severity,
                HEADERS,
                f"{file_result.name} [{error.sheet}]: {error.message}{suffix}",
            )


def check_rebuild_scope(report: Report) -> None:
    """Flag location-rebuild sources under data/ that lie outside the source folder.

    ``rebuild_locations`` searches all of data/ recursively for chronotopi and
    metadata workbooks, but the notebook reads only the source folder. A file
    found elsewhere can change the location workbook without any other check
    seeing it, so each one is reported as a WARNING.
    """
    for pattern in ("*chronotop*.xlsx", "metadati*.xlsx"):
        for path in sorted(DATA_DIR.rglob(pattern)):
            if path.parent != SOURCE_DIR:
                report.add(
                    WARNING,
                    PLACES,
                    f"{path.relative_to(REPO_ROOT)}: the location rebuild reads this workbook as a source, "
                    f"but it is outside {SOURCE_DIR.relative_to(REPO_ROOT)}",
                )


# --- Taxonomy ----------------------------------------------------------------


def taxonomy_label_set(path: Path) -> set[str]:
    return {
        _label_key(occurrence.label) for occurrence in parse_mermaid_occurrences(path)
    }


def check_taxonomy_copies(report: Report) -> None:
    """Compare the MAXQDA code-system copies; only MAXQDA_Code_System.mmd is imported.

    ``MAXQDA_Code_System.mmd`` is the only copy the import reads. The other
    ``*.mmd`` files in data/maxqda/ and the docs copy are not read, so a difference
    would go unnoticed. Each copy is parsed and its label set is compared with the
    used file. A copy that cannot be parsed, or whose labels differ, is a WARNING.
    """
    reference = taxonomy_label_set(MAXQDA_TAXONOMY)
    for copy in [DOCS_TAXONOMY, *sorted(MAXQDA_DIR.glob("*.mmd"))]:
        if copy == MAXQDA_TAXONOMY or not copy.is_file():
            continue
        try:
            labels = taxonomy_label_set(copy)
        except Exception as exc:
            report.add(
                WARNING,
                TAXONOMY,
                f"{copy.relative_to(REPO_ROOT)} cannot be parsed ({exc})",
            )
            continue
        if labels != reference:
            report.add(
                WARNING,
                TAXONOMY,
                f"{copy.relative_to(REPO_ROOT)} differs from {MAXQDA_TAXONOMY.name}: "
                f"only in the copy {sorted(labels - reference)[:5]}, "
                f"only in the used file {sorted(reference - labels)[:5]}",
            )


def dry_run_import(
    report: Report, taxonomy: Taxonomy, person_aliases: Mapping[str, str]
) -> TaxonomyImport | None:
    """Run the real import with fake IDs, so resolvers can be queried without writes.

    This calls ``taxonomy_import.import_taxonomy`` with a counter in place of the
    API adapter, so the parent-first ordering, the alias table, and the person
    alias checks are the same as in the import. The returned ``TaxonomyImport`` is
    what the place, sheet, and value checks query. It is None after an ERROR when
    the taxonomy or the Personae aliases are rejected.
    """
    ids = itertools.count(1)

    def create_concept(label: str, **_options: object) -> int:
        return next(ids)

    try:
        importer = import_taxonomy(
            taxonomy, create_concept, person_aliases=dict(person_aliases)
        )
    except (TaxonomyError, PersonTaxonomyError) as exc:
        report.add(ERROR, TAXONOMY, f"dry-run import failed: {exc}")
        return None
    report.add(
        INFO,
        TAXONOMY,
        f"dry-run import resolved {len(importer.ids_by_key)} taxonomy keys with fake IDs",
    )
    return importer


# --- Synonym files -----------------------------------------------------------


def load_strict_json(path: Path, report: Report) -> dict[str, Any] | None:
    """Parse a synonym file and report duplicate keys, which json.loads would hide.

    ``json.loads`` silently keeps the last value of a repeated key, and the
    production loaders use the same parser, so each repeated key is reported as an
    ERROR here. Returns the parsed object, or None after an ERROR when the file
    cannot be read, is not valid JSON, or is not a JSON object.
    """
    duplicates: list[str] = []

    def keep_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        """Build one JSON object from key pairs and note each repeated key.

        Repeated keys are appended to the enclosing ``duplicates`` list. As in
        ``json.loads``, the last value still wins.
        """
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                duplicates.append(key)
            result[key] = value
        return result

    try:
        data = json.loads(
            path.read_text(encoding="utf-8"), object_pairs_hook=keep_pairs
        )
    except (OSError, json.JSONDecodeError) as exc:
        report.add(ERROR, SYNONYMS, f"{path.name} cannot be read as JSON: {exc}")
        return None
    for key in dict.fromkeys(duplicates):
        report.add(
            ERROR,
            SYNONYMS,
            f"{path.name}: key {key!r} is defined more than once; only the last value is used",
        )
    if not isinstance(data, dict):
        report.add(ERROR, SYNONYMS, f"{path.name}: the top level must be a JSON object")
        return None
    return data


def pairs_of(
    section: object, file_name: str, key: str, report: Report
) -> list[SynonymPair]:
    """Return an object section as SynonymPair entries with string values.

    A section that is not a JSON object is an ERROR, and the result is empty. A
    value that is not a string is an ERROR and is skipped, so one bad entry does
    not hide the others. ``file_name`` and ``key`` only label the messages.
    """
    if not isinstance(section, dict):
        report.add(
            ERROR, SYNONYMS, f"{file_name}: section {key!r} must be a JSON object"
        )
        return []
    pairs: list[SynonymPair] = []
    for source, target in section.items():
        if isinstance(target, str):
            pairs.append(SynonymPair(str(source), target))
        else:
            report.add(
                ERROR,
                SYNONYMS,
                f"{file_name}: {key!r} value for {source!r} must be a string",
            )
    return pairs


def items_of(section: object, file_name: str, key: str, report: Report) -> list[str]:
    """Return an array section as strings.

    A section that is not a JSON array of strings is an ERROR, and the result is
    empty, so the other sections are still checked.
    """
    if not isinstance(section, list) or not all(
        isinstance(item, str) for item in section
    ):
        report.add(
            ERROR,
            SYNONYMS,
            f"{file_name}: section {key!r} must be a JSON array of strings",
        )
        return []
    return list(section)


def text_problems(text: str, *, punctuation: bool) -> list[str]:
    """Return the spacing and punctuation problems in one synonym string.

    Leading or trailing spaces and double spaces are always checked. Trailing
    punctuation is checked only when ``punctuation`` is true, which is the case
    for alias and parent entries. A stray comma or colon at the end of a location
    target stops it from matching its GO label, because taxonomy keys fold only
    case, dashes, and spaces.
    """
    problems: list[str] = []
    if text != text.strip():
        problems.append("leading or trailing spaces")
    if "  " in text:
        problems.append("double spaces")
    if punctuation and re.search(r"[,;:]$", text.strip()):
        problems.append("trailing punctuation")
    return problems


def check_text(
    report: Report, file_name: str, role: str, text: str, *, punctuation: bool
) -> None:
    for problem in text_problems(text, punctuation=punctuation):
        report.add(WARNING, SYNONYMS, f"{file_name}: {role} {text!r} has {problem}")


def report_collisions(
    report: Report,
    file_name: str,
    pairs: Sequence[SynonymPair],
    normalize: Callable[[str], str],
) -> None:
    """Flag sources that normalize to one key; conflicting targets are errors.

    Sources are grouped by their normalized key. Targets that still differ after
    normalization are an ERROR. The location loader silently keeps the last value,
    and the person loader raises PersonTaxonomyError. Identical targets are a
    WARNING, because they are redundant.
    """
    groups: defaultdict[str, list[SynonymPair]] = defaultdict(list)
    for pair in pairs:
        groups[normalize(pair.source)].append(pair)
    for key, members in groups.items():
        if len(members) < 2:
            continue
        conflicting = len({normalize(pair.target) for pair in members}) > 1
        severity = ERROR if conflicting else WARNING
        kind = "conflicting" if conflicting else "identical"
        listing = "; ".join(f"{pair.source!r} -> {pair.target!r}" for pair in members)
        report.add(
            severity,
            SYNONYMS,
            f"{file_name}: sources normalize to {key!r} with {kind} targets: {listing}",
        )


def read_location_synonyms(report: Report) -> LocationSynonyms | None:
    """Read the reviewed location synonyms and check them for problems.

    Unknown sections are WARNINGs and are ignored, as the production loader
    ignores them. ``query_overrides`` is a known section whose entries are not
    checked here. The alias, parent, excluded, and reviewed entries get the spacing
    and collision checks. Returns None when ``load_strict_json`` rejects the file,
    and the graph and parent checks are then skipped.
    """
    name = LOCATION_TAXONOMY_SYNONYMS_PATH.name
    data = load_strict_json(LOCATION_TAXONOMY_SYNONYMS_PATH, report)
    if data is None:
        return None
    for section in sorted(set(data) - LOCATION_SECTIONS):
        report.add(WARNING, SYNONYMS, f"{name}: unknown section {section!r} is ignored")
    synonyms = LocationSynonyms(
        aliases=pairs_of(data.get("aliases", {}), name, "aliases", report),
        parents=pairs_of(
            data.get("canonical_parents", {}), name, "canonical_parents", report
        ),
        excluded=items_of(
            data.get("excluded_contains", []), name, "excluded_contains", report
        ),
        reviewed=items_of(
            data.get("reviewed_unmapped", []), name, "reviewed_unmapped", report
        ),
    )
    for pair in [*synonyms.aliases, *synonyms.parents]:
        check_text(report, name, "key", pair.source, punctuation=True)
        check_text(report, name, "target", pair.target, punctuation=True)
    for term in [*synonyms.excluded, *synonyms.reviewed]:
        check_text(report, name, "entry", term, punctuation=False)
    report_collisions(report, name, synonyms.aliases, taxonomy_label_key)
    report_collisions(report, name, synonyms.parents, taxonomy_label_key)
    report_collisions(
        report,
        name,
        [SynonymPair(term, term) for term in synonyms.excluded],
        taxonomy_label_key,
    )
    report_collisions(
        report,
        name,
        [SynonymPair(term, term) for term in synonyms.reviewed],
        taxonomy_label_key,
    )
    return synonyms


def check_metadata_relatives(
    report: Report, file_name: str, section: object, synonyms: PersonSynonyms
) -> None:
    """Check that relative links point at a reviewed 'Surname, Given' alias target.

    ``metadata_relatives`` maps each metadata workbook to its relative links. A
    link target must contain a comma and must equal the target of some alias, so
    the link names a person that the reviewed aliases already use. A section that
    is not a JSON object is an ERROR, and each link that breaks the rule is an
    ERROR that names the relative.
    """
    if not isinstance(section, dict):
        report.add(
            ERROR,
            SYNONYMS,
            f"{file_name}: section 'metadata_relatives' must be a JSON object",
        )
        return
    alias_targets = synonyms.alias_target_keys()
    for workbook, links in section.items():
        pairs = pairs_of(links, file_name, f"metadata_relatives {workbook}", report)
        for link in pairs:
            check_text(
                report, file_name, "relative link", link.source, punctuation=False
            )
            target_key = person_taxonomy_key(link.target)
            if "," not in link.target or target_key not in alias_targets:
                report.add(
                    ERROR,
                    SYNONYMS,
                    f"{file_name}: relative {link.source!r} of {workbook} "
                    f"must link to a reviewed alias target, not {link.target!r}",
                )


def read_person_synonyms(report: Report) -> PersonSynonyms | None:
    """Read the reviewed person synonyms and check them for problems.

    The production loader in ``person_taxonomy.synonyms`` ignores unknown sections
    and keeps the last of a repeated key, so this function reports both. An alias
    target without a comma is a WARNING. A placeholder name part such as
    ``unbekannt`` is an ERROR, because augmentation could create a person with
    that name. Returns None when ``load_strict_json`` rejects the file.
    """
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    data = load_strict_json(PERSON_TAXONOMY_SYNONYMS_PATH, report)
    if data is None:
        return None
    for section in sorted(set(data) - PERSON_SECTIONS):
        report.add(WARNING, SYNONYMS, f"{name}: unknown section {section!r} is ignored")
    synonyms = PersonSynonyms(
        aliases=pairs_of(data.get("aliases", {}), name, "aliases", report),
        associations=pairs_of(
            data.get("metadata_by_chronotopi", {}),
            name,
            "metadata_by_chronotopi",
            report,
        ),
    )
    for pair in synonyms.aliases:
        target = pair.target
        check_text(report, name, "alias key", pair.source, punctuation=True)
        check_text(report, name, "alias target", target, punctuation=True)
        if "," not in target:
            report.add(
                WARNING,
                SYNONYMS,
                f"{name}: alias target {target!r} is not in 'Surname, Given' form",
            )
            continue
        family, given = (part.strip() for part in target.split(",", 1))
        if is_placeholder_name(family) or is_placeholder_name(given):
            report.add(
                ERROR,
                SYNONYMS,
                f"{name}: alias target {target!r} has a placeholder name part",
            )
    for pair in synonyms.associations:
        check_text(report, name, "association key", pair.source, punctuation=False)
        check_text(report, name, "association value", pair.target, punctuation=False)
    check_metadata_relatives(report, name, data.get("metadata_relatives", {}), synonyms)
    report_collisions(report, name, synonyms.aliases, person_taxonomy_key)
    return synonyms


# --- Places ------------------------------------------------------------------


def parent_cycles(parents: Mapping[str, str]) -> list[list[str]]:
    """Return each distinct cycle in the normalized parent map once.

    Each child is followed up its parent chain until the chain ends or a node
    repeats. A repeated node closes a cycle, and cycles with the same members are
    returned once, however many children lead into them. The import raises when it
    reaches a location in a cycle, so ``check_location_graph`` reports each one as
    an ERROR.
    """
    cycles: list[list[str]] = []
    seen: set[frozenset[str]] = set()
    for start in parents:
        path: list[str] = []
        node: str | None = start
        while node is not None and node in parents and node not in path:
            path.append(node)
            node = taxonomy_label_key(parents[node])
        if node is not None and node in path:
            cycle = path[path.index(node) :]
            signature = frozenset(cycle)
            if signature not in seen:
                seen.add(signature)
                cycles.append([*cycle, node])
    return cycles


def check_location_graph(
    report: Report, synonyms: LocationSynonyms, go_keys: set[str]
) -> None:
    """Check the normalized alias and parent maps for chains, cycles, and overlaps.

    ``go_keys`` holds the normalized GO labels. Lookups stop after one alias step,
    so an alias whose target is another alias, and not a GO label, is a WARNING. A
    cycle in the parent map is an ERROR. A reviewed exclusion that also has an
    alias or parent entry is a WARNING, because the exclusion wins and the entry is
    never used. An ``excluded_contains`` term matches by substring, so any entry
    whose key contains it gets the same WARNING.
    """
    name = LOCATION_TAXONOMY_SYNONYMS_PATH.name
    aliases = synonyms.alias_map()
    parents = synonyms.parent_map()
    keys = set(aliases) | set(parents)
    for source, target in aliases.items():
        target_key = taxonomy_label_key(target)
        if target_key != source and target_key in aliases and target_key not in go_keys:
            report.add(
                WARNING,
                PLACES,
                f"{name}: alias {source!r} -> {target!r}: the target is itself an alias and not a "
                "GO label, and lookups stop after one step",
            )
    for cycle in parent_cycles(parents):
        report.add(
            ERROR, PLACES, f"{name}: canonical_parents cycle {' -> '.join(cycle)}"
        )
    for term in sorted({taxonomy_label_key(term) for term in synonyms.reviewed}):
        if term in keys:
            report.add(
                WARNING,
                PLACES,
                f"{name}: reviewed_unmapped {term!r} also has an alias or parent entry; "
                "the exclusion wins and that entry is never used",
            )
    for term in sorted({taxonomy_label_key(term) for term in synonyms.excluded}):
        hits = sorted(key for key in keys if term and term in key)
        if hits:
            report.add(
                WARNING,
                PLACES,
                f"{name}: excluded_contains {term!r} matches {hits[:5]}; "
                "those alias or parent entries are never used",
            )


def check_places(
    report: Report, importer: TaxonomyImport | None, synonyms: LocationSynonyms | None
) -> None:
    """Dry-run the location linking for every name the rebuild and sync would see.

    Names come from the chronotopi and metadata workbooks that the location
    rebuild reads under data/, and from the rows of locations.xlsx. Each name is
    passed to ``ensure_location_concept``, the resolver that the notebook's
    location sync uses. A name excluded as reviewed-unmapped is listed in one INFO
    line, a name with no direct GO label is listed as a new supplemental concept,
    also in INFO, and a name that cannot be linked is an ERROR. Parent hints from
    the synonym file are dry-run too. Without a dry-run importer, the check is
    skipped with one INFO line.
    """
    if importer is None:
        report.add(
            INFO,
            PLACES,
            "place checks skipped: the dry-run taxonomy import did not complete",
        )
        return
    names = {source.name for source in read_chronotopi_sources(DATA_DIR).values()}
    names.update(source.name for source in read_metadata_sources(DATA_DIR).values())
    if LOCATION_WORKBOOK.is_file():
        _, records = read_existing_workbook(LOCATION_WORKBOOK)
        names.update(clean_str(record.get("location")) for record in records.values())
    names.discard("")
    excluded: list[str] = []
    supplemental: list[str] = []
    for name in sorted(names, key=str.casefold):
        if taxonomy_match_excluded(name):
            excluded.append(name)
            continue
        direct = importer.resolve_location_exact(name) is not None
        try:
            importer.ensure_location_concept(name)
        except TaxonomyError as exc:
            report.add(ERROR, PLACES, f"place {name!r} cannot be linked: {exc}")
            continue
        if not direct:
            supplemental.append(name)
    if synonyms is not None:
        for pair in synonyms.parents:
            try:
                importer.ensure_location_concept(pair.target)
            except TaxonomyError as exc:
                report.add(
                    ERROR,
                    PLACES,
                    f"parent hint {pair.target!r} cannot be linked: {exc}",
                )
    if excluded:
        report.add(
            INFO,
            PLACES,
            f"{len(excluded)} place names are reviewed as unmapped: "
            f"{', '.join(sorted(excluded, key=str.casefold))}",
        )
    if supplemental:
        report.add(
            INFO,
            PLACES,
            f"{len(supplemental)} place names have no direct GO label or alias and become "
            f"supplemental concepts: {', '.join(sorted(supplemental, key=str.casefold))}",
        )


# --- Persons -----------------------------------------------------------------


def check_person_associations(
    report: Report, workbooks: Sequence[Path], synonyms: PersonSynonyms
) -> None:
    """Check that each chronotopi workbook has a reviewed metadata workbook.

    Augmentation stops at the first chronotopi workbook without an association,
    so a missing association is an ERROR. An association that names a metadata
    file missing from the source folder is an ERROR. An association for a workbook
    that is not in the source folder is a WARNING. A metadata workbook that no
    association names is a WARNING, because its person is not placed under
    Personae.
    """
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    mapping = synonyms.association_map()
    chrono_names = {path.name for path in workbooks}
    metadata_files = {path.name for path in SOURCE_DIR.glob("metadati*.xlsx")}
    for chrono in sorted(chrono_names - mapping.keys()):
        report.add(
            ERROR,
            PERSONS,
            f"{chrono}: no metadata association in {name}; augmentation stops at this workbook",
        )
    for chrono in sorted(mapping.keys() - chrono_names):
        report.add(
            WARNING,
            PERSONS,
            f"{name}: association for {chrono} matches no chronotopi workbook in {SOURCE_DIR.name}",
        )
    for chrono, metadata in sorted(mapping.items()):
        if metadata not in metadata_files:
            report.add(
                ERROR,
                PERSONS,
                f"{name}: {chrono} points to {metadata}, which does not exist in {SOURCE_DIR.name}",
            )
    for metadata in sorted(metadata_files - set(mapping.values())):
        report.add(
            WARNING,
            PERSONS,
            f"{metadata}: not associated with any chronotopi workbook; its person is not placed under Personae",
        )


def raw_birth_date(path: Path) -> str:
    """Return the Geburtsdatum text as written in the 'Scheda individuale' sheet.

    ``read_person_record`` turns this field into a date interval, and a value it
    cannot parse leaves that interval as None. The raw text lets the caller tell an
    unparseable date from a missing one. Returns an empty string when the row or
    its value is missing. A workbook without the 'Scheda individuale' sheet raises
    KeyError.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for row in workbook["Scheda individuale"].iter_rows(values_only=True):
            if row and clean_str(row[0]) == "Geburtsdatum":
                return clean_str(row[1]) if len(row) > 1 else ""
        return ""
    finally:
        workbook.close()


def read_metadata_records(report: Report) -> dict[str, dict[str, Any]]:
    """Read each metadati workbook in the source folder, reporting unreadable ones.

    Only the top level of the source folder is searched, which matches what the
    augmentation reads. The result is keyed by file name. A workbook that
    ``read_person_record`` cannot parse is an ERROR and is left out.
    """
    records: dict[str, dict[str, Any]] = {}
    for path in sorted(SOURCE_DIR.glob("metadati*.xlsx")):
        try:
            records[path.name] = read_person_record(path)
        except Exception as exc:
            report.add(ERROR, PERSONS, f"{path.name}: metadata cannot be read ({exc})")
    return records


def check_metadata_records(
    report: Report, records: Mapping[str, Mapping[str, Any]]
) -> None:
    """Check each metadata record's identifier, main name, and birth date.

    A record without a metadata identifier (Metadaten ID) is a WARNING, and so is
    a main name that is missing or a placeholder. A birth date that the parser
    cannot read is a WARNING that shows the raw text, and the date is then not
    compared with chronotopi dates. Two records with the same identifier are an
    ERROR, because the import looks up people by this identifier.
    """
    owners: defaultdict[str, list[str]] = defaultdict(list)
    for file_name, record in records.items():
        person = record.get("person") or {}
        identifier = clean_str(record.get("identifier"))
        if identifier:
            owners[identifier].append(file_name)
        else:
            report.add(
                WARNING,
                PERSONS,
                f"{file_name}: no metadata identifier (Metadaten ID) found",
            )
        given = person.get("given_name")
        family = person.get("family_name")
        if is_placeholder_name(given) or is_placeholder_name(family):
            report.add(
                WARNING,
                PERSONS,
                f"{file_name}: main name is missing or a placeholder "
                f"(given {clean_str(given)!r}, family {clean_str(family)!r})",
            )
        if person.get("birth_date_interval") is None:
            raw = raw_birth_date(SOURCE_DIR / file_name)
            if raw:
                report.add(
                    WARNING,
                    PERSONS,
                    f"{file_name}: birth date {raw!r} cannot be parsed; it is not compared with chronotopi",
                )
    for identifier, files in owners.items():
        if len(files) > 1:
            report.add(
                ERROR,
                PERSONS,
                f"identifier {identifier} is used by {', '.join(sorted(files))}",
            )


def augment_dry_run(
    report: Report, taxonomy: Taxonomy, overview: Mapping[str, Sequence[str]]
) -> PersonaeAugmentation | None:
    """Run the Personae augmentation in memory and report a failure as an ERROR.

    The augmentation reads the associated metadata workbooks and merges their
    people under each surname branch. It writes nothing. A missing association, a
    conflict between person names, or a metadata record without a usable surname
    raises an error, which is reported here. Returns None in that case, and
    ``main`` then skips the dry-run import and the placement checks.
    """
    try:
        return augment_personae_taxonomy(taxonomy, overview, data_root=SOURCE_DIR)
    except (PersonTaxonomyError, OSError, KeyError, ValueError) as exc:
        report.add(ERROR, PERSONS, f"augmentation failed: {exc}")
        return None


def check_person_placement(
    report: Report,
    overview: Mapping[str, Sequence[str]],
    taxonomy: Taxonomy,
    synonyms: PersonSynonyms,
    records: Mapping[str, Mapping[str, Any]],
    person_aliases: Mapping[str, str],
) -> None:
    """Find aliases the augmentation silently drops and sheets that become raw leaves.

    The augmentation drops an alias whose target is not a Personae leaf and that no
    sheet uses, so each such alias is a WARNING. A chronotopi sheet that matches no
    metadata name, MAXQDA person, or alias becomes a raw Personae leaf, which is
    also a WARNING. Sheets whose names start with an underscore are skipped, as in
    the notebook.
    """
    name = PERSON_TAXONOMY_SYNONYMS_PATH.name
    for pair in synonyms.aliases:
        if person_taxonomy_key(pair.source) not in person_aliases:
            report.add(
                WARNING,
                PERSONS,
                f"{name}: alias {pair.source!r} -> {pair.target!r} is ignored; "
                "its target is not a Personae leaf and no sheet uses it",
            )
    occurrences = _occurrences_for(taxonomy)
    parent_indexes = {
        occurrence.parent_index
        for occurrence in occurrences
        if occurrence.parent_index is not None
    }
    known = {
        person_taxonomy_key(occurrence.label)
        for index, occurrence in enumerate(occurrences)
        if _is_personae_occurrence(occurrence)
        and index not in parent_indexes
        and not _is_collective_label(occurrence.label)
    }
    for record in records.values():
        for variant in _metadata_name_variants(record.get("person") or {}):
            known.add(person_taxonomy_key(variant))
    for pair in synonyms.aliases:
        known.add(person_taxonomy_key(pair.source))
        known.add(person_taxonomy_key(pair.target))
    for path_text, sheets in overview.items():
        workbook = Path(path_text).name
        for sheet in sheets:
            if sheet.startswith("_"):
                continue
            if (
                person_taxonomy_key(sheet) in known
                or person_taxonomy_alias(sheet) is not None
            ):
                continue
            report.add(
                WARNING,
                PERSONS,
                f"{workbook} [{sheet}]: no metadata record, MAXQDA person, or alias matches; "
                "the sheet becomes a raw Personae leaf",
            )


# --- Chronotopi sheets -------------------------------------------------------


def read_sheet(report: Report, path: Path, sheet: str) -> pd.DataFrame | None:
    """Load one sheet as cell 9 does, locating the header by its event_label cell.

    The header is searched for in the first ten rows. If none has ``event_label``,
    an ERROR is reported and row 1 is used, as the notebook does. A legacy sheet
    with unnamed columns is skipped with a WARNING, where the notebook renames the
    columns. The frame gets ``name``, ``source_workbook``, and ``excel_row``
    columns, where ``excel_row`` is the spreadsheet row of each data row. Repeated
    headers are renamed by pandas (for example ``place_name.1``) before the
    duplicate check runs, so that WARNING is not reached in practice.
    """
    where = f"{path.name} [{sheet}]"
    preview = pd.read_excel(path, sheet_name=sheet, header=None, nrows=10)
    preview_rows = preview.itertuples(index=False, name=None)
    header_row = next(
        (
            position
            for position, values in enumerate(preview_rows)
            if "event_label"
            in {str(value).strip() for value in values if pd.notna(value)}
        ),
        None,
    )
    if header_row is None:
        report.add(
            ERROR,
            SHEETS,
            f"{where}: no 'event_label' header in the first 10 rows; the notebook reads row 1 as the header",
        )
        header_row = 0
    frame = pd.read_excel(path, sheet_name=sheet, header=header_row)
    if "Unnamed:" in str(frame.columns[0]):
        report.add(
            WARNING,
            SHEETS,
            f"{where}: legacy layout without named columns; the notebook renames them and this check skips the sheet",
        )
        return None
    duplicated = [str(column) for column in frame.columns[frame.columns.duplicated()]]
    if duplicated:
        report.add(
            WARNING,
            SHEETS,
            f"{where}: duplicate column headers {duplicated}; only the first column of each name is read",
        )
    frame = frame.astype(str)
    frame["name"] = sheet
    frame["source_workbook"] = path.name
    frame["excel_row"] = [header_row + index + 2 for index in range(len(frame))]
    return frame


def load_chronotopi(report: Report, workbooks: Sequence[Path]) -> ChronotopiData:
    """Read every chronotopi sheet into one table, as cells 5 and 9 do.

    A workbook that openpyxl cannot open is left out here, because ``check_files``
    already reports it. Sheets whose names start with an underscore are listed in
    ``overview`` but not read. When no sheet is read, the frame is empty but keeps
    the four key columns that the later checks use.
    """
    overview: dict[str, list[str]] = {}
    loaded: list[LoadedSheet] = []
    frames: list[pd.DataFrame] = []
    for path in workbooks:
        try:
            workbook = openpyxl.load_workbook(path, read_only=True)
        except Exception:
            continue  # already reported under Files
        sheet_names = list(workbook.sheetnames)
        workbook.close()
        overview[str(path)] = sheet_names
        for sheet in sheet_names:
            if sheet.startswith("_"):
                continue
            frame = read_sheet(report, path, sheet)
            if frame is not None:
                frames.append(frame)
                loaded.append(LoadedSheet(path.name, sheet))
    if frames:
        combined = pd.concat(frames, ignore_index=True).astype(str)
        add_locations(combined)
    else:
        combined = pd.DataFrame(
            columns=["name", "source_workbook", "excel_row", "place_name"]
        )
    return ChronotopiData(overview, combined, loaded)


def add_locations(frame: pd.DataFrame) -> None:
    """Mirror the end of cell 9: start_location is the previous place of the same sheet.

    The rows must already be in sheet order. ``end_location`` is ``place_name``.
    A blank place reads as the text ``nan`` (see ``read_sheet``). This function
    treats that text as a place, as the notebook does, so the start location of the
    row after a blank place can be ``nan``.
    """
    if "place_name" not in frame.columns:
        frame["place_name"] = ""
    previous: dict[str, str] = {}
    starts: list[str] = []
    for name, place in zip(frame["name"], frame["place_name"]):
        end = str(place).strip()
        starts.append(previous.get(str(name), end))
        if end:
            previous[str(name)] = end
    frame["start_location"] = starts
    frame["end_location"] = frame["place_name"]


def check_sheets(
    report: Report,
    chrono: ChronotopiData,
    workbook_count: int,
    importer: TaxonomyImport | None,
) -> None:
    """Check the loaded sheets: counts, stray spaces, duplicates, and people.

    The counts of loaded sheets and of underscore sheets that are skipped are INFO
    lines. A sheet name with stray spaces is a WARNING. A sheet name found in more
    than one workbook is a WARNING, because events from both are attached to one
    person. With a dry-run importer, each sheet name is resolved to a Personae
    concept: an ambiguous name is an ERROR, and a name that resolves to nothing is
    a WARNING, since its events get no person. Without an importer, the person
    checks are skipped with one INFO line.
    """
    report.add(
        INFO,
        SHEETS,
        f"{len(chrono.loaded)} sheets loaded from {workbook_count} chronotopi workbooks",
    )
    skipped = [
        f"{Path(path).name} [{sheet}]"
        for path, sheets in chrono.overview.items()
        for sheet in sheets
        if sheet.startswith("_")
    ]
    if skipped:
        report.add(
            INFO,
            SHEETS,
            f"{len(skipped)} sheets starting with '_' are skipped by design: {', '.join(skipped)}",
        )
    workbooks_by_sheet: defaultdict[str, set[str]] = defaultdict(set)
    for item in chrono.loaded:
        workbooks_by_sheet[item.sheet].add(item.workbook)
        if item.sheet != item.sheet.strip() or "  " in item.sheet:
            report.add(
                WARNING,
                SHEETS,
                f"{item.workbook} [{item.sheet!r}]: sheet name has stray spaces",
            )
    for sheet, workbooks in sorted(workbooks_by_sheet.items()):
        if len(workbooks) > 1:
            report.add(
                WARNING,
                SHEETS,
                f"sheet {sheet!r} appears in {', '.join(sorted(workbooks))}; "
                "events from both are attached to one person",
            )
    if importer is None:
        report.add(
            INFO,
            SHEETS,
            "person checks skipped: the dry-run taxonomy import did not complete",
        )
        return
    for item in chrono.loaded:
        try:
            person = importer.resolve_person(item.sheet)
        except TaxonomyError as exc:
            report.add(ERROR, SHEETS, f"{item.label}: person name is ambiguous ({exc})")
            continue
        if person is None:
            report.add(
                WARNING,
                SHEETS,
                f"{item.label}: the sheet name does not resolve to a Personae concept; "
                "its events get no person",
            )


def check_rows(report: Report, rows: Sequence[Mapping[Any, Any]]) -> None:
    """Find rows the import skips and events that repeat an earlier row.

    ``prepare_event_data`` in ``events`` returns nothing for a row without an event
    label or without a place, so the import skips that row. Such rows are listed
    as one WARNING per sheet and kind, with their row numbers. A row that repeats
    an earlier one, with the same sheet, event, place, and dates, is reported for
    review.
    """
    missing_place: defaultdict[str, list[int]] = defaultdict(list)
    missing_event: defaultdict[str, list[int]] = defaultdict(list)
    first_seen: dict[tuple[str, ...], str] = {}
    for row in rows:
        event = cell(row, "event_label")
        place = cell(row, "place_name")
        sheet_key = f"{row['source_workbook']} [{row['name']}]"
        if event and not place:
            missing_place[sheet_key].append(int(row["excel_row"]))
        elif place and not event:
            missing_event[sheet_key].append(int(row["excel_row"]))
        elif event and place:
            key = (
                str(row["name"]),
                event,
                place,
                cell(row, TIME_COLUMN),
                cell(row, "date_start"),
                cell(row, "date_end"),
            )
            if key in first_seen:
                report.add(
                    WARNING,
                    SHEETS,
                    f"{location_of(row)} repeats {first_seen[key]} (same person, event, place, and dates)",
                )
            else:
                first_seen[key] = location_of(row)
    for sheet_key, numbers in missing_place.items():
        report.add(
            WARNING,
            SHEETS,
            f"{sheet_key}: {format_rows(numbers)} have an event but no place_name; the import skips them",
        )
    for sheet_key, numbers in missing_event.items():
        report.add(
            WARNING,
            SHEETS,
            f"{sheet_key}: {format_rows(numbers)} have a place but no event_label; the import skips them",
        )


# --- Chronotopi values -------------------------------------------------------


def closest_label(value: str, labels: Mapping[str, str]) -> str | None:
    """Suggest the taxonomy label closest to a value that matches nothing.

    Only a close match is suggested: the normalized value must reach a similarity
    of 0.85 with a normalized label. A value with no near neighbour is reported as
    dropped instead. ``labels`` maps normalized labels to their display form.

    >>> closest_label("Gebuurt", {"geburt": "Geburt"})
    'Geburt'
    >>> closest_label("Reise", {"geburt": "Geburt"}) is None
    True
    """
    matches = difflib.get_close_matches(
        _label_key(value), list(labels), n=1, cutoff=0.85
    )
    return labels[matches[0]] if matches else None


def check_column(
    report: Report,
    column: str,
    values: Counter[str],
    *,
    allowed: frozenset[str] | None,
    resolve: Callable[[str], int | None] | None,
    labels: Mapping[str, str],
) -> None:
    """Check one column's distinct values: spelling, layout, vocabulary, taxonomy.

    Values that differ only in case, dashes, or spacing are reported together as
    spelling variants. Multi-line values match nothing, so each is a WARNING, and
    so is each URL in a value column. A column with an ``allowed`` vocabulary is
    checked against it. Otherwise each value goes to ``resolve``, the dry-run
    lookup, and a value it finds ambiguous is an ERROR. A value it cannot resolve
    gets a closest-label WARNING from ``labels``, or is listed as dropped in one
    INFO line.
    """
    grouped: defaultdict[str, list[str]] = defaultdict(list)
    for value in values:
        grouped[_label_key(value)].append(value)
    for forms in grouped.values():
        if len(forms) > 1:
            listing = ", ".join(f"{form!r} ({values[form]})" for form in sorted(forms))
            report.add(
                WARNING, VALUES, f"{column}: spelling variants of one value: {listing}"
            )
    unknown: list[str] = []
    for value in sorted(values):
        count = values[value]
        if "\n" in value:
            report.add(
                WARNING,
                VALUES,
                f"{column}: multi-line value {value!r} ({count} rows) is not split and matches nothing",
            )
            continue
        if value.startswith(("http://", "https://")):
            report.add(
                WARNING,
                VALUES,
                f"{column}: URL {value!r} ({count} rows) in a value column",
            )
            continue
        if allowed is not None:
            if value not in allowed:
                report.add(
                    WARNING,
                    VALUES,
                    f"{column}: {value!r} ({count} rows) is outside the vocabulary {sorted(allowed)}",
                )
            continue
        if resolve is None:
            continue
        try:
            found = resolve(value)
        except TaxonomyError as exc:
            report.add(
                ERROR,
                VALUES,
                f"{column}: {value!r} is ambiguous in the taxonomy ({exc})",
            )
            continue
        if found is not None:
            continue
        suggestion = closest_label(value, labels)
        if suggestion:
            report.add(
                WARNING,
                VALUES,
                f"{column}: {value!r} ({count} rows) is not a taxonomy label; closest is {suggestion!r}",
            )
        else:
            unknown.append(f"{value!r} ({count})")
    if unknown:
        report.add(
            INFO,
            VALUES,
            f"{column}: {len(unknown)} values have no taxonomy concept and are dropped: {', '.join(unknown)}",
        )


def check_values(
    report: Report,
    rows: Sequence[Mapping[Any, Any]],
    columns: Sequence[str],
    importer: TaxonomyImport | None,
    labels: Mapping[str, str],
) -> None:
    """Check the value columns that map to taxonomy concepts or fixed vocabularies.

    The concept columns (event type, transport mode, place category) and the date
    label are checked against the dry-run importer. The date certainty and date
    precision columns are checked against their fixed vocabularies. Only columns
    present in ``columns`` are checked. Cells with leading or trailing spaces get
    one INFO line per column. Without an importer, the taxonomy lookups are skipped
    without a message, and only the vocabulary and text checks run.
    """
    for column in (*CONCEPT_COLUMNS, TIME_COLUMN, *ENUM_COLUMNS):
        if column not in columns:
            continue
        values = Counter(
            value for value in (cell(row, column) for row in rows) if value
        )
        spaced = sum(1 for row in rows if has_outer_spaces(row.get(column)))
        if spaced:
            report.add(
                INFO,
                VALUES,
                f"{column}: {spaced} cells have leading or trailing spaces",
            )
        allowed = ENUM_COLUMNS.get(column)
        resolve: Callable[[str], int | None] | None = None
        if allowed is None and importer is not None:
            resolve = (
                importer.resolve_time_period
                if column == TIME_COLUMN
                else importer.resolve_concept
            )
        check_column(
            report, column, values, allowed=allowed, resolve=resolve, labels=labels
        )


def check_fields(report: Report, rows: Sequence[Mapping[Any, Any]]) -> None:
    """Check dates, coordinates, and IDs, reporting unparseable values in aggregate.

    A date that ``parse_date`` cannot read costs the event its timespan, so the
    unparseable values are counted per column and reported in one WARNING per
    column. Within a row, a start after the end, a latitude without a longitude (or
    the reverse), and coordinates outside the valid range are each a WARNING. The
    rebuild ignores invalid coordinates. Wikidata values that are not Q-identifiers
    and GeoNames IDs that are not plain digits are counted the same way.
    """
    unparsed: dict[str, Counter[str]] = {column: Counter() for column in DATE_COLUMNS}
    qids: Counter[str] = Counter()
    geonames: Counter[str] = Counter()
    for row in rows:
        for column in DATE_COLUMNS:
            value = cell(row, column)
            if value and parse_date(value) is None:
                unparsed[column][value] += 1
        start_text = cell(row, "date_start")
        end_text = cell(row, "date_end")
        start = parse_date(start_text) if start_text else None
        end = parse_date(end_text) if end_text else None
        if start and end and start > end:
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: date_start {start} is after date_end {end}",
            )
        lat = cell(row, "place_lat")
        lon = cell(row, "place_lon")
        if bool(lat) != bool(lon):
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: only one of place_lat and place_lon is set",
            )
        elif lat and coordinate_pair(lat, lon) is None:
            report.add(
                WARNING,
                VALUES,
                f"{location_of(row)}: coordinates {lat!r}, {lon!r} are invalid and are ignored",
            )
        qid = cell(row, "wikidata_qid")
        if qid and not WIKIDATA_ID.fullmatch(qid):
            qids[qid] += 1
        geonames_id = cell(row, "geonames_id")
        if geonames_id and not geonames_id.isdigit():
            geonames[geonames_id] += 1
    for column, counter in unparsed.items():
        if counter:
            report.add(
                WARNING,
                VALUES,
                f"{column}: {sum(counter.values())} values cannot be parsed as dates, "
                f"so those events lose the timespan: {summarize(counter)}",
            )
    if qids:
        report.add(
            WARNING,
            VALUES,
            f"wikidata_qid: values that are not Q-identifiers: {summarize(qids)}",
        )
    if geonames:
        report.add(
            WARNING,
            VALUES,
            f"geonames_id: values that are not plain numeric IDs: {summarize(geonames)}",
        )


# --- Biographies -------------------------------------------------------------


def check_biographies(report: Report, chrono: ChronotopiData) -> None:
    """Run the biography planner over the chronotopi sheets and report its outcome.

    The planner checks each person's biographical facts against the metadata
    workbook for the same chronotopi workbook. It makes no API calls and writes
    nothing. Each conflict it finds is one ERROR, marked as not fixed. Any other
    planner error, such as a missing association, is one ERROR. On success, one
    INFO line gives the number of people with validated facts.
    """
    try:
        plan = build_person_biography_plan(chrono.overview, chrono.frame, SOURCE_DIR)
    except PersonBiographyConflictError as exc:
        for conflict in exc.conflicts:
            report.add(ERROR, BIOGRAPHIES, f"{conflict} (reported only; not fixed)")
        return
    except PersonBiographyError as exc:
        report.add(ERROR, BIOGRAPHIES, str(exc))
        return
    report.add(INFO, BIOGRAPHIES, f"{len(plan)} people have validated biography facts")


# --- Entry point -------------------------------------------------------------


def finish(report: Report, show_all: bool) -> int:
    print(report.render(None if show_all else DEFAULT_LIMIT))
    return 1 if report.count(ERROR) else 0


def main(argv: Sequence[str] | None = None) -> int:
    """Run every check, print the findings report, and return the exit status.

    ``--all`` prints every finding instead of capping each list. The MAXQDA
    taxonomy is loaded before the checks that need it. If it cannot be loaded,
    the run stops after the first four checks and reports that. Nothing is written.
    ``argv`` defaults to ``sys.argv[1:]``. Returns 1 when the report has at least
    one ERROR, and 0 otherwise.
    """
    parser = argparse.ArgumentParser(
        description="Report data inconsistencies without changing anything."
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="print every finding instead of capping each list",
    )
    args = parser.parse_args(argv)

    report = Report()
    workbooks = check_files(report)
    check_headers(report)
    check_rebuild_scope(report)
    check_taxonomy_copies(report)

    try:
        taxonomy = load_taxonomy(MAXQDA_TAXONOMY)
        concept_nodes = taxonomy.concept_nodes()
    except (OSError, TaxonomyError, ValueError) as exc:
        report.add(ERROR, TAXONOMY, f"{MAXQDA_TAXONOMY.name} cannot be loaded: {exc}")
        return finish(report, args.all)
    go_keys = {
        taxonomy_label_key(node.label)
        for node in concept_nodes
        if _label_key(_extract_key(node.top_level)) == _label_key("GO")
    }
    labels = {_label_key(node.label): node.label for node in concept_nodes}

    location = read_location_synonyms(report)
    person = read_person_synonyms(report)
    if location is not None:
        check_location_graph(report, location, go_keys)
    if person is not None:
        check_person_associations(report, workbooks, person)

    chrono = load_chronotopi(report, workbooks)
    rows = chrono.frame.to_dict("records")
    records = read_metadata_records(report)
    check_metadata_records(report, records)

    importer: TaxonomyImport | None = None
    augmentation = (
        augment_dry_run(report, taxonomy, chrono.overview)
        if person is not None
        else None
    )
    if augmentation is not None and person is not None:
        importer = dry_run_import(report, augmentation.taxonomy, augmentation.aliases)
        check_person_placement(
            report, chrono.overview, taxonomy, person, records, augmentation.aliases
        )
    check_places(report, importer, location)
    check_sheets(report, chrono, len(workbooks), importer)
    check_rows(report, rows)
    check_fields(report, rows)
    check_values(report, rows, list(chrono.frame.columns), importer, labels)
    check_biographies(report, chrono)
    return finish(report, args.all)


if __name__ == "__main__":
    sys.exit(main())
