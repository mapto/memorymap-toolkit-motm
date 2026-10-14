"""Tests for location coordinate resolution and workbook backfilling."""

import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest
import requests

import locations
from locations import db, entities, geo


HEADERS = [
    "location",
    "bag_of_words",
    "super_region",
    "lat",
    "long",
    "label",
    "www.geonames.org",
    "www.wikidata.org",
    "www.example.org",
]


def test_prepare_location_metadata_uses_custom_normalizer(monkeypatch):
    monkeypatch.setattr(
        entities,
        "resolve_location_coordinates",
        lambda entry: (50.0, 8.0),
    )

    metadata = locations._prepare_location_metadata(
        "A Place",
        locations_db={"custom-key": {"lat": None, "long": None}},
        normalize_fn=lambda value: "custom-key",
    )

    assert metadata["key"] == "custom-key"
    assert metadata["coordinates"] == (50.0, 8.0)


def test_canonical_location_name_maps_variants_to_workbook_entity():
    workbook = {
        "koblenz": {"name": "Koblenz"},
        "grosen linden": {"name": "Großen-Linden"},
    }

    assert (
        locations.canonical_location_name("Koblenz (Bundesarchiv)", workbook)
        == "Koblenz"
    )
    assert (
        locations.canonical_location_name("Großen-Linden", workbook) == "Großen-Linden"
    )
    assert locations.canonical_location_name("?", workbook) is None
    assert (
        locations.canonical_location_name(
            "Association of Israelis of Central European Origin (AICEO) (irgun-jeckes.org/familie-alsberg/)",
            workbook,
        )
        is None
    )


def test_canonical_location_name_matches_aliases_to_workbook_rows():
    workbook = {
        locations.normalize_location("Frankfurt am Main"): {
            "name": "Frankfurt am Main"
        },
        locations.normalize_location("Marktplatz 11"): {"name": "Marktplatz 11"},
    }

    assert (
        locations.canonical_location_name("Frankfurt_Main", workbook)
        == "Frankfurt am Main"
    )
    assert (
        locations.canonical_location_name("Gießen (Marktplatz 11)", workbook)
        == "Marktplatz 11"
    )
    assert locations.taxonomy_label_key("Frankfurt_Main") == "frankfurt main"
    assert locations.taxonomy_label_key("FRANKFURT-MAIN") == "frankfurt main"


def test_existing_location_updates_include_only_missing_values():
    metadata = {
        "coordinates": (50.0, 8.0),
        "wikidata_id": "Q1",
        "geonames_id": "123",
    }
    match = {"location": None, "wikidata_id": "Q2", "concept": 9}

    assert locations._existing_location_updates(metadata, match, 10) == {
        "latitude": 50.0,
        "longitude": 8.0,
        "geonames_id": "123",
    }


def test_new_location_payload_contains_taxonomy_concept():
    metadata = {
        "name": "Place",
        "coordinates": (50.0, 8.0),
        "wikidata_id": "Q1",
        "geonames_id": "123",
    }

    assert locations._new_location_payload(metadata, 10) == {
        "current_name": "Place",
        "latitude": 50.0,
        "longitude": 8.0,
        "wikidata_id": "Q1",
        "geonames_id": "123",
        "concept": 10,
    }


def write_locations(path, row):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(row)
    workbook.save(path)


def read_coordinates(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return sheet.cell(2, 4).value, sheet.cell(2, 5).value
    finally:
        workbook.close()


def test_load_locations_prefers_table_coordinates(tmp_path, monkeypatch):
    path = tmp_path / "locations.xlsx"
    write_locations(
        path,
        [
            "Berlin",
            "berlin",
            "",
            52.5,
            13.4,
            "",
            "https://www.geonames.org/2950159/berlin.html",
            "https://www.wikidata.org/wiki/Q64",
            "",
        ],
    )
    locations._locations_db.clear()
    monkeypatch.setattr(
        db,
        "resolve_location_coordinates",
        lambda url: (_ for _ in ()).throw(AssertionError("fallback should not run")),
    )

    database = locations.load_locations_db(path)

    entry = database[locations.normalize_location("Berlin")]
    assert (entry["lat"], entry["long"]) == (52.5, 13.4)
    assert read_coordinates(path) == (52.5, 13.4)


def test_load_locations_reads_coordinate_fields_by_header_order(tmp_path, monkeypatch):
    path = tmp_path / "locations.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "location",
            "bag_of_words",
            "super_region",
            "wider_region",
            "coordinate_source",
            "lat",
            "long",
            "label",
            "www.geonames.org",
            "www.wikidata.org",
        ]
    )
    sheet.append(
        [
            "Berlin",
            "berlin",
            "Deutschland",
            "GO > Deutschland > Berlin",
            "chronotopi",
            52.5,
            13.4,
            "",
            "",
            "https://www.wikidata.org/wiki/Q64",
        ]
    )
    workbook.save(path)

    locations._locations_db.clear()
    database = locations.load_locations_db(path, enrich_coordinates=False)

    entry = database[locations.normalize_location("Berlin")]
    assert entry["lat"] == 52.5
    assert entry["long"] == 13.4
    assert entry["wikidata_id"] == "Q64"
    assert entry["wider_region"] == "GO > Deutschland > Berlin"


def test_load_locations_collects_every_populated_column_after_geonames(tmp_path):
    path = tmp_path / "locations-extra-urls.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "location",
            "lat",
            "long",
            "www.geonames.org",
            "www.wikidata.org",
            "www.example.org",
            "archive.example",
        ]
    )
    sheet.append(
        [
            "Berlin",
            52.5,
            13.4,
            "https://www.geonames.org/2950159/berlin.html",
            "https://www.wikidata.org/wiki/Q64",
            "https://www.example.org/berlin",
            "https://archive.example/item/42",
        ]
    )
    workbook.save(path)

    locations._locations_db.clear()
    database = locations.load_locations_db(path, enrich_coordinates=False)
    entry = database[locations.normalize_location("Berlin")]

    assert entry["wikidata_id"] == "Q64"
    assert entry["extra_url"] == {
        "www.wikidata.org": "https://www.wikidata.org/wiki/Q64",
        "www.example.org": "https://www.example.org/berlin",
        "archive.example": "https://archive.example/item/42",
    }


def test_save_locations_preserves_all_extra_url_columns_by_header(tmp_path):
    path = tmp_path / "locations-save-extra-urls.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "location",
            "lat",
            "long",
            "www.geonames.org",
            "www.wikidata.org",
            "www.example.org",
        ]
    )
    sheet.append(["Berlin", 52.5, 13.4, "", "", ""])
    workbook.save(path)

    locations._locations_db.clear()
    database = locations.load_locations_db(path, enrich_coordinates=False)
    entry = database[locations.normalize_location("Berlin")]
    entry["extra_url"] = {
        "www.wikidata.org": "https://www.wikidata.org/wiki/Q64",
        "www.example.org": "https://www.example.org/berlin",
        "archive.example": "https://archive.example/item/42",
    }
    locations.save_locations_db(path)

    saved = load_workbook(path, read_only=True, data_only=True).active
    headers = [cell.value for cell in saved[1]]
    row = {
        header: saved.cell(2, index + 1).value for index, header in enumerate(headers)
    }
    assert row["www.wikidata.org"] == "https://www.wikidata.org/wiki/Q64"
    assert row["www.example.org"] == "https://www.example.org/berlin"
    assert row["archive.example"] == "https://archive.example/item/42"


def test_extract_geonames_id_uses_numeric_url_segment():
    assert (
        locations.extract_geonames_id(
            "https://www.geonames.org/3017382/republic-of-france.html"
        )
        == "3017382"
    )
    assert locations.extract_geonames_id("3017382") == "3017382"
    assert locations.extract_geonames_id("republic-of-france.html") is None


def test_cached_coordinates_skip_external_lookup_and_send_numeric_geonames_id(
    monkeypatch,
):
    locations._location_cache.clear()
    locations_db = {
        locations.normalize_location("Frankreich"): {
            "name": "Frankreich",
            "lat": 46,
            "long": 2,
            "geonames_id": "republic-of-france.html",
            "geonames_url": (
                "https://www.geonames.org/3017382/republic-of-france.html"
            ),
        }
    }
    posts = []
    monkeypatch.setattr(
        entities,
        "resolve_location_coordinates",
        lambda entry: (_ for _ in ()).throw(AssertionError("lookup should not run")),
    )
    monkeypatch.setattr("api_client.api_get", lambda endpoint, params=None: [])
    monkeypatch.setattr(
        "api_client.api_post",
        lambda endpoint, payload, fail_on_error=True: (
            posts.append(payload) or {"id": 42}
        ),
    )

    assert (
        locations.get_or_create_location("Frankreich", locations_db=locations_db) == 42
    )
    assert posts == [
        {
            "current_name": "Frankreich",
            "latitude": 46.0,
            "longitude": 2.0,
            "geonames_id": "3017382",
        }
    ]


def test_geonames_coordinates_fall_back_to_public_page(monkeypatch):
    class Response:
        def __init__(self, text="", error=None):
            self.text = text
            self.error = error

        def raise_for_status(self):
            if self.error:
                raise self.error

        def json(self):
            return {"status": {"value": 10}}

    calls = []

    def fake_get(url, **kwargs):
        calls.append(url)
        if "api.geonames.org" in url:
            return Response(error=requests.RequestException("unauthorized"))
        return Response(
            text='{"@type":"GeoCoordinates","latitude":52.52437,"longitude":13.41053}'
        )

    monkeypatch.setattr(requests, "get", fake_get)

    assert locations.extract_geonames_coordinates(
        "https://www.geonames.org/2950159/berlin.html"
    ) == {"lat": 52.52437, "long": 13.41053}
    assert calls == [
        "https://api.geonames.org/getJSON",
        "https://www.geonames.org/2950159/berlin.html",
    ]


def test_query_geocoder_parses_whereami_coordinate_response(monkeypatch):
    class Response:
        text = "50.59, 8.67"

        def raise_for_status(self):
            return None

    monkeypatch.setattr(time, "sleep", lambda seconds: None)
    monkeypatch.setattr(
        requests,
        "get",
        lambda url, **kwargs: Response(),
    )

    assert locations.query_geocoder("Gießen") == {"lat": 50.59, "long": 8.67}


def test_load_locations_falls_back_to_geonames_and_persists(tmp_path, monkeypatch):
    path = tmp_path / "locations.xlsx"
    write_locations(
        path,
        [
            "Gießen",
            "giessen",
            "",
            None,
            None,
            "",
            "https://www.geonames.org/2938913/giessen.html",
            "https://www.wikidata.org/wiki/Q2973",
            "",
        ],
    )
    locations._locations_db.clear()
    calls = []

    def geonames(url):
        calls.append("geonames")
        return {"lat": 50.59, "long": 8.67}

    monkeypatch.setattr(geo, "extract_geonames_coordinates", geonames)
    monkeypatch.setattr(
        geo,
        "extract_wikidata_coordinates",
        lambda url: (_ for _ in ()).throw(AssertionError("Wikidata should not run")),
    )

    database = locations.load_locations_db(path)

    entry = database[locations.normalize_location("Gießen")]
    assert (entry["lat"], entry["long"]) == (50.59, 8.67)
    assert calls == ["geonames"]
    assert read_coordinates(path) == (50.59, 8.67)


def test_load_locations_falls_back_to_wikidata_when_geonames_fails(
    tmp_path, monkeypatch
):
    path = tmp_path / "locations.xlsx"
    write_locations(
        path,
        [
            "Kibbuz Ein HaNetziv",
            "ein hanetziv kibbuz",
            "",
            None,
            None,
            "",
            "https://www.geonames.org/295211/Ein-HaNaziv.html",
            "https://www.wikidata.org/wiki/Q2095906",
            "",
        ],
    )
    locations._locations_db.clear()
    calls = []

    monkeypatch.setattr(
        geo,
        "extract_geonames_coordinates",
        lambda url: calls.append("geonames") or None,
    )
    monkeypatch.setattr(
        geo,
        "extract_wikidata_coordinates",
        lambda url: calls.append("wikidata") or {"lat": 32.45, "long": 35.51},
    )

    database = locations.load_locations_db(path)

    entry = database[locations.normalize_location("Kibbuz Ein HaNetziv")]
    assert (entry["lat"], entry["long"]) == (32.45, 35.51)
    assert calls == ["geonames", "wikidata"]
    assert read_coordinates(path) == (32.45, 35.51)


def test_resolve_location_uses_name_geocoder_without_source_ids(monkeypatch):
    monkeypatch.setattr(
        geo,
        "query_geocoder",
        lambda name: {"lat": 50.59, "long": 8.67},
    )
    monkeypatch.setattr(
        geo,
        "extract_geonames_coordinates",
        lambda url: (_ for _ in ()).throw(AssertionError("no source ID fallback")),
    )
    monkeypatch.setattr(
        geo,
        "extract_wikidata_coordinates",
        lambda url: (_ for _ in ()).throw(AssertionError("no source ID fallback")),
    )

    entry = {"name": "Gießen"}
    assert locations.resolve_location_coordinates(entry) == (50.59, 8.67)
    assert (entry["lat"], entry["long"]) == (50.59, 8.67)


def test_resolve_location_does_not_name_geocode_when_source_id_exists(monkeypatch):
    called = []
    monkeypatch.setattr(geo, "extract_geonames_coordinates", lambda url: None)
    monkeypatch.setattr(geo, "extract_wikidata_coordinates", lambda url: None)
    monkeypatch.setattr(
        geo,
        "query_geocoder",
        lambda name: called.append(name) or {"lat": 50.59, "long": 8.67},
    )

    entry = {
        "name": "Gießen",
        "geonames_id": "2938913",
        "geonames_url": "https://www.geonames.org/2938913/giessen.html",
    }
    assert locations.resolve_location_coordinates(entry) is None
    assert called == []


def test_resolve_location_logs_failed_name_geocoder(monkeypatch, caplog):
    monkeypatch.setattr(geo, "query_geocoder", lambda name: None)

    with caplog.at_level("WARNING", logger="locations"):
        assert locations.resolve_location_coordinates({"name": "Unknown place"}) is None

    assert (
        "Name geocoder returned no usable coordinates for Unknown place" in caplog.text
    )


def test_coordinate_lookup_logs_failure_and_continues_to_next_source(
    monkeypatch, caplog
):
    def failed_geonames(url):
        raise requests.RequestException("GeoNames unavailable")

    monkeypatch.setattr(geo, "extract_geonames_coordinates", failed_geonames)
    monkeypatch.setattr(
        geo,
        "extract_wikidata_coordinates",
        lambda url: {"lat": 32.45, "long": 35.51},
    )

    with caplog.at_level("WARNING", logger="locations"):
        coordinates = locations.resolve_location_coordinates(
            {
                "name": "Kibbuz Ein HaNetziv",
                "geonames_url": "https://www.geonames.org/295211",
                "wikidata_url": "https://www.wikidata.org/wiki/Q2095906",
            }
        )

    assert coordinates == (32.45, 35.51)
    assert "GeoNames unavailable" in caplog.text


def test_location_api_failure_is_logged_and_returns_none(monkeypatch, caplog):
    locations._location_cache.clear()
    monkeypatch.setattr(
        "api_client.api_get",
        lambda endpoint, params=None: (_ for _ in ()).throw(
            requests.RequestException("API unavailable")
        ),
    )

    with caplog.at_level("ERROR", logger="locations"):
        result = locations.get_or_create_location("Unavailable place")

    assert result is None
    assert "Location API lookup failed for Unavailable place" in caplog.text


def test_get_or_create_location_uses_normalized_db_key_and_repairs_existing_api_location(
    monkeypatch,
):
    locations._location_cache.clear()
    key = locations.normalize_location("Kibbuz Ein HaNetziv")
    locations_db = {
        key: {
            "name": "Kibbuz Ein HaNetziv",
            "lat": 32.45,
            "long": 35.51,
            "wikidata_id": "Q2095906",
            "geonames_id": "295211",
        }
    }
    patches = []
    monkeypatch.setattr(
        "api_client.api_get",
        lambda endpoint, params=None: [
            {"id": 7, "current_name": "Kibbuz Ein HaNetziv", "location": None}
        ],
    )
    monkeypatch.setattr(
        "api_client.api_patch",
        lambda endpoint, object_id, payload: patches.append(
            (endpoint, object_id, payload)
        ),
    )

    result = locations.get_or_create_location(
        "Kibbuz Ein HaNetziv", locations_db=locations_db
    )

    assert result == 7
    assert patches == [
        (
            "locations",
            7,
            {
                "latitude": 32.45,
                "longitude": 35.51,
                "wikidata_id": "Q2095906",
                "geonames_id": "295211",
            },
        )
    ]


def test_get_or_create_location_reuses_point_already_linked_to_concept(monkeypatch):
    locations._location_cache.clear()
    calls = []

    def api_get(endpoint, params=None):
        calls.append((endpoint, params))
        if params and params.get("search") == "Jerusalem":
            return []
        return [
            {
                "id": 15,
                "current_name": "Jerusalem (Hebrew University of Jerusalem)",
                "concept": 547,
            }
        ]

    monkeypatch.setattr("api_client.api_get", api_get)
    monkeypatch.setattr(
        "api_client.api_post",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("the concept-linked point should be reused")
        ),
    )

    result = locations.get_or_create_location("Jerusalem", concept_id=547)

    assert result == 15
    assert calls == [
        ("locations", {"search": "Jerusalem"}),
        ("locations", None),
    ]


def test_rebuild_location_workbook_delegates_to_shared_rebuild(monkeypatch):
    calls = []

    monkeypatch.setattr(
        "rebuild_locations.rebuild_workbook",
        lambda **kwargs: calls.append(kwargs) or {"rows": 1},
    )

    result = locations.rebuild_location_workbook(
        input_path="input.xlsx",
        output_path="output.xlsx",
        data_root="data",
        taxonomy_path="taxonomy.mmd",
    )

    assert result == {"rows": 1}
    assert calls == [
        {
            "input_path": Path("input.xlsx"),
            "output_path": Path("output.xlsx"),
            "data_root": Path("data"),
            "taxonomy_path": Path("taxonomy.mmd"),
        }
    ]


def test_sync_location_concepts_patches_unlinked_resolvable_points():
    locations = [
        {"id": 1, "current_name": "Gießen", "concept": None},
        {"id": 2, "current_name": "Koblenz", "concept": None},
        {"id": 3, "current_name": "Already linked", "concept": 9},
    ]
    patches = []

    result = __import__("locations").sync_location_concepts(
        lambda name: {"Gießen": 42}.get(name),
        api_get=lambda endpoint: locations,
        api_patch=lambda endpoint, object_id, payload: patches.append(
            (endpoint, object_id, payload)
        ),
    )

    assert result == {"matched": [(1, 42)], "unmapped": ["Koblenz"]}
    assert patches == [("locations", 1, {"concept": 42})]


def test_sync_location_concepts_enables_aliases_without_embedded_matching():
    locations = [{"id": 1, "current_name": "Frankfurt/M", "concept": None}]
    calls = []
    patches = []

    def resolve_location(name, **kwargs):
        calls.append((name, kwargs))
        return 42

    result = __import__("locations").sync_location_concepts(
        resolve_location,
        api_get=lambda endpoint: locations,
        api_patch=lambda endpoint, object_id, payload: patches.append(
            (endpoint, object_id, payload)
        ),
    )

    assert result == {"matched": [(1, 42)], "unmapped": []}
    assert calls == [
        (
            "Frankfurt/M",
            {"allow_aliases": True, "allow_embedded": False},
        )
    ]
    assert patches == [("locations", 1, {"concept": 42})]


def test_sync_location_concepts_does_not_steal_an_existing_unique_concept():
    locations = [
        {"id": 1, "current_name": "Gießen", "concept": None},
        {"id": 2, "current_name": "Gießen Marktplatz 15", "concept": 42},
    ]
    patches = []

    result = __import__("locations").sync_location_concepts(
        lambda name, **kwargs: 42,
        api_get=lambda endpoint: locations,
        api_patch=lambda endpoint, object_id, payload: patches.append(
            (endpoint, object_id, payload)
        ),
    )

    assert result == {"matched": [], "unmapped": ["Gießen"]}
    assert patches == []


def test_seed_location_entities_reuses_workbook_rows_and_resolves_concepts(monkeypatch):
    calls = []
    resolver_calls = []
    monkeypatch.setattr(
        "api_client.api_get",
        lambda endpoint: {"concept": 449} if endpoint == "locations/1" else [],
    )
    monkeypatch.setattr("api_client.api_patch", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        entities,
        "get_or_create_location",
        lambda name, **kwargs: calls.append((name, kwargs)) or len(calls),
    )
    workbook = {
        "berlin": {
            "name": "Berlin",
            "wikidata_id": "Q64",
            "geonames_id": "2950159",
        }
    }

    def resolve_location(name, **kwargs):
        resolver_calls.append((name, kwargs))
        return 449

    result = locations.seed_location_entities(
        workbook,
        resolve_location=resolve_location,
    )

    assert result == [("Berlin", 1, 449)]
    assert resolver_calls == [
        (
            "Berlin",
            {"allow_aliases": True, "allow_embedded": False},
        )
    ]
    assert calls == [
        (
            "Berlin",
            {
                "wikidata_qid": "Q64",
                "geonames_id": "2950159",
                "locations_db": workbook,
                "concept_id": 449,
            },
        )
    ]


def test_seed_location_entities_fails_before_patching_unmapped_canonical_rows(
    monkeypatch,
):
    monkeypatch.setattr(
        "api_client.api_get",
        lambda endpoint: (_ for _ in ()).throw(
            AssertionError("unmapped rows should fail before API writes")
        ),
    )
    workbook = {"unknown": {"name": "Unknown Canonical Place"}}

    with pytest.raises(RuntimeError, match="has no GO concept mapping"):
        locations.seed_location_entities(workbook, resolve_location=lambda name: None)
