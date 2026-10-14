"""Unit tests for metadata workbook parsing and pure helpers."""

import json

import pytest
from openpyxl import Workbook

import api_client
import metadati
import metadati.interviews
import metadati.people
import metadati.relationships
import person_taxonomy
import person_taxonomy.synonyms


def test_clean_value_returns_none_for_empty_values():
    assert metadati.clean_value(None) is None
    assert metadati.clean_value("  ") is None
    assert metadati.clean_value("value") == "value"


@pytest.mark.parametrize(
    ("full_name", "expected"),
    [
        ("Stern, Josef", ("Josef", "Stern")),
        ("Josef Stern", ("Josef", "Stern")),
        ("Josef", ("Josef", "UNKNOWN")),
    ],
)
def test_parse_name(full_name, expected):
    assert metadati.parse_name(full_name) == expected


@pytest.mark.parametrize(
    "source_label",
    [
        "Anne Betten, IS_E_00124",
        "Anne Betten + IS_E_00124",
    ],
)
def test_parse_source_label_extracts_interviewer_and_corpus_id(source_label):
    assert metadati.parse_source_label(source_label) == (
        "Anne Betten",
        "IS_E_00124",
    )


def test_read_person_record_starts_family_rows_at_row_four(tmp_path):
    workbook = Workbook()
    person_sheet = workbook.active
    person_sheet.title = "Scheda individuale"
    person_sheet.append(["Metadaten - P001"])
    person_sheet.append(["ID Sprecher", "P001"])
    person_sheet.append(["Vorname", "Josef"])
    person_sheet.append(["Nachname", "Stern"])

    family_sheet = workbook.create_sheet("Familie")
    family_sheet.append(["Family header"])
    family_sheet.append(["Column header"])
    family_sheet.append(["This row is not a family member", "Header"])
    family_sheet.append(
        [
            1,
            "Mutter",
            "Maria",
            "Stern",
            "12.03.1920",
            "Berlin",
            "Family note",
        ]
    )

    workbook_path = tmp_path / "metadati_P001.xlsx"
    workbook.save(workbook_path)

    record = metadati.read_person_record(workbook_path)

    assert record["identifier"] == "P001"
    assert record["person"]["given_name"] == "Josef"
    assert record["family"] == [
        {
            "relation": "Mutter",
            "given_name": "Maria",
            "family_name": "Stern",
            "birth_date": "1920-03-12",
            "birth_date_interval": ("1920-03-12", "1920-03-12"),
            "death_date": None,
            "death_date_interval": None,
            "birth_place": "Berlin",
            "notes": "Family note",
        }
    ]


def test_read_person_record_skips_empty_family_rows(tmp_path):
    workbook = Workbook()
    person_sheet = workbook.active
    person_sheet.title = "Scheda individuale"
    family_sheet = workbook.create_sheet("Familie")
    family_sheet.append(["Family header"])
    family_sheet.append(["Column header"])
    family_sheet.append(["Template row"])
    family_sheet.append([None] * 7)
    family_sheet.append([1, "Bruder", "Hans", "Stern", None, None, None])

    workbook_path = tmp_path / "metadati_P002.xlsx"
    workbook.save(workbook_path)

    record = metadati.read_person_record(workbook_path)

    assert len(record["family"]) == 1
    assert record["family"][0]["relation"] == "Bruder"


def test_create_interview_uses_shared_normalized_upsert(monkeypatch):
    calls = []
    monkeypatch.setattr(
        metadati.interviews,
        "get_or_create_person",
        lambda data: 8,
    )

    def fake_upsert(payloads):
        calls.extend(payloads)
        return {"IS_E_00170": 17}, []

    monkeypatch.setattr(metadati.interviews, "upsert_interviews", fake_upsert)
    monkeypatch.setattr(metadati.interviews, "get_or_create_url", lambda value: 31)

    interview_id = metadati.create_interview(
        7,
        {
            "archive_id": "IS_E_170",
            "type": "Interview (oral)",
            "date": "08.05.1991",
            "place": "Berlin",
            "interviewer": "Doe, Jane",
            "url": "https://example.test/interview/1",
        },
    )

    assert interview_id == 17
    assert calls == [
        {
            "archive_id": "IS_E_00170",
            "interviewee": 7,
            "interviewer": 8,
            "interview_type": "oral",
            "date": "1991-05-08",
            "place": "Berlin",
            "urls": [31],
        }
    ]


def test_read_person_record_extracts_source_url(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scheda individuale"
    sheet.append(["ID Sprecher", "P001"])
    sheet.append(["Quelle 1", None])
    sheet.append(["Quelle 1 – Bezeichnung", "Anne Betten, IS_E_00124"])
    sheet.append(["Quelle 1 – Ort", "Jerusalem"])
    sheet.append(["Quelle 1 – Handle/URL", "https://example.test/interview/1"])

    workbook_path = tmp_path / "metadati_P001.xlsx"
    workbook.save(workbook_path)

    record = metadati.read_person_record(workbook_path)

    assert record["interviews"] == [
        {
            "type": None,
            "place": "Jerusalem",
            "date": None,
            "interviewer": "Anne Betten",
            "archive_id": "IS_E_00124",
            "url": "https://example.test/interview/1",
        }
    ]


def test_read_person_record_extracts_place_metadata(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scheda individuale"
    sheet.append(
        [
            "Geburtsort",
            "Frankfurt an der Oder\nLand: Deutsches Reich | "
            "Geokoordinaten (Breite/Länge): [52.35 / 14.55] | "
            "Ortsname: Frankfurt an der Oder",
        ]
    )

    workbook_path = tmp_path / "metadati_place.xlsx"
    workbook.save(workbook_path)

    record = metadati.read_person_record(workbook_path)

    assert record["birth_place"] == {
        "name": "Frankfurt an der Oder",
        "coordinates": {"lat": 52.35, "lon": 14.55},
        "wikidata_id": None,
        "regions": ["Deutsches Reich"],
    }


def test_read_person_record_extracts_person_metadata(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scheda individuale"
    sheet.append(["Schulbildung / Abschluss", "Gymnasium\nmit Abschluss"])
    sheet.append(["Beruf (zum Zeitpunkt der Emigration)", "Schüler"])
    sheet.append(["Beruf (zum Zeitpunkt der Quelle)", "Bibliothekar; Historiker"])
    sheet.append(["Wohnort zum Zeitpunkt der Quelle", "Haifa"])

    workbook_path = tmp_path / "metadati_metadata.xlsx"
    workbook.save(workbook_path)

    record = metadati.read_person_record(workbook_path)

    assert record["person"] == {
        "given_name": None,
        "previous_given_name": None,
        "family_name": None,
        "previous_family_name": None,
        "gender": None,
        "birth_date": None,
        "birth_date_interval": None,
        "education": "Gymnasium mit Abschluss",
        "occupation_at_emigration": "Schüler",
        "occupation_at_source": "Bibliothekar; Historiker",
        "residence_at_source": "Haifa",
        "attributes": {},
    }


def test_create_interview_skips_missing_archive_id(monkeypatch):
    monkeypatch.setattr(
        metadati.interviews,
        "upsert_interviews",
        lambda payloads: pytest.fail("should not upsert without archive ID"),
    )

    assert metadati.create_interview(7, {"archive_id": None}) is None


def test_maybe_update_lifespan_keeps_existing_date_confirmed_by_partial_metadata(
    monkeypatch,
):
    responses = {
        f"{api_client.BASE_URL}/persons/7/": {"lifespan": 9},
        f"{api_client.BASE_URL}/timespans/9/": {
            "start": "1924-10-30",
            "end": None,
        },
    }

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def json(self):
            return self.payload

    monkeypatch.setattr(
        api_client.SESSION,
        "get",
        lambda url: Response(responses[url]),
    )
    monkeypatch.setattr(
        api_client.SESSION,
        "patch",
        lambda *args, **kwargs: pytest.fail("a compatible date should not be patched"),
    )

    metadati._maybe_update_lifespan(
        7,
        {
            "birth_date": "1924-01-01",
            "birth_date_interval": ("1924-01-01", "1924-12-31"),
        },
    )


def test_maybe_update_lifespan_rejects_a_conflicting_existing_date(monkeypatch):
    responses = {
        f"{api_client.BASE_URL}/persons/7/": {"lifespan": 9},
        f"{api_client.BASE_URL}/timespans/9/": {
            "start": "1923-01-01",
            "end": None,
        },
    }

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def json(self):
            return self.payload

    monkeypatch.setattr(
        api_client.SESSION,
        "get",
        lambda url: Response(responses[url]),
    )

    with pytest.raises(ValueError, match="conflicting metadata value"):
        metadati._maybe_update_lifespan(
            7,
            {
                "birth_date": "1924-10-30",
                "birth_date_interval": ("1924-10-30", "1924-10-30"),
            },
        )


def test_get_or_create_person_uses_reviewed_alias_resolution_with_identifier(
    monkeypatch,
):
    identity_calls = []
    patches = []
    existing = {
        "id": 20,
        "identifier": None,
        "given_name": "Johanna",
        "family_name": "Kaminka",
        "gender": "",
        "previous_given_name": "",
        "previous_family_name": "",
    }

    def fake_get_person_id(protagonist, name, identifier=None):
        identity_calls.append((protagonist, name, identifier))
        return 20

    monkeypatch.setattr(api_client, "get_person_id", fake_get_person_id)
    monkeypatch.setattr(
        metadati.people,
        "api_get",
        lambda endpoint, params=None: existing if endpoint == "persons/20" else [],
    )
    monkeypatch.setattr(
        metadati.people, "api_patch", lambda *args: patches.append(args)
    )
    monkeypatch.setattr(
        metadati.people,
        "_maybe_update_lifespan",
        lambda person_id, data: None,
    )
    metadati._person_cache.clear()

    person_id = metadati.get_or_create_person(
        {
            "given_name": "Johanna",
            "family_name": "Stern",
            "gender": "female",
            "occupation_at_source": "Researcher",
        },
        identifier="META-020",
    )

    assert person_id == 20
    assert identity_calls == [
        ("Johanna Stern", "Johanna Stern", "META-020"),
    ]
    assert patches == [
        (
            "persons",
            20,
            {
                "occupation_at_source": "Researcher",
                "gender": "female",
                "identifier": "META-020",
                "previous_family_name": "Stern",
            },
        ),
    ]


def test_create_relationship_links_unnamed_relative_and_merges_placeholder(
    tmp_path, monkeypatch
):
    synonyms_path = tmp_path / "person_synonyms.json"
    synonyms_path.write_text(
        json.dumps(
            {
                "metadata_relatives": {
                    "metadati_brodsky.xlsx": {
                        "Tante 1 (Schwester der Mutter)": "Brodsky, Tante 1",
                    }
                }
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        person_taxonomy.synonyms, "PERSON_TAXONOMY_SYNONYMS_PATH", synonyms_path
    )
    persons = {
        14: {"id": 14, "identifier": None, "given_name": "?", "family_name": "?"},
        30: {
            "id": 30,
            "identifier": None,
            "given_name": "Maria",
            "family_name": "Brodsky",
        },
    }
    relationships = [
        {"id": 11, "person_from": 5, "person_to": 14, "relationship_type": 7},
        {"id": 12, "person_from": 5, "person_to": 15, "relationship_type": 8},
        {"id": 13, "person_from": 5, "person_to": 30, "relationship_type": 7},
        {"id": 16, "person_from": 6, "person_to": 16, "relationship_type": 7},
    ]
    calls = []

    def fake_get(endpoint, params=None):
        if endpoint == "relationships":
            return [dict(row) for row in relationships]
        return dict(persons[int(endpoint.split("/")[1])])

    def fake_patch(endpoint, obj_id, payload):
        calls.append(("patch", endpoint, obj_id, payload))
        persons[obj_id].update(payload)

    def fake_post(endpoint, payload, fail_on_error=True):
        calls.append(("post", endpoint, payload))
        assert endpoint == "persons/49/merge"
        duplicate_ids = payload["duplicate_ids"]
        for row in relationships:
            if row["person_to"] in duplicate_ids:
                row["person_to"] = 49
        for duplicate_id in duplicate_ids:
            del persons[duplicate_id]
        return {"canonical_id": 49, "merged_ids": duplicate_ids}

    def fake_get_or_create_person(data, identifier=None):
        calls.append(("person", data["given_name"], data["family_name"]))
        return 49

    monkeypatch.setattr(metadati.relationships, "_relationship_cache", {})
    monkeypatch.setattr(
        metadati.relationships, "get_or_create_person", fake_get_or_create_person
    )
    monkeypatch.setattr(
        metadati.relationships, "get_or_create_relationship_type", lambda label: 7
    )
    monkeypatch.setattr(metadati.relationships, "api_get", fake_get)
    monkeypatch.setattr(metadati.relationships, "api_patch", fake_patch)
    monkeypatch.setattr(metadati.relationships, "api_post", fake_post)

    relationship_id = metadati.create_relationship(
        5,
        {
            "relation": "Tante 1 (Schwester der Mutter)",
            "given_name": "?",
            "family_name": "?",
        },
        "Brodsky",
        source_workbook="metadati_brodsky.xlsx",
    )

    assert relationship_id == 11
    assert calls == [
        ("person", "Tante 1", "Brodsky"),
        ("patch", "persons", 14, {"given_name": "Tante 1", "family_name": "Brodsky"}),
        ("post", "persons/49/merge", {"duplicate_ids": [14]}),
    ]
    assert [(row["id"], row["person_to"]) for row in relationships] == [
        (11, 49),
        (12, 15),
        (13, 30),
        (16, 16),
    ]


def test_create_relationship_keeps_unlinked_placeholder_unmerged(monkeypatch):
    calls = []

    def fake_get_or_create_person(data, identifier=None):
        calls.append(("person", data["given_name"], data["family_name"]))
        return 14

    def fake_post(endpoint, payload, fail_on_error=True):
        calls.append(("post", endpoint, payload))
        return {"id": 21}

    monkeypatch.setattr(metadati.relationships, "_relationship_cache", {})
    monkeypatch.setattr(
        metadati.relationships, "get_or_create_person", fake_get_or_create_person
    )
    monkeypatch.setattr(
        metadati.relationships, "get_or_create_relationship_type", lambda label: 7
    )
    monkeypatch.setattr(
        metadati.relationships, "api_get", lambda endpoint, params=None: []
    )
    monkeypatch.setattr(metadati.relationships, "api_post", fake_post)

    relationship_id = metadati.create_relationship(
        5,
        {
            "relation": "Tante 1 (Schwester der Mutter)",
            "given_name": "?",
            "family_name": "?",
        },
        "Brodsky",
        source_workbook="metadati_unlinked.xlsx",
    )

    assert relationship_id == 21
    assert calls == [
        ("person", "?", "?"),
        (
            "post",
            "relationships",
            {
                "relationship_type": 7,
                "person_from": 5,
                "person_to": 14,
                "description": "",
            },
        ),
    ]


def test_import_all_metadati_imports_workbooks_and_skips_templates(
    tmp_path, monkeypatch
):
    """Import each metadati workbook in a folder, skipping templates."""
    (tmp_path / "metadati_stern.xlsx").touch()
    (tmp_path / "metadati_template.xlsx").touch()
    (tmp_path / "chronotopi_stern.xlsx").touch()
    read_names = []

    def fake_read_person_record(xlsx_path):
        read_names.append(xlsx_path.name)
        return {
            "identifier": None,
            "person": {},
            "birth_place": {},
            "interviews": [],
            "family": [],
        }

    class Response:
        def json(self):
            return {"family_name": "Stern"}

    class FakeSession:
        def get(self, url):
            return Response()

    monkeypatch.setattr(
        metadati.importer, "read_person_record", fake_read_person_record
    )
    monkeypatch.setattr(
        metadati.importer,
        "get_or_create_person",
        lambda data, identifier=None: 5,
    )
    monkeypatch.setattr(
        metadati.importer, "get_or_create_location", lambda birth_place: None
    )
    monkeypatch.setattr(metadati.importer, "SESSION", FakeSession())

    results = metadati.import_all_metadati(directory=tmp_path)

    assert results == [5]
    assert read_names == ["metadati_stern.xlsx"]


def test_get_or_create_person_creates_person_with_identifier_and_lifespan(
    monkeypatch,
):
    """Create a new Person with its identifier and lifespan when none matches."""
    posts = []
    timespan_requests = []

    def fake_get_or_create_timespan(start=None, end=None):
        timespan_requests.append((start, end))
        return 3

    def fake_post(endpoint, payload, fail_on_error=True):
        posts.append((endpoint, payload))
        return {"id": 40}

    monkeypatch.setattr(metadati.people, "_person_cache", {})
    monkeypatch.setattr(
        metadati.people, "canonical_person_name", lambda given, family: None
    )
    monkeypatch.setattr(metadati.people, "person_taxonomy_alias", lambda value: None)
    monkeypatch.setattr(metadati.people, "api_get", lambda endpoint, params=None: [])
    monkeypatch.setattr(
        metadati.people,
        "api_patch",
        lambda *args: pytest.fail("a new person should not be patched"),
    )
    monkeypatch.setattr(
        metadati.people, "get_or_create_timespan", fake_get_or_create_timespan
    )
    monkeypatch.setattr(metadati.people, "api_post", fake_post)

    person_id = metadati.get_or_create_person(
        {
            "given_name": "Josef",
            "family_name": "Stern",
            "gender": "male",
            "birth_date": "1900-01-01",
            "death_date": "1970-01-01",
            "education": "Gymnasium",
        },
        identifier="META-040",
    )

    assert person_id == 40
    assert timespan_requests == [("1900-01-01", "1970-01-01")]
    assert posts == [
        (
            "persons",
            {
                "identifier": "META-040",
                "given_name": "Josef",
                "family_name": "Stern",
                "education": "Gymnasium",
                "gender": "male",
                "lifespan": 3,
            },
        ),
    ]


def test_get_or_create_location_links_region_chain(monkeypatch):
    """Link the innermost region of a birth place to its location."""
    calls = []
    region_ids = {"Deutsches Reich": 5, "Brandenburg": 6}

    def fake_get(endpoint, params=None):
        calls.append(("get", endpoint, params))
        return []

    def fake_post(endpoint, payload, fail_on_error=True):
        calls.append(("post", endpoint, payload))
        return {"id": region_ids[payload["name"]]}

    def fake_patch(endpoint, obj_id, payload):
        calls.append(("patch", endpoint, obj_id, payload))

    monkeypatch.setattr(metadati.places, "_location_workbook_cache", None)
    monkeypatch.setattr(
        metadati.places,
        "load_locations_db",
        lambda enrich_coordinates=True: {},
    )
    monkeypatch.setattr(
        metadati.places,
        "get_or_create_location_point",
        lambda name, wikidata_qid=None, locations_db=None: 77,
    )
    monkeypatch.setattr(metadati.places, "api_get", fake_get)
    monkeypatch.setattr(metadati.places, "api_post", fake_post)
    monkeypatch.setattr(metadati.places, "api_patch", fake_patch)

    location_id = metadati.get_or_create_location(
        {
            "name": "Frankfurt an der Oder",
            "wikidata_id": None,
            "regions": ["Deutsches Reich", "Brandenburg"],
        }
    )

    assert location_id == 77
    assert calls == [
        ("get", "regions", {"search": "Deutsches Reich"}),
        ("post", "regions", {"name": "Deutsches Reich"}),
        ("get", "regions", {"search": "Brandenburg"}),
        ("post", "regions", {"name": "Brandenburg", "part_of": 5}),
        ("patch", "locations", 77, {"region": 6}),
    ]


def test_maybe_update_person_place_sets_an_empty_birth_place(monkeypatch):
    """Set the birth place of a Person that has none, with a single PATCH."""
    patches = []

    def fake_patch(endpoint, obj_id, payload):
        patches.append((endpoint, obj_id, payload))

    monkeypatch.setattr(
        metadati.places, "api_get", lambda endpoint, params=None: {"id": 7}
    )
    monkeypatch.setattr(metadati.places, "api_patch", fake_patch)

    metadati._maybe_update_person_place(7, "birth_place", 77, "Frankfurt an der Oder")

    assert patches == [("persons", 7, {"birth_place": 77})]


def test_maybe_update_lifespan_fills_missing_lifespan_with_new_timespan(monkeypatch):
    """Create and link a Timespan when a Person has no lifespan yet."""
    searches = []
    posts = []
    patches = []

    class Response:
        def json(self):
            return {"lifespan": None}

    def fake_get(endpoint, params=None):
        searches.append(endpoint)
        return []

    def fake_post(endpoint, payload, fail_on_error=True):
        posts.append((endpoint, payload))
        return {"id": 9}

    monkeypatch.setattr(api_client.SESSION, "get", lambda url: Response())
    monkeypatch.setattr(
        api_client.SESSION,
        "patch",
        lambda url, json: patches.append((url, json)),
    )
    monkeypatch.setattr(metadati.timespans, "_timespan_cache", {})
    monkeypatch.setattr(metadati.timespans, "api_get", fake_get)
    monkeypatch.setattr(metadati.timespans, "api_post", fake_post)

    metadati._maybe_update_lifespan(7, {"birth_date": "1924-10-30", "death_date": None})

    assert searches == ["timespans"]
    assert posts == [("timespans", {"start": "1924-10-30", "end": None})]
    assert patches == [(f"{api_client.BASE_URL}/persons/7/", {"lifespan": 9})]
