"""Rebuild the location workbook from chronotopi and geospatial sources.

The rebuild is deliberately separate from the API import path. It produces a
stable workbook that can be consumed by ``locations.load_locations_db`` and
keeps the source provenance in ``coordinate_source`` while retaining every
existing workbook column. Geography is derived from the imported GO taxonomy;
addresses and descriptive source rows use their city context as the nearest
parent.

The chronotopi notebook calls ``locations.rebuild_location_workbook()``, which
delegates here, before it creates taxonomy concepts. Geocoding is slow and rate
limited, so it is done in this step, ahead of the API import. The workbook can
be inspected before anything is sent to the API.

Inputs are the existing ``locations.xlsx``, the ``*chronotop*.xlsx`` and
``metadati*.xlsx`` workbooks under ``data_root`` (by default
``data/nuove schede``), and the MAXQDA code system. The output is one sheet
named ``locations``. Rows without coordinates are never written; the rebuild
raises ``RuntimeError`` instead. The provider order is documented on
``resolve_coordinates``.

Usage, from the ``import`` folder::

    uv run python rebuild_locations.py

The default output path is the input path, so a run overwrites
``locations.xlsx``. Copy the workbook first. The script prints the summary
returned by ``rebuild_workbook``.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
import time
from typing import Callable

import openpyxl
import pandas as pd
import requests
from unidecode import unidecode

from locations import (
    LOCATIONS_WORKBOOK_PATH,
    extract_geonames_coordinates,
    extract_geonames_id,
    extract_wikidata_coordinates,
    normalize_location,
    query_geocoder,
    taxonomy_alias,
    taxonomy_location_parent,
    taxonomy_match_excluded,
    taxonomy_query_override,
    taxonomy_reviewed_unmapped,
    tokenize_location,
)
from taxonomy import Taxonomy, _extract_key, _label_key, load_taxonomy
from utils import extract_id_from_url, extract_url_map


DEFAULT_INPUT = LOCATIONS_WORKBOOK_PATH
DEFAULT_OUTPUT = DEFAULT_INPUT
DEFAULT_DATA_ROOT = Path(__file__).resolve().parents[1] / "data"
DEFAULT_TAXONOMY = DEFAULT_DATA_ROOT / "maxqda" / "MAXQDA_Code_System.mmd"
DEFAULT_WORKBOOKS_ROOT = DEFAULT_DATA_ROOT / "nuove schede"

BASE_COLUMNS = (
    "location",
    "bag_of_words",
    "super_region",
    "lat",
    "long",
    "label",
    "www.geonames.org",
    "www.wikidata.org",
)

DERIVED_COLUMNS = ("wider_region", "coordinate_source")


@dataclass
class SourceLocation:
    """Hold merged source evidence for one normalized location name.

    One instance is built per ``location_key`` by the ``read_*_sources``
    functions and combined by ``merge_source_maps``. Coordinates, GeoNames and
    Wikidata IDs, place details, and categories are collected in lists or sets,
    so repeated mentions accumulate. URLs are keyed by host. ``name`` is the
    display name from the first row seen. ``resolve_coordinates`` and
    ``wider_region_for`` use this evidence to place the location.
    """

    name: str
    direct_coordinates: list[tuple[float, float]] = field(default_factory=list)
    geonames_ids: set[str] = field(default_factory=set)
    wikidata_ids: set[str] = field(default_factory=set)
    details: set[str] = field(default_factory=set)
    categories: set[str] = field(default_factory=set)
    urls: dict[str, str] = field(default_factory=dict)


def clean_value(value) -> str:
    value = str(value or "").strip()
    return "" if value.lower() in {"nan", "none"} else value


def latinize_text(value: str) -> str:
    """Transliterate a region label so the exported value uses Latin text.

    Used to write ``super_region`` and to compare labels in ``_geo_path_for``.
    ``wider_region`` keeps the unmodified taxonomy path. Blank, ``nan``, and
    ``none`` inputs return an empty string.
    """
    return " ".join(unidecode(clean_value(value)).split())


def location_key(value: str) -> str:
    """Normalize punctuation and uncertain suffixes for source merging.

    The key identifies one location across workbook rows and source evidence.
    It drops a trailing ``?``, applies the reviewed aliases from
    ``taxonomy_alias``, and then uses ``normalize_location`` to lowercase, strip
    punctuation, and sort the words. Names that differ only in those respects,
    or by a reviewed alias, therefore share one key.
    """
    cleaned = clean_value(value).rstrip("?").strip()
    cleaned = taxonomy_alias(cleaned)
    return normalize_location(cleaned)


def coordinate_pair(latitude, longitude) -> tuple[float, float] | None:
    """Return a validated ``(lat, lon)`` float pair, or None when invalid.

    Accepts numbers or numeric strings, because spreadsheet cells and geocoder
    responses supply either. Non-numeric values, NaN, and coordinates outside
    -90 to 90 (latitude) or -180 to 180 (longitude) return None, so invalid
    values never reach the workbook. ``validate_data.py`` uses this check to
    report bad coordinates in the existing workbook.
    """
    try:
        latitude = float(latitude)
        longitude = float(longitude)
    except (TypeError, ValueError):
        return None
    if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
        return None
    return latitude, longitude


def read_existing_workbook(path: Path) -> tuple[list[str], dict[str, dict]]:
    """Read meaningful existing rows while retaining all workbook columns.

    Opens the active sheet of ``path`` read-only, with cached cell values.
    Rows with no location name, rows whose location cell is a URL, and rows
    naming a reviewed non-location are skipped. The rest are keyed by
    ``location_key``, so curated values and URL columns survive when
    ``merge_sources`` combines them with source evidence. If two rows share a
    key, the later row replaces the earlier one. ``validate_data.py`` also uses
    this function to list the names already in the workbook.

    Returns:
        The cleaned header names in column order, and a mapping from location
        key to the full row, with each header mapped to its cell value.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [
        clean_value(sheet.cell(1, column).value)
        for column in range(1, sheet.max_column + 1)
    ]
    records = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = {header: value for header, value in zip(headers, values)}
        name = clean_value(record.get("location"))
        if not name or name.startswith("http") or taxonomy_reviewed_unmapped(name):
            continue
        records[location_key(name)] = record
    return headers, records


def _header_row(path: Path, sheet_name: str) -> int | None:
    """Return the zero-based row index of the ``event_label`` header, if any.

    Chronotopi sheets have title rows above the header. Instead of a per-sheet
    skip count, such as the skip table in ``chronotopi.ipynb``, this searches
    the first 12 rows for an ``event_label`` cell. Returns None when no row
    matches, and the caller then skips the sheet.
    """
    preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=12)
    for row_number, (_, row) in enumerate(preview.iterrows()):
        values = {clean_value(value) for value in row if pd.notna(value)}
        if "event_label" in values:
            return row_number
    return None


def _add_url(urls: dict[str, str], value: str) -> None:
    for host, url in extract_url_map(value)[0].items():
        if host and url:
            urls[host] = url


def read_chronotopi_sources(data_root: Path) -> dict[str, SourceLocation]:
    """Collect coordinates, identifiers, context, and URLs from chronotopi.

    Chronotopi rows hold the coordinates that ``resolve_coordinates`` tries
    first. Every ``*chronotop*.xlsx`` under ``data_root`` is read in sorted
    order. Sheets whose names start with an underscore are skipped, as are
    sheets without an ``event_label`` header or a ``place_name`` column. Each
    place row adds to the entry for its ``location_key``. Reviewed
    non-locations are ignored, and a workbook that cannot be opened is skipped
    without a message. Returns the entries keyed by ``location_key``.
    """
    sources: dict[str, SourceLocation] = {}
    for path in sorted(data_root.rglob("*chronotop*.xlsx")):
        try:
            workbook = pd.ExcelFile(path)
        except Exception:
            continue
        for sheet_name in workbook.sheet_names:
            if sheet_name.startswith("_"):
                continue
            header = _header_row(path, sheet_name)
            if header is None:
                continue
            frame = pd.read_excel(path, sheet_name=sheet_name, header=header)
            if "place_name" not in frame.columns:
                continue
            for _, row in frame.iterrows():
                name = clean_value(row.get("place_name"))
                if not name or taxonomy_reviewed_unmapped(name):
                    continue
                key = location_key(name)
                source = sources.setdefault(
                    key, SourceLocation(name=name.rstrip("?").strip())
                )
                coordinates = coordinate_pair(
                    clean_value(row.get("place_lat")),
                    clean_value(row.get("place_lon")),
                )
                if coordinates:
                    source.direct_coordinates.append(coordinates)
                geonames_id = extract_geonames_id(clean_value(row.get("geonames_id")))
                if geonames_id:
                    source.geonames_ids.add(geonames_id)
                wikidata_id = clean_value(row.get("wikidata_qid")) or None
                if wikidata_id:
                    source.wikidata_ids.add(wikidata_id)
                detail = clean_value(row.get("place_detail"))
                if detail:
                    source.details.add(detail.rstrip("?").strip())
                category = clean_value(row.get("place_category"))
                if category:
                    source.categories.update(
                        part.strip() for part in category.splitlines() if part.strip()
                    )
                _add_url(source.urls, clean_value(row.get("external_links")))
    return sources


def read_metadata_sources(data_root: Path) -> dict[str, SourceLocation]:
    """Collect birth-place evidence from all metadata workbooks.

    Each ``metadati*.xlsx`` under ``data_root`` is parsed with
    ``metadati.read_person_record``. The birth place name keys the entry, and
    its coordinates, Wikidata ID, and regions are merged into it as direct
    coordinates, identifiers, and details. Files that fail to parse, records
    without a birth place name, and reviewed non-locations are skipped
    silently. Returns the entries keyed by ``location_key``.
    """
    from metadati import read_person_record

    sources: dict[str, SourceLocation] = {}
    for path in sorted(data_root.rglob("metadati*.xlsx")):
        try:
            record = read_person_record(path)
        except Exception:
            continue
        birth_place = record.get("birth_place") or {}
        name = clean_value(birth_place.get("name"))
        if not name or taxonomy_reviewed_unmapped(name):
            continue
        key = location_key(name)
        source = sources.setdefault(key, SourceLocation(name=name))
        coordinates = birth_place.get("coordinates") or {}
        pair = coordinate_pair(coordinates.get("lat"), coordinates.get("lon"))
        if pair:
            source.direct_coordinates.append(pair)
        wikidata_id = clean_value(birth_place.get("wikidata_id"))
        if wikidata_id:
            source.wikidata_ids.add(wikidata_id)
        source.details.update(
            clean_value(region)
            for region in birth_place.get("regions", [])
            if clean_value(region)
        )
    return sources


def merge_source_maps(
    target: dict[str, SourceLocation], extra: dict[str, SourceLocation]
) -> None:
    """Merge source evidence from another import source into the location map.

    ``target`` is changed in place, and nothing is returned. A key that is new
    to ``target`` takes the ``SourceLocation`` from ``extra`` without copying
    it, so ``extra`` should not be reused. For a key already present, coordinates,
    identifiers, details, and categories are added, and a URL host from
    ``extra`` replaces the same host in ``target``. ``rebuild_workbook`` uses
    this to add metadati birth places to the chronotopi evidence.
    """
    for key, source in extra.items():
        if key not in target:
            target[key] = source
            continue
        current = target[key]
        current.direct_coordinates.extend(source.direct_coordinates)
        current.geonames_ids.update(source.geonames_ids)
        current.wikidata_ids.update(source.wikidata_ids)
        current.details.update(source.details)
        current.categories.update(source.categories)
        current.urls.update(source.urls)


def merge_sources(
    headers: list[str],
    existing: dict[str, dict],
    sources: dict[str, SourceLocation],
) -> dict[str, dict]:
    """Union existing workbook metadata with chronotopi source evidence.

    Returns a new mapping from location key to row dict; the input rows are
    copied, not changed. A source with no workbook row gets one, with every
    header set to an empty string. Source evidence only fills blanks, such as
    a missing GeoNames or Wikidata URL, an empty URL column, or a missing
    name. A label that is empty or a bare URL is replaced by the source
    categories.
    """
    merged = {key: dict(record) for key, record in existing.items()}
    for key, source in sources.items():
        record = merged.setdefault(key, {header: "" for header in headers})
        record.setdefault("location", source.name)
        if not clean_value(record.get("location")):
            record["location"] = source.name
        if source.geonames_ids and not clean_value(record.get("www.geonames.org")):
            record["www.geonames.org"] = (
                f"https://www.geonames.org/{sorted(source.geonames_ids)[0]}"
            )
        if source.wikidata_ids and not clean_value(record.get("www.wikidata.org")):
            record["www.wikidata.org"] = (
                f"https://www.wikidata.org/wiki/{sorted(source.wikidata_ids)[0]}"
            )
        for host, url in source.urls.items():
            if host not in record or not clean_value(record.get(host)):
                record[host] = url
        if source.categories:
            current_label = clean_value(record.get("label"))
            if not current_label or current_label.startswith("http"):
                record["label"] = ",".join(sorted(source.categories))
    return merged


def extend_headers(headers: list[str], sources: dict[str, SourceLocation]) -> list[str]:
    """Retain existing columns and append new URL hosts found in chronotopi.

    Existing columns keep their positions. Each URL host in ``sources`` that is
    not yet a column is appended, in order of first appearance. A new list is
    returned, and ``headers`` is not changed. ``rebuild_workbook`` calls this
    before merging rows, so every URL host gets a column.
    """
    extended = list(headers)
    for source in sources.values():
        for host in source.urls:
            if host not in extended:
                extended.append(host)
    return extended


def order_output_headers(headers: list[str]) -> list[str]:
    """Place derived geography columns directly after ``super_region``.

    ``wider_region`` and ``coordinate_source`` are removed from their current
    positions and inserted right after ``super_region``, which keeps the
    geography columns together. Other columns keep their order. Raises
    ``ValueError`` when ``super_region`` is not among ``headers``.
    """
    ordered = [header for header in headers if header not in DERIVED_COLUMNS]
    insert_at = ordered.index("super_region") + 1
    return ordered[:insert_at] + list(DERIVED_COLUMNS) + ordered[insert_at:]


def build_geo_paths(taxonomy: Taxonomy) -> dict[str, tuple[str, ...]]:
    """Index GO taxonomy labels by their broad-to-specific source path.

    Only occurrences whose top-level category is GO are indexed. Keys are the
    labels normalized by ``_label_key``, which ignores case, dashes, and
    underscores. Each value is the occurrence path, from the GO root down to
    the label itself. A repeated normalized label replaces the earlier entry.
    """
    paths = {}
    for occurrence in taxonomy.occurrences:
        if _label_key(_extract_key(occurrence.top_level)) != "go":
            continue
        paths[_label_key(occurrence.label)] = occurrence.path
    return paths


def choose_direct_coordinates(source: SourceLocation) -> tuple[float, float] | None:
    """Pick the most frequent direct coordinate pair for a location.

    Pairs are rounded to six decimal places before they are counted, so float
    noise compares equal. Taking the most frequent pair keeps one outlying row
    from deciding the position. Returns None when the source has no direct
    coordinates.
    """
    if not source.direct_coordinates:
        return None
    rounded = Counter(
        (round(lat, 6), round(lon, 6)) for lat, lon in source.direct_coordinates
    )
    return rounded.most_common(1)[0][0]


def source_ids(
    record: dict, source: SourceLocation | None
) -> tuple[str | None, str | None]:
    """Return the GeoNames and Wikidata IDs to try for a location.

    IDs already in the workbook's ``www.geonames.org`` and ``www.wikidata.org``
    columns come first. When the workbook has none, the first ID in sorted
    order from the source evidence is used. Either element may be None. The
    IDs feed the lookups in ``resolve_coordinates``.
    """
    geonames_url = clean_value(record.get("www.geonames.org"))
    wikidata_url = clean_value(record.get("www.wikidata.org"))
    geonames_id = extract_geonames_id(geonames_url)
    wikidata_id = extract_id_from_url(wikidata_url)
    if source:
        geonames_id = geonames_id or next(iter(sorted(source.geonames_ids)), None)
        wikidata_id = wikidata_id or next(iter(sorted(source.wikidata_ids)), None)
    return geonames_id, wikidata_id


def build_query(
    name: str,
    source: SourceLocation | None,
    record: dict,
    geo_paths: dict[str, tuple[str, ...]] | None = None,
) -> str:
    """Build the geocoder query for a location, adding context when available.

    A reviewed query override for the exact label takes precedence. When
    ``geo_paths`` is given, a name that is already a GO label is returned
    unchanged. Otherwise the first source detail in sorted order is used as
    context if it is under 80 characters and is not a URL. If it is not, the
    workbook's ``super_region`` is used unless it is a GO root label. The
    context is appended with ``Germany`` to help the geocoder choose the place.
    Without context, the bare name is returned.
    """
    override = taxonomy_query_override(name)
    if override is not None:
        return override
    if geo_paths:
        normalized_name = _label_key(name)
        if normalized_name in geo_paths:
            return name
    context = (
        next(iter(sorted(source.details)), "") if source and source.details else ""
    )
    if context and len(context) < 80 and not context.startswith("http"):
        return f"{name}, {context}, Germany"
    existing_parent = clean_value(record.get("super_region"))
    if existing_parent and not existing_parent.startswith("GO –"):
        return f"{name}, {existing_parent}, Germany"
    return name


def geocode_nominatim(
    query: str, *, request_get: Callable = requests.get, delay: float = 1.1
):
    """Geocode a query with Nominatim and return coordinates and address parts.

    This is the external OpenStreetMap geocoder. ``resolve_coordinates`` calls
    it when the local geocoder returns no usable coordinates. Each call makes
    one HTTP request with an identifying User-Agent and a 20-second timeout.
    HTTP errors are not caught, so they stop the rebuild. The function sleeps
    ``delay`` seconds after a non-empty response, and not after an empty one.
    ``request_get`` can be replaced in tests.

    Returns:
        ``{"coordinates": (lat, lon), "address": {...}}``, or None when there
        is no result or the result has no valid coordinates.
    """
    response = request_get(
        "https://nominatim.openstreetmap.org/search",
        params={"format": "jsonv2", "addressdetails": 1, "limit": 1, "q": query},
        headers={"User-Agent": "memorymap-toolkit/1.0"},
        timeout=20,
    )
    response.raise_for_status()
    results = response.json()
    if not results:
        return None
    time.sleep(delay)
    result = results[0]
    coordinates = coordinate_pair(result.get("lat"), result.get("lon"))
    if not coordinates:
        return None
    return {"coordinates": coordinates, "address": result.get("address", {})}


def reverse_geocode_nominatim(
    coordinates: tuple[float, float],
    *,
    request_get: Callable = requests.get,
    delay: float = 1.1,
):
    """Reverse-geocode coordinates to recover city/country containment.

    ``rebuild_workbook`` calls this only when the containment for a row is still
    empty after the taxonomy, context, and forward address checks. It returns
    the Nominatim ``address`` dictionary, such as city, state, or country. The
    function sleeps ``delay`` seconds after each successful response. Errors
    propagate, and the caller treats them as an empty address. ``request_get``
    can be replaced in tests.
    """
    latitude, longitude = coordinates
    response = request_get(
        "https://nominatim.openstreetmap.org/reverse",
        params={
            "format": "jsonv2",
            "addressdetails": 1,
            "lat": latitude,
            "lon": longitude,
        },
        headers={"User-Agent": "memorymap-toolkit/1.0"},
        timeout=20,
    )
    response.raise_for_status()
    time.sleep(delay)
    return response.json().get("address", {})


def resolve_coordinates(
    name: str,
    record: dict,
    source: SourceLocation | None,
    *,
    own_geocoder: Callable = query_geocoder,
    nominatim: Callable = geocode_nominatim,
    geo_paths: dict[str, tuple[str, ...]] | None = None,
) -> tuple[tuple[float, float], str, dict]:
    """Resolve coordinates in source-priority order and return provenance.

    Direct and identifier-based sources come before the network geocoders,
    which are slower and rate limited. The order is: direct source coordinates
    (chronotopi, or metadati birth places); for a reviewed query override, the
    local geocoder and then Nominatim; GeoNames and Wikidata by the IDs from
    ``source_ids``; then the local geocoder and Nominatim with the query from
    ``build_query``. A failed GeoNames or Wikidata lookup is ignored. The
    function writes no files.

    Returns:
        The ``(lat, lon)`` pair, the provenance label for ``coordinate_source``
        (``chronotopi`` for any direct source coordinates, otherwise
        ``own-geocoder``, ``nominatim``, ``geonames``, or ``wikidata``), and the
        Nominatim address dictionary, which is empty unless Nominatim supplied
        the coordinates.

    Raises:
        RuntimeError: When no source yields valid coordinates.
    """
    direct = choose_direct_coordinates(source) if source else None
    if direct:
        return direct, "chronotopi", {}

    query = build_query(name, source, record, geo_paths)
    if taxonomy_query_override(name) is not None:
        own = own_geocoder(query)
        if own:
            pair = coordinate_pair(own.get("lat"), own.get("long"))
            if pair:
                return pair, "own-geocoder", {}
        result = nominatim(query)
        if result:
            return result["coordinates"], "nominatim", result.get("address", {})

    geonames_id, wikidata_id = source_ids(record, source)
    if geonames_id:
        url = f"https://www.geonames.org/{geonames_id}"
        try:
            coordinates = extract_geonames_coordinates(url)
        except Exception:
            coordinates = None
        if coordinates:
            pair = coordinate_pair(coordinates.get("lat"), coordinates.get("long"))
            if pair:
                return pair, "geonames", {}
    if wikidata_id:
        url = f"https://www.wikidata.org/wiki/{wikidata_id}"
        try:
            coordinates = extract_wikidata_coordinates(url)
        except Exception:
            coordinates = None
        if coordinates:
            pair = coordinate_pair(coordinates.get("lat"), coordinates.get("long"))
            if pair:
                return pair, "wikidata", {}

    own = own_geocoder(query)
    if own:
        pair = coordinate_pair(own.get("lat"), own.get("long"))
        if pair:
            return pair, "own-geocoder", {}

    result = nominatim(query)
    if result:
        return result["coordinates"], "nominatim", result.get("address", {})
    raise RuntimeError(f"No coordinates found for {name!r} using {query!r}")


def _taxonomy_parent_chain(name: str, geo_paths: dict[str, tuple[str, ...]]):
    """Return the GO path that contains a location name, from broad to narrow.

    Reviewed exclusions, and names that are not in the GO taxonomy, give an
    empty list. A reviewed parent hint is used when its label has a GO path, so
    the location is placed under that parent. Otherwise a name that is itself a
    GO label gets only its ancestors, while a reviewed alias gets the full path,
    which ends with the canonical label.
    """
    if taxonomy_match_excluded(name):
        return []
    canonical = taxonomy_alias(name)
    parent_label = taxonomy_location_parent(canonical) or taxonomy_location_parent(name)
    if parent_label:
        parent_path = _geo_path_for(taxonomy_alias(parent_label), geo_paths)
        if parent_path:
            return list(parent_path)
    path = _geo_path_for(canonical, geo_paths)
    if not path:
        return []
    if _label_key(canonical) == _label_key(name):
        return list(path[:-1])
    return list(path)


def _geo_path_for(label: str, geo_paths: dict[str, tuple[str, ...]]):
    """Find a GO path using exact or Latin-transliterated label keys.

    The exact normalized key is tried first. If it fails, the label and every
    GO label are transliterated and casefolded, so a spelling such as
    ``Gießen`` matches ``Giessen``. Returns the first match, or None.
    """
    path = geo_paths.get(_label_key(label))
    if path:
        return path
    latin_key = latinize_text(label).casefold()
    return next(
        (
            path
            for candidate, path in geo_paths.items()
            if latinize_text(candidate).casefold() == latin_key
        ),
        None,
    )


def wider_region_for(
    name: str,
    record: dict,
    source: SourceLocation | None,
    geo_paths: dict[str, tuple[str, ...]],
    address: dict,
) -> str:
    """Return broad-to-nearest containment for a location row.

    The first step that produces a path wins. The steps are the path from
    ``_taxonomy_parent_chain``; the longest GO label that occurs in the
    normalized name; the source details and the existing ``super_region``, each
    passed through the alias table; and the address parts from geocoding. If no
    address part is a GO label, the path is the GO root followed by the address
    parts, which can include names outside the taxonomy. Source details form a
    set, so when several of them match, the one tried first can differ between
    runs. Returns an empty string when no step matches.
    """
    exact = _taxonomy_parent_chain(name, geo_paths)
    if exact:
        return " > ".join(exact)

    normalized_name = _label_key(name)
    embedded_candidates = [
        (label, path)
        for label, path in geo_paths.items()
        if label and label in normalized_name
    ]
    embedded = max(embedded_candidates, key=lambda item: len(item[0]), default=None)
    if embedded:
        return " > ".join(embedded[1])

    context_candidates = []
    if source:
        context_candidates.extend(source.details)
    existing_parent = clean_value(record.get("super_region"))
    if existing_parent:
        context_candidates.append(existing_parent)
    for candidate in context_candidates:
        alias = taxonomy_alias(candidate)
        path = _geo_path_for(alias, geo_paths)
        if path:
            return " > ".join(path)

    address_names = [
        address.get("city"),
        address.get("town"),
        address.get("village"),
        address.get("municipality"),
        address.get("state"),
        address.get("country"),
    ]
    for candidate in address_names:
        if not candidate:
            continue
        alias = taxonomy_alias(candidate)
        path = _geo_path_for(alias, geo_paths)
        if path:
            return " > ".join(path)
    values = [value for value in address_names if value]
    if values:
        return " > ".join(["GO – geographischer Ort", *reversed(values)])
    return ""


def rebuild_workbook(
    input_path: Path = DEFAULT_INPUT,
    output_path: Path = DEFAULT_OUTPUT,
    data_root: Path = DEFAULT_WORKBOOKS_ROOT,
    taxonomy_path: Path = DEFAULT_TAXONOMY,
    *,
    own_geocoder: Callable = query_geocoder,
    nominatim: Callable = geocode_nominatim,
    reverse_nominatim: Callable = reverse_geocode_nominatim,
) -> dict[str, int]:
    """Rebuild the workbook and raise if any retained row lacks coordinates.

    Reads the existing rows from ``input_path``, adds the chronotopi and
    metadati evidence under ``data_root``, and resolves coordinates and
    containment against the GO taxonomy at ``taxonomy_path``. The result is
    written to ``output_path`` as one sheet named ``locations``. The default
    output is the input path, so the file is overwritten. The existing rows are
    read into memory before the file is saved. Network calls go through the
    keyword arguments and the GeoNames and Wikidata helpers. Nothing is sent to
    the Memory Map API.

    Returns:
        Counts with the keys ``rows`` (rows written), ``dropped`` (rows with no
        usable name or reviewed as non-locations), and ``coordinates_<source>``
        for each coordinate provenance label.

    Raises:
        RuntimeError: When a row cannot be geocoded, or the output would be
            empty or hold invalid coordinates. Nothing is written in that case.
        ValueError: When the input has no ``super_region`` column.
    """
    headers, existing = read_existing_workbook(input_path)
    sources = read_chronotopi_sources(data_root)
    merge_source_maps(sources, read_metadata_sources(data_root))
    headers = extend_headers(headers, sources)
    merged = merge_sources(headers, existing, sources)
    taxonomy = load_taxonomy(taxonomy_path)
    geo_paths = build_geo_paths(taxonomy)
    headers = order_output_headers(headers)

    output_rows = []
    dropped = 0
    source_counts = Counter()
    for key in sorted(
        merged, key=lambda item: clean_value(merged[item].get("location")).casefold()
    ):
        record = merged[key]
        name = clean_value(record.get("location"))
        if not name or taxonomy_reviewed_unmapped(name):
            dropped += 1
            continue
        source = sources.get(key)
        coordinates, coordinate_source, address = resolve_coordinates(
            name,
            record,
            source,
            own_geocoder=own_geocoder,
            nominatim=nominatim,
            geo_paths=geo_paths,
        )
        record["lat"], record["long"] = coordinates
        record["bag_of_words"] = " ".join(sorted(tokenize_location(name)))
        if not wider_region_for(name, record, source, geo_paths, address):
            try:
                address = reverse_nominatim(coordinates)
            except Exception:
                address = {}
        wider = wider_region_for(name, record, source, geo_paths, address)
        record["wider_region"] = wider
        record["super_region"] = latinize_text(wider.split(" > ")[-1]) if wider else ""
        record["coordinate_source"] = coordinate_source
        for header in headers:
            record.setdefault(header, "")
        output_rows.append(record)
        source_counts[coordinate_source] += 1

    if not output_rows or any(
        not coordinate_pair(row.get("lat"), row.get("long")) for row in output_rows
    ):
        raise RuntimeError(
            "Rebuilt workbook still contains rows without valid coordinates"
        )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "locations"
    sheet.append(headers)
    for record in output_rows:
        sheet.append([record.get(header, "") for header in headers])
    workbook.save(output_path)
    return {
        "rows": len(output_rows),
        "dropped": dropped,
        **{f"coordinates_{key}": value for key, value in sorted(source_counts.items())},
    }


if __name__ == "__main__":
    print(rebuild_workbook())
