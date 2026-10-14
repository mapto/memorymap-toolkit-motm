"""Region tokens and super-region enrichment for the workbook."""

import regex as re


# --- Bag-of-words & super-region logic ---


def tokenize_location(name):
    """Tokenize a location name into a set of lowercase words, stripping punctuation."""
    name = str(name).strip().lower()
    name = re.sub(r"[^\w\s]", "", name)
    return set(name.split()) - {""}


def compute_super_regions(locations):
    """Given a dict of {name: bag_of_words}, return {name: super_region_name}.

    A location A is the super-region of B if A's bag is a strict subset of B's bag.
    Among all subsets, pick the one with the most words (closest ancestor).
    """
    items = [(name, bow) for name, bow in locations.items() if bow]
    result = {}
    for name, bow in items:
        best_parent = None
        best_len = 0
        for other_name, other_bow in items:
            if other_name == name:
                continue
            if other_bow < bow and len(other_bow) > best_len:
                best_parent = other_name
                best_len = len(other_bow)
        result[name] = best_parent or ""
    return result


def enrich_locations_xlsx(xlsx_path):
    """Add bag_of_words and super_region columns to an existing locations.xlsx.

    Only adds information; never removes existing data.
    """
    import openpyxl as _xl

    wb = _xl.load_workbook(xlsx_path)
    ws = wb.active
    assert ws is not None

    # Read header row
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    # Ensure new columns exist
    bow_col = None
    sr_col = None
    if "bag_of_words" in headers:
        bow_col = headers.index("bag_of_words") + 1
    if "super_region" in headers:
        sr_col = headers.index("super_region") + 1

    if bow_col is None:
        bow_col = 2  # insert after column 1
        ws.insert_cols(2)
        ws.cell(1, 2).value = "bag_of_words"  # type: ignore[reportAttributeAccessIssue]
        # Shift sr_col if it existed
        if sr_col is not None:
            sr_col += 1
        # Re-read headers after insert
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    if sr_col is None:
        sr_col = bow_col + 1
        ws.insert_cols(sr_col)
        ws.cell(1, sr_col).value = "super_region"  # type: ignore[reportAttributeAccessIssue]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    # Build bag-of-words for all locations
    loc_bows = {}  # row -> (name, bow_set)
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        bow = tokenize_location(name)
        ws.cell(r, bow_col).value = " ".join(sorted(bow))  # type: ignore[reportAttributeAccessIssue]
        loc_bows[r] = (str(name).strip(), bow)

    # Compute super-regions
    name_to_bow = {name: bow for _, (name, bow) in loc_bows.items()}
    super_regions = compute_super_regions(name_to_bow)

    # Write super-region column (only if currently empty)
    for r, (name, _) in loc_bows.items():
        existing = ws.cell(r, sr_col).value
        if not existing or str(existing).strip() in ("", "nan"):
            sr = super_regions.get(name, "")
            if sr:
                ws.cell(r, sr_col).value = sr

    wb.save(xlsx_path)
