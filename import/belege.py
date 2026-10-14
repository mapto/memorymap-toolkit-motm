"""Belege (Extraction) import logic.

A Belege workbook has one row per citation: a quoted passage stored as an
``Extraction`` that points at the ``Interview`` it came from. Interviews must
exist before the citations that reference them, so ``import_belege`` upserts
them first and then links each citation to its Interview's database ID. Rows
also name speakers, mentioned persons, and topics, which are resolved to
database IDs. Interviews and Extractions are matched by archive ID and
identifier, so rerunning an import updates existing records.

A row's ``mentioned_persons`` cell usually names a Person by its
``identifier``, but some sheets instead use a per-interview role label such as
``IS_E_00002_self``. The ``_self`` label is resolved to the row's own
Interview's interviewee; other role labels (mother, father, ...) have no
Person record to link to and are left unresolved.

The ``markierung`` (classification) cell names a highlight color, which is
converted to an editorial priority rank at import time rather than stored as
text: gelb=1 (high), cyan=2 (medium), ohne=3 (low). See
``CLASSIFICATION_PRIORITY``/``resolve_priority``.

``import_belege`` is the entry point. ``belege.ipynb`` calls it for each
``belege_*.xlsx`` file under ``data/nuove schede``, and ``import.sh`` runs that
notebook.

Usage::

    from api_client import login
    from belege import import_belege

    login()
    import_belege("path/to/belege_<name>.xlsx")
"""

from __future__ import annotations

import re
from pathlib import Path

from fast_langdetect import detect as detect_language
from openpyxl import load_workbook
from tqdm.auto import tqdm

from api_client import (
    api_bulk_upsert,
    api_get,
    build_interview_payload,
    get_or_create_concept,
    normalize_interview_id,
    parse_interview_id_from_quelle,
    upsert_interviews,
)
from utils import clean_str, is_empty

# Some belege sheets name mentioned persons with a per-interview role label
# (e.g. "IS_E_00002_self", "IS_E_00002_mother") instead of a Person.identifier.
# Only the "_self" role unambiguously maps to a known Person: the interview's
# own interviewee. Other roles (mother, father, ...) have no Person record to
# link to and are left unresolved.
SELF_MENTION_RE = re.compile(r"_self\s*$", re.IGNORECASE)

# The ``markierung`` cell marks a quote's highlight color, which is imported
# as an editorial priority rank instead of the raw color name: gelb (yellow
# highlight) is most important, cyan is secondary, ohne (no highlight) is
# least important. A value outside this map (including blank) has no priority.
CLASSIFICATION_PRIORITY = {
    "gelb": 1,
    "cyan": 2,
    "ohne": 3,
}


def resolve_priority(raw_classification) -> int | None:
    """Map a sheet's raw ``markierung`` value to its editorial priority rank.

    Comparison is case-insensitive and ignores surrounding whitespace. Values
    outside ``CLASSIFICATION_PRIORITY`` (including blank cells) resolve to
    ``None``, so the Extraction keeps no priority rather than an invented one.
    """
    classification = clean_str(raw_classification).lower()
    return CLASSIFICATION_PRIORITY.get(classification)


def read_belege_rows(xlsx_path: str | Path, sheet_name: str = "Belege") -> list[dict]:
    """Read the Belege sheet into named row dictionaries.

    The first three sheet rows (title, spacer, and header) are skipped, and blank
    rows are dropped. Cells are mapped to the names in ``fields`` by position
    only, so the header text is never checked and a reordered sheet would be read
    wrongly. Short rows are padded with ``None`` and extra cells are ignored.
    Keeping this mapping here prevents the import orchestration and the payload
    builders from depending on positional Excel indexes.
    """
    workbook = load_workbook(xlsx_path)
    try:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()

    fields = (
        "identifier",
        "quelle",
        "interview_id",
        "interview_date",
        "interviewee_identifier",
        "mentioned_persons",
        "timecode",
        "concepts",
        "quote",
        "classification",
        "event_ids",
        "event_confidence",
        "notes",
    )
    named_rows = []
    for row in rows[3:]:
        if not row or all(cell is None for cell in row):
            continue
        cells = (list(row) + [None] * len(fields))[: len(fields)]
        named_rows.append({field: cells[index] for index, field in enumerate(fields)})
    return named_rows


def load_person_map() -> dict[str, int]:
    """Load all source identifiers needed by citation imports.

    Makes one paginated ``persons`` request and maps each Person ``identifier``
    to its database ID, skipping Persons without an identifier. The map resolves
    ``interviewee_identifier`` and the first ``mentioned_persons`` entry of each
    row. An identifier that is not in the map leaves the link empty, with no
    error.
    """
    return {
        person["identifier"]: person["id"]
        for person in api_get("persons")
        if person.get("identifier")
    }


def load_interview_map() -> dict[str, int]:
    """Load existing Interview archive IDs for idempotent reimports.

    Makes one paginated ``interviews`` request and skips Interviews without an
    ``archive_id``. Keys are normalized so they match the IDs written by this
    run. The map lets citations link to Interviews that earlier runs created,
    even when this run upserts none for them.
    """
    return {
        normalize_interview_id(interview["archive_id"]): interview["id"]
        for interview in api_get("interviews")
        if interview.get("archive_id")
    }


def load_interview_interviewee_map() -> dict[str, int]:
    """Map normalized archive IDs to their Interview's interviewee Person ID.

    Makes one paginated ``interviews`` request and skips Interviews without an
    ``archive_id`` or ``interviewee``. Used to resolve ``mentioned_persons``
    tokens written as a role label (``<archive_id>_self``) rather than a
    Person identifier.
    """
    return {
        normalize_interview_id(interview["archive_id"]): interview["interviewee"]
        for interview in api_get("interviews")
        if interview.get("archive_id") and interview.get("interviewee")
    }


def archive_id_for_row(row: dict) -> str | None:
    """Resolve a row's explicit archive ID or derive it from ``quelle``.

    The ``interview_id`` cell takes priority. Its value is cleaned with
    ``clean_str`` but not zero-padded, so ``IS_E_170`` is looked up as written and
    misses an Interview stored as ``IS_E_00170``. Without that cell the ID comes
    from ``parse_interview_id_from_quelle``. Returns ``None`` when neither source
    gives an ID.
    """
    archive_id = clean_str(row["interview_id"])
    if archive_id:
        return archive_id
    return parse_interview_id_from_quelle(clean_str(row["quelle"]))


def build_row_interview_payload(row: dict, person_map: dict[str, int]) -> dict | None:
    """Build the shared Interview payload for one citation row.

    The archive ID comes from ``archive_id_for_row``, the interviewee is the
    ``person_map`` entry for ``interviewee_identifier``, and the date goes through
    ``build_interview_payload``, which normalizes it. Returns ``None`` when the
    row has no archive ID. This function makes no API call.
    """
    return build_interview_payload(
        archive_id_for_row(row),
        interviewee=person_map.get(clean_str(row["interviewee_identifier"])),
        interview_date=row["interview_date"],
    )


def detect_quote_language(quote: str) -> str | None:
    """Detect a quote language when the detector has enough confidence.

    Quotes shorter than 20 characters are not classified, so a fragment does not
    receive a language label from a guess. Otherwise the top ``fast_langdetect``
    result is used only when its score is above 0.7. Detector exceptions are
    swallowed, so a detection failure leaves ``language`` unset instead of
    stopping the import.
    """
    if len(quote) < 20:
        return None
    try:
        result = detect_language(quote)
        if result and float(result[0]["score"]) > 0.7:
            return str(result[0]["lang"])
    except Exception:
        pass
    return None


def resolve_people_mentioned(
    row: dict,
    mentioned: str,
    person_map: dict[str, int],
    interviewee_map: dict[str, int],
) -> int | None:
    """Resolve the ``mentioned_persons`` cell to a single Person ID.

    Takes the first identifier in the comma- or newline-separated cell, because
    the model field holds a single person. When that first token is a known
    Person identifier, it wins, matching older sheets that name people
    directly. Otherwise, if any token in the cell is a per-interview "self"
    role label (``<archive_id>_self``), it resolves to the row's own
    Interview's interviewee via ``interviewee_map``. Other role labels
    (mother, father, ...) have no Person record and are dropped.
    """
    tokens = [token.strip() for token in mentioned.replace("\n", ",").split(",") if token.strip()]
    if not tokens:
        return None
    person_id = person_map.get(tokens[0])
    if person_id:
        return person_id
    if any(SELF_MENTION_RE.search(token) for token in tokens):
        archive_id = archive_id_for_row(row)
        if archive_id:
            return interviewee_map.get(normalize_interview_id(archive_id))
    return None


def build_extraction_payload(
    row: dict,
    person_map: dict[str, int],
    interviewee_map: dict[str, int] | None = None,
) -> dict | None:
    """Build one Extraction payload, excluding its Interview foreign key.

    Returns ``None`` when the row has no ``identifier`` (the ``beleg_id`` cell).
    ``language`` is added only when ``detect_quote_language`` finds one.
    ``people_mentioned`` is resolved by ``resolve_people_mentioned``.
    ``priority`` is resolved by ``resolve_priority`` from the raw
    ``classification`` cell and omitted when it doesn't map to a known rank.
    Each comma-separated topic goes through ``get_or_create_concept``, which
    makes API calls, and topics that resolve to no ID are omitted. The
    Interview foreign key is added later by ``attach_interview_ids``.
    """
    identifier = clean_str(row["identifier"])
    if not identifier:
        return None

    payload: dict = {
        "identifier": identifier,
        "timecode": clean_str(row["timecode"]),
        "quote": clean_str(row["quote"]),
        "event_confidence": clean_str(row["event_confidence"]),
        "notes": clean_str(row["notes"]),
    }
    priority = resolve_priority(row["classification"])
    if priority:
        payload["priority"] = priority
    language = detect_quote_language(payload["quote"])
    if language:
        payload["language"] = language

    mentioned = clean_str(row["mentioned_persons"])
    if mentioned:
        person_id = resolve_people_mentioned(
            row, mentioned, person_map, interviewee_map or {}
        )
        if person_id:
            payload["people_mentioned"] = person_id

    concept_ids = []
    for topic in clean_str(row["concepts"]).split(","):
        if not is_empty(topic):
            concept_id = get_or_create_concept(topic.strip())
            if concept_id:
                concept_ids.append(concept_id)
    if concept_ids:
        payload["concepts"] = concept_ids
    return payload


def prepare_belege_payloads(
    rows: list[dict],
    person_map: dict[str, int],
    interviewee_map: dict[str, int] | None = None,
) -> tuple[list[dict], list[tuple[str | None, dict]]]:
    """Prepare Interview payloads and citation payloads separately.

    Every row with an archive ID yields one Interview payload. Repeats are kept
    here because ``upsert_interviews`` merges them. Each citation is paired with
    its archive ID, so ``attach_interview_ids`` can add the foreign key once the
    Interviews exist. Rows without a usable payload are skipped. A tqdm progress
    bar is shown, and concept lookups happen in ``build_extraction_payload``.

    Returns:
        A tuple of the Interview payload list and the citation list, where each
        citation entry is an ``(archive_id, payload)`` pair.
    """
    interview_payloads = []
    extraction_payloads = []
    for row in tqdm(rows, desc="Belege", unit="row"):
        interview_payload = build_row_interview_payload(row, person_map)
        if interview_payload:
            interview_payloads.append(interview_payload)
        extraction_payload = build_extraction_payload(row, person_map, interviewee_map)
        if extraction_payload:
            extraction_payloads.append((archive_id_for_row(row), extraction_payload))
    return interview_payloads, extraction_payloads


def attach_interview_ids(
    prepared_payloads: list[tuple[str | None, dict]],
    interview_map: dict[str, int],
) -> tuple[list[dict], list[str]]:
    """Add Interview foreign keys to prepared citations and report misses.

    A citation whose archive ID is in *interview_map* gets that database ID as
    ``interview``. A citation whose ID is missing gets an error message but is
    still returned without the foreign key, so the Extraction is written and the
    gap is reported. Citations without an archive ID pass through unchanged and
    without a message.

    Returns:
        A tuple of the payload list and the list of error strings.
    """
    payloads = []
    errors = []
    for archive_id, payload in prepared_payloads:
        if archive_id:
            interview_id = interview_map.get(archive_id)
            if interview_id:
                payload["interview"] = interview_id
            else:
                errors.append(
                    f"citation {payload['identifier']}: "
                    f"Interview {archive_id} was not created"
                )
        payloads.append(payload)
    return payloads, errors


def bulk_upsert_extractions(payloads: list[dict]) -> tuple[list[int], list[str]]:
    """Write citation payloads in batches of 200 and collect failures.

    Each batch is one ``api_bulk_upsert`` request. The server matches Extractions
    by ``identifier``, so a rerun updates existing records and fills in any
    Interview link they lack. A batch that raises is recorded as an error and the
    next batch still runs, so one failure does not discard the rest. Per-item
    errors from the response are added to the same list.

    Returns:
        A tuple of the IDs of created and updated items and the list of error
        strings.
    """
    results = []
    errors = []
    for start in range(0, len(payloads), 200):
        try:
            response = api_bulk_upsert("extractions", payloads[start : start + 200])
            imported_items = [
                *response.get("created_items", []),
                *response.get("updated_items", []),
            ]
            results.extend(
                item["id"] for item in imported_items if item.get("id") is not None
            )
            errors.extend(
                f"bulk item {item.get('index', '?')}: {item.get('error', item)}"
                for item in response.get("errors", [])
            )
        except Exception as exc:
            errors.append(f"bulk batch {start}: {exc}")
    return results, errors


def import_belege(xlsx_path: str | Path, sheet_name: str = "Belege") -> list[int]:
    """Import Belege rows and link each citation to its Interview source.

    The steps run in order: read the rows, load the Person, Interview, and
    interviewee maps with three list requests, prepare the payloads (which
    resolves topics to concept IDs), upsert the Interviews, link each citation,
    and upsert the citations in batches. Interviews go first because each
    Extraction stores its Interview's database ID. Rerunning the importer
    therefore also repairs existing citations whose ``interview_id`` was empty.

    Errors from the bulk writes are collected and printed as a summary rather than
    raised. A missing file raises ``FileNotFoundError``, a missing sheet raises
    ``KeyError``, and HTTP errors from the list requests or from concept creation
    propagate. Returns the IDs of the created and updated Extractions.
    """
    xlsx_path = Path(xlsx_path)
    rows = read_belege_rows(xlsx_path, sheet_name)
    person_map = load_person_map()
    interview_map = load_interview_map()
    interviewee_map = load_interview_interviewee_map()
    interview_payloads, prepared_payloads = prepare_belege_payloads(
        rows, person_map, interviewee_map
    )

    created_interviews, errors = upsert_interviews(interview_payloads)
    interview_map.update(created_interviews)
    payloads, linkage_errors = attach_interview_ids(prepared_payloads, interview_map)
    results, extraction_errors = bulk_upsert_extractions(payloads)
    errors.extend(linkage_errors)
    errors.extend(extraction_errors)

    if errors:
        print(f"\n\u26a0 {len(errors)} errors:")
        for error in errors:
            print(f"  \u274c {error}")
    print(f"\u2705 Imported {len(results)} extractions from {xlsx_path.name}")
    return results
